"""Unidad 10 · Costeo por absorción y costeo variable.

Recalcula todas las cifras de la unidad desde los datos, genera la planilla
`planillas/u10_absorcion_variable.xlsx` con los dos estados de resultados del caso y del ejemplo
(con fórmulas, no valores pegados) y comprueba sus resultados evaluando esas fórmulas en Python.
Con --revisar busca cada resultado en el .md.

    py -3 -P u10_absorcion_variable.py            # imprime los resultados y arma la planilla
    py -3 -P u10_absorcion_variable.py --revisar  # además los busca en el .md

Todos los datos son originales de esta base (pasados por herramientas/colisiones.py). El caso es
el de caso_estampa_sur_contable.py, que se importa sin tocarlo. Los inventarios salen en FIFO.
"""
from pathlib import Path
import importlib.util, re, sys

sys.stdout.reconfigure(encoding="utf-8")
AQUI = Path(__file__).resolve().parent


def fmt(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


# ---------------------------------------------------------------- el mes con los dos métodos

def mes(p, prod, vend, inv0=0, inv0_var=0.0, inv0_abs=0.0, fin=0.0):
    """Un producto, un mes. p: precio, mpd, mod, cifv, ciff, gavv, gavf. Inventario inicial en FIFO."""
    cvu = p["mpd"] + p["mod"] + p["cifv"]
    fu = p["ciff"] / prod if prod else 0.0
    absu = cvu + fu
    de_inv = min(vend, inv0)
    nuevos = vend - de_inv
    ventas = vend * p["precio"]
    cvv = de_inv * inv0_var + nuevos * cvu
    cva = de_inv * inv0_abs + nuevos * absu
    inv_fin = inv0 + prod - vend
    fin_var = (inv0 - de_inv) * inv0_var + (prod - nuevos) * cvu
    fin_abs = (inv0 - de_inv) * inv0_abs + (prod - nuevos) * absu
    gavv = vend * p["gavv"]
    mc = ventas - cvv - gavv
    ro_v = mc - p["ciff"] - p["gavf"]
    mb = ventas - cva
    ro_a = mb - gavv - p["gavf"]
    fijo_fin = fin_abs - fin_var
    fijo_ini = inv0 * (inv0_abs - inv0_var)
    assert abs((ro_a - ro_v) - (fijo_fin - fijo_ini)) < 1e-6
    return dict(cvu=cvu, fu=fu, absu=absu, ventas=ventas, cvv=cvv, cva=cva, gavv=gavv, mc=mc, ro_v=ro_v,
                mb=mb, ro_a=ro_a, dif=ro_a - ro_v, inv_fin=inv_fin, fin_var=fin_var, fin_abs=fin_abs,
                fijo_fin=fijo_fin, fijo_ini=fijo_ini, uai_v=ro_v - fin, uai_a=ro_a - fin,
                fijos=p["ciff"] + p["gavf"], gav=gavv + p["gavf"])


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · fábrica de velas aromáticas
VELAS = dict(precio=6_500, mpd=1_450, mod=620, cifv=280, ciff=2_640_000, gavv=190, gavf=1_150_000)
VELAS_MESES = dict(mes1=(4_400, 4_000), mes2=(3_600, 4_000), igual=(4_000, 4_000))
# 10.1 · jabones artesanales
E1 = dict(q=6_000, mpd=640, mod=310, cifv=150, ciff=2_100_000, inv=800)
# 10.2, 10.3, 10.5, 10.6 · cajas de cartón para delivery, tres meses seguidos
CAJAS = dict(precio=515, mpd=162, mod=47, cifv=23, ciff=1_396_800, gavv=19, gavf=683_000)
CAJAS_MESES = [(11_640, 11_640), (14_550, 11_640), (9_700, 12_610)]
# 10.4 · taller de cuchillos
CUCH = dict(precio=24_900, mpd=6_140, mod=3_460, cifv=1_170, ciff=1_692_000, gavv=830, gavf=915_000)
CUCH_Q = (470, 385)
# 10.7 · fábrica de sillas y el bono del gerente
SILLAS = dict(precio=38_000, mpd=17_000, mod=0, cifv=0, ciff=16_000_000, gavv=0, gavf=6_000_000)
SILLAS_Q = dict(vend=2_000, opciones=[2_000, 2_800], bono=0.05)
# 10.8 y 10.9 · bloqueador solar
BLOQ = dict(precio=7_940, mpd=1_870, mod=715, cifv=395, ciff=8_316_000, gavv=585, gavf=4_270_000)
BLOQ_Q = dict(prod=9_240, vend=7_630, fin=517_000)
PEDIDO = dict(q=1_050, precio=3_690, flete=145)
# 10.10 · problema inverso
E10 = dict(r_abs=4_380_000, r_var=3_660_000, prod=8_000, vend=6_200, prod2=6_000, vend2=7_800)


# ---------------------------------------------------------------- resultados

def resultados():
    R = {}
    m1 = mes(VELAS, *VELAS_MESES["mes1"])
    m2 = mes(VELAS, *VELAS_MESES["mes2"], inv0=m1["inv_fin"], inv0_var=m1["cvu"], inv0_abs=m1["absu"])
    ig = mes(VELAS, *VELAS_MESES["igual"])
    assert abs((m1["ro_a"] + m2["ro_a"]) - (m1["ro_v"] + m2["ro_v"])) < 1e-6
    R["ej"] = dict(m1=m1, m2=m2, ig=ig, suma=m1["ro_a"] + m2["ro_a"],
                   costo_m2_inv=m1["inv_fin"] * m1["absu"], costo_m2_nuevos=(VELAS_MESES["mes2"][1] - m1["inv_fin"]) * m2["absu"])

    e = E1
    cvu = e["mpd"] + e["mod"] + e["cifv"]
    fu = e["ciff"] / e["q"]
    R["e1"] = dict(cvu=cvu, fu=fu, absu=cvu + fu, inv_var=e["inv"] * cvu, inv_abs=e["inv"] * (cvu + fu),
                   dif=e["inv"] * fu)

    c1 = mes(CAJAS, *CAJAS_MESES[0])
    c2 = mes(CAJAS, *CAJAS_MESES[1], inv0=c1["inv_fin"], inv0_var=c1["cvu"], inv0_abs=c1["absu"])
    c3 = mes(CAJAS, *CAJAS_MESES[2], inv0=c2["inv_fin"], inv0_var=c2["cvu"], inv0_abs=c2["absu"])
    assert c3["inv_fin"] == 0
    R["cajas"] = dict(c1=c1, c2=c2, c3=c3, suma_a=c1["ro_a"] + c2["ro_a"] + c3["ro_a"],
                      suma_v=c1["ro_v"] + c2["ro_v"] + c3["ro_v"],
                      c3_inv=c2["inv_fin"] * c2["absu"], c3_nuevos=CAJAS_MESES[2][0] * c3["absu"])
    assert abs(R["cajas"]["suma_a"] - R["cajas"]["suma_v"]) < 1e-6

    k = mes(CUCH, *CUCH_Q)
    incurrido = CUCH_Q[0] * k["cvu"] + CUCH["ciff"] + k["gavv"] + CUCH["gavf"]
    a_res_v = k["cvv"] + k["gavv"] + CUCH["ciff"] + CUCH["gavf"]
    a_res_a = k["cva"] + k["gavv"] + CUCH["gavf"]
    assert abs(a_res_v + k["fin_var"] - incurrido) < 1e-6 and abs(a_res_a + k["fin_abs"] - incurrido) < 1e-6
    R["e4"] = dict(k=k, incurrido=incurrido, a_res_v=a_res_v, a_res_a=a_res_a)

    s = [mes(SILLAS, q, SILLAS_Q["vend"]) for q in SILLAS_Q["opciones"]]
    R["e7"] = dict(s=s, bono=[x["ro_a"] * SILLAS_Q["bono"] for x in s],
                   caja_inv=(SILLAS_Q["opciones"][1] - SILLAS_Q["vend"]) * SILLAS["mpd"])

    b = mes(BLOQ, BLOQ_Q["prod"], BLOQ_Q["vend"], fin=BLOQ_Q["fin"])
    R["e8"] = b
    # 10.9 · el pedido: se fabrica y se vende además de lo normal; su flete reemplaza al GAV variable normal
    p = PEDIDO
    cvu = b["cvu"]
    contrib_u = p["precio"] - cvu - p["flete"]
    prod2 = BLOQ_Q["prod"] + p["q"]
    fu2 = BLOQ["ciff"] / prod2
    absu2 = cvu + fu2
    vend2 = BLOQ_Q["vend"] + p["q"]
    ventas2 = b["ventas"] + p["q"] * p["precio"]
    gavv2 = b["gavv"] + p["q"] * p["flete"]
    ro_a2 = ventas2 - vend2 * absu2 - gavv2 - BLOQ["gavf"]
    ro_v2 = b["ro_v"] + contrib_u * p["q"]
    inv = BLOQ_Q["prod"] - BLOQ_Q["vend"]
    assert abs((ro_a2 - ro_v2) - inv * fu2) < 1e-6
    R["e9"] = dict(contrib_u=contrib_u, contrib=contrib_u * p["q"], aparente=(p["precio"] - b["absu"]) * p["q"],
                   prod2=prod2, fu2=fu2, absu2=absu2, ro_a2=ro_a2, ro_v2=ro_v2, sube_a=ro_a2 - b["ro_a"],
                   sube_v=ro_v2 - b["ro_v"], menos_inv=inv * (b["fu"] - fu2), cvu_ped=cvu + p["flete"],
                   ventas2=ventas2, cventas2=vend2 * absu2, dif2=ro_a2 - ro_v2)

    e = E10
    dif = e["r_abs"] - e["r_var"]
    inv = e["prod"] - e["vend"]
    fu = dif / inv
    R["e10"] = dict(dif=dif, inv=inv, fu=fu, ciff=fu * e["prod"], dif2=0 - inv * fu,
                    r_abs_pct=dif / e["r_var"])
    return R


# ---------------------------------------------------------------- caso

def cargar_caso():
    spec = importlib.util.spec_from_file_location("caso", AQUI / "caso_estampa_sur_contable.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m, m.calcular()


def caso(m, C):
    cv = [p - mg for p, mg in zip(m.PRECIO, m.MARGEN)]
    dep = sum((q["costo"] - q["residual"]) / q["vida_meses"] for q in m.MAQUINAS.values())
    cif = m.SUELDOS + dep + m.SEGURO_ANUAL / 12 + m.BALANCE_INICIAL["activo"]["Arriendo pagado por anticipado"]
    P, V = m.PRODUCIDAS, m.VENDIDAS
    seg = sum(q * s for q, s in zip(P, m.SEG_EST))
    tasa = cif / seg
    cif_u = [tasa * s for s in m.SEG_EST]
    absu = [c + x for c, x in zip(cv, cif_u)]
    inv = [p - v for p, v in zip(P, V)]
    ventas = sum(v * pr for v, pr in zip(V, m.PRECIO))
    cvv = sum(v * c for v, c in zip(V, cv))
    cva = sum(v * a for v, a in zip(V, absu))
    ga = m.OTROS_FIJOS
    interes = m.PRESTAMO["monto"] * m.PRESTAMO["tasa_mensual"]
    mc = ventas - cvv
    ro_v = mc - cif - ga
    mb = ventas - cva
    ro_a = mb - ga
    fijo_inv = [q * x for q, x in zip(inv, cif_u)]
    var_inv = [q * c for q, c in zip(inv, cv)]
    var_prod = sum(q * c for q, c in zip(P, cv))
    # segundo camino: lo producido = lo vendido + lo que queda, en cada método
    assert abs(var_prod + cif - (cva + sum(var_inv) + sum(fijo_inv))) < 1e-6
    return dict(cv=cv, cif=cif, seg=seg, tasa=tasa, cif_u=cif_u, absu=absu, inv=inv, inv_tot=sum(inv),
                ventas=ventas, cvv=cvv, cva=cva, ga=ga, interes=interes, mc=mc, ro_v=ro_v, mb=mb, ro_a=ro_a,
                uai_v=ro_v - interes, uai_a=ro_a - interes, dif=ro_a - ro_v, fijo_inv=fijo_inv,
                fijo_inv_tot=sum(fijo_inv), var_inv=var_inv, var_inv_tot=sum(var_inv),
                inv_abs_tot=sum(var_inv) + sum(fijo_inv), cif_gasto_abs=cif - sum(fijo_inv),
                var_prod=var_prod, P=P, V=V, precio=m.PRECIO, seg_u=m.SEG_EST,
                prod_tot=sum(P), vend_tot=sum(V))


# ---------------------------------------------------------------- planilla (con fórmulas) y su evaluación

def recta(pts):
    n = len(pts)
    mx = sum(x for x, _ in pts) / n
    my = sum(y for _, y in pts) / n
    b = sum((x - mx) * (y - my) for x, y in pts) / sum((x - mx) ** 2 for x, _ in pts)
    return my - b * mx, b


def evaluar(ws):
    """Evalúa las fórmulas de una hoja (+ − × ÷, SUM, SUMPRODUCT, MIN, MAX)."""
    from openpyxl.utils import range_boundaries, get_column_letter
    memo = {}

    def v(c):
        if c not in memo:
            x = ws[c].value
            memo[c] = calc(x[1:]) if isinstance(x, str) and x.startswith("=") else (x or 0)
        return memo[c]

    def r(a, b):
        c1, r1, c2, r2 = range_boundaries(f"{a}:{b}")
        return [v(f"{get_column_letter(c)}{f}") for f in range(r1, r2 + 1) for c in range(c1, c2 + 1)]

    plano = lambda a: [y for x in a for y in (x if isinstance(x, list) else [x])]
    env = dict(_v=v, _r=r, SUM=lambda *a: sum(plano(a)), SUMPRODUCT=lambda a, b: sum(x * y for x, y in zip(a, b)),
               MIN=lambda *a: min(plano(a)), MAX=lambda *a: max(plano(a)))

    def calc(expr):
        e = re.sub(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)", r'_r("\1\2","\3\4")', expr)
        e = re.sub(r'(?<![A-Z"_])\$?([A-Z]{1,3})\$?(\d+)(?![\d"])', r'_v("\1\2")', e)
        return eval(e, {"__builtins__": {}}, env)

    return v


def planilla(R, K):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    neg = Font(bold=True)

    # Hoja 1 · caso
    ws = wb.active
    ws.title = "Caso Estampa Sur"
    ws.append(["Producto", "Producidas", "Vendidas", "Inventario final", "Costo variable", "Segundos de estampado",
               "CIF fijo por prenda", "Costo por absorción", "Precio"])
    nombres = ["Poleras básicas", "Polerones", "Poleras estampado doble", "Bolsas", "Gorros"]
    for i, n in enumerate(nombres):
        f = i + 2
        ws.append([n, K["P"][i], K["V"][i], f"=B{f}-C{f}", K["cv"][i], K["seg_u"][i], f"=F{f}*$L$4",
                   f"=E{f}+G{f}", K["precio"][i]])
    ws.append(["Total", "=SUM(B2:B6)", "=SUM(C2:C6)", "=SUM(D2:D6)"])
    datos = [("CIF fijos del mes", "=2880000+9840000/96+(14300000-1800000)/60+948000/12+580000"),
             ("Segundos de estampado del mes", "=SUMPRODUCT(B2:B6,F2:F6)"),
             ("CIF fijo por segundo (unidad 11)", "=L2/L3"),
             ("Gastos de administración", K["ga"]), ("Interés del préstamo", "=6400000*0.0115")]
    for j, (t, x) in enumerate(datos, start=2):
        ws[f"K{j}"], ws[f"L{j}"] = t, x
    ws["K9"], ws["L9"] = "COSTEO VARIABLE", ""
    var = [("Ventas", "=SUMPRODUCT(C2:C6,I2:I6)"), ("Costo variable de lo vendido", "=SUMPRODUCT(C2:C6,E2:E6)"),
           ("Margen de contribución", "=L10-L11"), ("CIF fijos (gasto del período)", "=L2"),
           ("Gastos de administración", "=L5"), ("Resultado operacional", "=L12-L13-L14"),
           ("Interés", "=L6"), ("Resultado antes de impuesto", "=L15-L16")]
    for j, (t, x) in enumerate(var, start=10):
        ws[f"K{j}"], ws[f"L{j}"] = t, x
    ws["K19"], ws["L19"] = "COSTEO POR ABSORCIÓN", ""
    ab = [("Ventas", "=SUMPRODUCT(C2:C6,I2:I6)"), ("Costo de ventas", "=SUMPRODUCT(C2:C6,H2:H6)"),
          ("Margen bruto", "=L20-L21"), ("Gastos de administración", "=L5"), ("Resultado operacional", "=L22-L23"),
          ("Interés", "=L6"), ("Resultado antes de impuesto", "=L24-L25")]
    for j, (t, x) in enumerate(ab, start=20):
        ws[f"K{j}"], ws[f"L{j}"] = t, x
    ws["K28"], ws["L28"] = "Diferencia (absorción − variable)", "=L24-L15"
    ws["K29"], ws["L29"] = "CIF fijo en el inventario final", "=SUMPRODUCT(D2:D6,G2:G6)"
    ws["K30"], ws["L30"] = "Inventario final a costo variable", "=SUMPRODUCT(D2:D6,E2:E6)"
    ws["K31"], ws["L31"] = "Inventario final por absorción", "=SUMPRODUCT(D2:D6,H2:H6)"
    for c in ("K9", "K19", "K15", "K24", "K28"):
        ws[c].font = neg
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["K"].width = 36
    ws.column_dimensions["L"].width = 16
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.column > 1 and c.column != 11:
                c.number_format = "#,##0.00" if c.column in (7, 8) or (c.column == 12 and c.row == 4) else "#,##0"

    # Hoja 2 · ejemplo de las velas
    wv = wb.create_sheet("Ejemplo velas")
    par = [("Precio", VELAS["precio"]), ("MPD por vela", VELAS["mpd"]), ("MOD por vela", VELAS["mod"]),
           ("CIF variable por vela", VELAS["cifv"]), ("CIF fijo del mes", VELAS["ciff"]),
           ("GAV variable por vela vendida", VELAS["gavv"]), ("GAV fijo del mes", VELAS["gavf"])]
    wv["A1"], wv["B1"] = "Dato", "Valor"
    for j, (t, x) in enumerate(par, start=2):
        wv[f"A{j}"], wv[f"B{j}"] = t, x
    wv["D1"], wv["E1"], wv["F1"], wv["G1"] = "Partida", "Mes 1", "Mes 2", "Producción = ventas"
    # filas 2 a 19; columnas E (mes 1), F (mes 2) y G (producción = ventas)
    plantilla = [("Producidas", None), ("Vendidas", None), ("Inventario inicial (unidades)", None),
                 ("Costo variable unitario del inventario inicial", None),
                 ("Costo por absorción unitario del inventario inicial", None),
                 ("Costo variable unitario del mes", "=$B$3+$B$4+$B$5"),
                 ("CIF fijo por unidad producida", "=$B$6/{c}2"),
                 ("Costo por absorción unitario del mes", "={c}7+{c}8"),
                 ("Inventario final (unidades)", "={c}4+{c}2-{c}3"),
                 ("Ventas", "={c}3*$B$2"),
                 ("Costo variable de lo vendido (FIFO)", "=MIN({c}3,{c}4)*{c}5+({c}3-MIN({c}3,{c}4))*{c}7"),
                 ("GAV variable", "={c}3*$B$7"),
                 ("Margen de contribución", "={c}11-{c}12-{c}13"),
                 ("Resultado operacional variable", "={c}14-$B$6-$B$8"),
                 ("Costo de ventas por absorción (FIFO)", "=MIN({c}3,{c}4)*{c}6+({c}3-MIN({c}3,{c}4))*{c}9"),
                 ("Margen bruto", "={c}11-{c}16"),
                 ("Resultado operacional por absorción", "={c}17-{c}13-$B$8"),
                 ("Diferencia (absorción − variable)", "={c}18-{c}15")]
    entradas = {"E": [VELAS_MESES["mes1"][0], VELAS_MESES["mes1"][1], 0, 0, 0],
                "F": [VELAS_MESES["mes2"][0], VELAS_MESES["mes2"][1], "=E10", "=E7", "=E9"],
                "G": [VELAS_MESES["igual"][0], VELAS_MESES["igual"][1], 0, 0, 0]}
    for j, (t, f) in enumerate(plantilla, start=2):
        wv[f"D{j}"] = t
        for col in "EFG":
            wv[f"{col}{j}"] = entradas[col][j - 2] if f is None else f.format(c=col)
            wv[f"{col}{j}"].number_format = "#,##0.00" if 5 <= j <= 9 else "#,##0"
    wv.column_dimensions["A"].width = 30
    wv.column_dimensions["D"].width = 48

    destino = AQUI.parent / "planillas" / "u10_absorcion_variable.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)

    v1, v2 = evaluar(ws), evaluar(wv)
    j = R["ej"]
    pares = [(v1("L2"), K["cif"]), (v1("L3"), K["seg"]), (v1("L4"), K["tasa"]), (v1("L10"), K["ventas"]),
             (v1("L11"), K["cvv"]), (v1("L12"), K["mc"]), (v1("L15"), K["ro_v"]), (v1("L17"), K["uai_v"]),
             (v1("L21"), K["cva"]), (v1("L22"), K["mb"]), (v1("L24"), K["ro_a"]), (v1("L26"), K["uai_a"]),
             (v1("L28"), K["dif"]), (v1("L29"), K["fijo_inv_tot"]), (v1("L30"), K["var_inv_tot"]),
             (v1("L31"), K["inv_abs_tot"]), (v1("D7"), K["inv_tot"]), (v1("L6"), K["interes"])]
    pares += [(v1(f"G{i + 2}"), K["cif_u"][i]) for i in range(5)] + [(v1(f"H{i + 2}"), K["absu"][i]) for i in range(5)]
    for col, k in (("E", "m1"), ("F", "m2"), ("G", "ig")):
        mm = j[k]
        pares += [(v2(f"{col}15"), mm["ro_v"]), (v2(f"{col}18"), mm["ro_a"]), (v2(f"{col}16"), mm["cva"]),
                  (v2(f"{col}12"), mm["cvv"]), (v2(f"{col}19"), mm["dif"]), (v2(f"{col}10"), mm["inv_fin"])]
    ok = all(abs(a - b) < 1e-6 for a, b in pares)
    print(f"PLANILLA: {'OK' if ok else 'NO calza'}, {len(pares)} celdas con fórmula comprobadas · {destino.name}")
    return ok


# ---------------------------------------------------------------- textos esperados

def esperados(R, K):
    p = fmt
    e = []
    j = R["ej"]
    for k in ("m1", "m2", "ig"):
        mm = j[k]
        e += [p(mm[x]) for x in ("ventas", "cvv", "cva", "gavv", "mc", "ro_v", "mb", "ro_a", "fijos", "gav")]
        e += [p(abs(mm["dif"])), p(mm["fu"]), p(mm["absu"]), p(mm["fin_abs"]), p(mm["fin_var"])]
    e += [fmt(j["m2"]["fu"], 2), p(j["m1"]["cvu"]), p(j["suma"]), p(j["costo_m2_inv"]), p(j["costo_m2_nuevos"]),
          fmt(j["m2"]["absu"], 2)]
    j = R["e1"]; e += [p(x) for x in j.values()]
    c = R["cajas"]
    for k in ("c1", "c2", "c3"):
        mm = c[k]
        e += [p(mm[x]) for x in ("ventas", "cvv", "cva", "gavv", "mc", "ro_v", "mb", "ro_a", "fu", "absu", "gav")]
        e += [p(abs(mm["dif"]))]
    e += [p(c["suma_a"]), p(c["c3_inv"]), p(c["c3_nuevos"]), p(c["c1"]["fijos"]), p(c["c1"]["cvu"])]
    j = R["e4"]; k = j["k"]
    e += [p(k[x]) for x in ("cvu", "fu", "absu", "cvv", "cva", "gavv", "fin_var", "fin_abs", "ro_v", "ro_a", "dif",
                            "ventas")]
    e += [p(j["incurrido"]), p(j["a_res_v"]), p(j["a_res_a"])]
    j = R["e7"]
    for s in j["s"]:
        e += [p(s["fu"]), p(s["absu"]), p(s["cva"]), p(s["mb"]), p(s["ro_a"]), p(s["ro_v"])]
    e += [p(x) for x in j["bono"]] + [p(j["s"][1]["fin_abs"]), p(j["caja_inv"])]
    k = R["e8"]
    e += [p(k[x]) for x in ("ventas", "cvu", "fu", "absu", "cvv", "gavv", "mc", "fijos", "ro_v", "uai_v", "cva",
                            "mb", "gav", "ro_a", "uai_a", "dif")]
    j = R["e9"]
    e += [p(j[x]) for x in ("contrib_u", "contrib", "ro_a2", "ro_v2", "sube_a", "sube_v", "menos_inv", "cvu_ped",
                            "ventas2", "cventas2", "dif2")]
    e += [p(-j["aparente"]), fmt(j["fu2"], 2), fmt(j["absu2"], 2), p(j["prod2"])]
    j = R["e10"]; e += [p(j["dif"]), p(j["inv"]), p(j["fu"]), p(j["ciff"]), p(-j["dif2"])]
    # caso
    e += [p(K[x]) for x in ("cif", "seg", "ventas", "cvv", "cva", "ga", "interes", "mc", "ro_v", "mb", "ro_a",
                            "uai_v", "uai_a", "dif", "fijo_inv_tot", "var_inv_tot", "inv_abs_tot", "cif_gasto_abs",
                            "var_prod", "prod_tot", "vend_tot", "inv_tot")]
    e += [fmt(K["tasa"], 4), fmt(K["cif"], 2), fmt(K["fijo_inv_tot"], 2)]
    e += [fmt(x, 2) for x in K["cif_u"]] + [fmt(x, 2) for x in K["absu"]] + [p(x) for x in K["fijo_inv"]]
    e += [p(x) for x in K["var_inv"]] + [str(x) for x in K["inv"]]
    # «las filas redondeadas suman un peso menos» que el total exacto redondeado
    assert sum(round(x) for x in K["fijo_inv"]) == round(K["fijo_inv_tot"]) - 1
    e += [p(K["cif"] + K["ga"]), p(K["var_prod"] + K["cif"])]
    e += [fmt(100 * (c["c2"]["ro_a"] / c["c2"]["ro_v"] - 1), 0) + " %", p(R["e9"]["sube_v"] - R["e9"]["sube_a"]),
          p(R["e7"]["s"][1]["ro_a"] - R["e7"]["s"][0]["ro_a"]), p(R["e7"]["bono"][1] - R["e7"]["bono"][0])]
    return e


def imprimir(R, K):
    def show(d, ind="   "):
        for kk, v in d.items():
            if isinstance(v, dict):
                print(f"{ind}{kk}:")
                show(v, ind + "   ")
            elif isinstance(v, list):
                print(f"{ind}{kk}: {[round(x, 4) if isinstance(x, float) else (x if not isinstance(x, dict) else '{...}') for x in v]}")
                for x in v:
                    if isinstance(x, dict):
                        show(x, ind + "   ")
            else:
                print(f"{ind}{kk}: {round(v, 4) if isinstance(v, float) else v}")
    for k, d in R.items():
        print(f"== {k}")
        show(d)
    print("== caso (propio)")
    show(K)


if __name__ == "__main__":
    R = resultados()
    m, C = cargar_caso()
    K = caso(m, C)
    imprimir(R, K)
    u10, u8 = C["u10"], C["u8"]
    ok = (abs(K["ro_v"] - u10["resultado_variable"]) < 1e-6 and abs(K["ro_a"] - u10["resultado_absorcion"]) < 1e-6
          and abs(K["dif"] - u10["diferencia"]) < 1e-6 and K["inv"] == u10["inventario_final"]
          and abs(K["fijo_inv_tot"] - u10["cif_en_inventario"]) < 1e-6
          and abs(K["cva"] - u8["costo_ventas"]) < 1e-6 and abs(K["uai_a"] - u8["utilidad_antes_impuesto"]) < 1e-6
          and abs(K["cif"] - C["u9"]["cif_fijos_mes"]) < 1e-6
          and all(abs(a - b) < 1e-9 for a, b in zip(K["absu"], C["u11"]["costo_absorcion_por_prenda"])))
    print("CADENA:", "OK, la pieza u10 calza con caso_estampa_sur_contable.py (y con u8, u9 y u11)" if ok else "NO calza")
    if not ok:
        sys.exit(1)
    if not planilla(R, K):
        sys.exit(1)
    if "--revisar" in sys.argv:
        md = (AQUI.parent / "unidades" / "10-absorcion-y-variable.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %")
        faltan = [x for x in esperados(R, K) if x not in md]
        print("REVISIÓN:", "OK, todas las cifras están en el texto" if not faltan else f"faltan {len(faltan)}: {faltan}")
        sys.exit(1 if faltan else 0)
