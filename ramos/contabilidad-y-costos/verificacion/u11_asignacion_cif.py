"""Unidad 11 · Asignación de costos indirectos.

Recalcula todas las cifras de la unidad desde los datos, comprueba que cada asiento cuadra, genera la
planilla `planillas/u11_asignacion_cif.xlsx` (con fórmulas), evalúa esas fórmulas con un evaluador
propio y las compara con el cálculo en Python. Con --revisar busca cada resultado en el .md.

    py -3 -P u11_asignacion_cif.py            # imprime los resultados y genera la planilla
    py -3 -P u11_asignacion_cif.py --revisar  # además los busca en el .md

Todos los datos son originales de esta base. Ninguna empresa nombrada existe.
"""
from pathlib import Path
import importlib.util, re, sys

sys.stdout.reconfigure(encoding="utf-8")
AQUI = Path(__file__).resolve().parent


# ---------------------------------------------------------------- herramientas

def fmt(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, dec=1):
    return fmt(100 * x, dec) + " %"


def cuadra(asiento):
    d = sum(m for c, m, lado in asiento["lineas"] if lado == "D")
    h = sum(m for c, m, lado in asiento["lineas"] if lado == "H")
    return abs(d - h) < 1e-6, d


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · la mueblería de Ignacia (presupuesto de mayo y cierre)
EJ = {
    "productos": ["mesas de centro", "sillas", "repisas"],
    "cif_items": {"depreciación y mantención del router CNC": 3_094_000, "energía": 947_300,
                  "arriendo del galpón": 1_480_000, "supervisor": 708_000},
    "unidades": [125, 460, 310],
    "hm": [2.4, 0.5, 0.9],          # horas máquina por unidad
    "hh": [3.2, 1.4, 1.1],          # horas hombre por unidad
    "directo": [58_400, 21_700, 17_900],   # MPD + MOD por unidad
    "recargo": 0.40,
    "reales": [116, 505, 296],      # unidades producidas en mayo
    "cif_real": 6_318_400,
}

# Sección 5 · departamentos de servicio (método directo)
DEP = {
    "servicio": {"Mantención": {"cif": 1_170_000, "base": {"Corte": 312, "Armado": 138, "Bodega": 46}},
                 "Bodega": {"cif": 640_000, "base": {"Corte": 115, "Armado": 205, "Mantención": 35}}},
    "productivo": {"Corte": {"cif": 2_408_800, "base": 1_150, "unidad": "HM"},
                   "Armado": {"cif": 1_761_200, "base": 1_840, "unidad": "HH"}},
}

# Sección 6 · ABC en una imprenta digital
ABC = {
    "productos": ["tarjetas", "volantes"],
    "volumen": [2_400, 1_350],                       # cajas de 100 tarjetas · millares de volantes
    "horas_impresion": [165, 435],
    "cif": 4_560_000,
    "actividades": {"preparar trabajos": (1_776_000, [222, 18]),
                    "revisar archivos": (912_000, [170, 22]),
                    "imprimir": (1_872_000, [165, 435])},
}

# 11.1 · taller mecánico
E1 = {"cif_anual": 18_900_000, "hh_anual": 4_200, "ordenes": {"A": 6.5, "B": 2.0, "C": 11.0}}
# 11.2 · banquetería
E2 = {"mpd": 410_000, "hh": 24, "costo_hh": 5_200, "tasa": 3_600, "recargo": 0.35, "invitados": 80}
# 11.3 · fábrica de cecinas
E3 = {"tasa": 2_150, "hm_real": 3_380, "cif_real": 6_980_000}
# 11.4 · envases plásticos
E4 = {"cif": 9_600_000, "unidades": [40_000, 8_000], "hm": [0.01, 0.075]}
# 11.5 · prorrateo de una subaplicación
E5 = {"sub": 348_000, "aplicado": {"Productos en proceso": 640_000, "Productos terminados": 1_160_000,
                                    "Costo de ventas": 4_000_000}}
# 11.6 · lavandería industrial: método directo y escalonado
E6 = {"servicio": {"Mantención": {"cif": 840_000, "base": {"Lavado": 210, "Planchado": 70, "Caldera": 40}},
                   "Caldera": {"cif": 1_260_000, "base": {"Lavado": 5_400, "Planchado": 3_600, "Mantención": 0}}},
      "productivo": {"Lavado": {"cif": 2_100_000, "base": 42_000}, "Planchado": {"cif": 1_500_000, "base": 1_230}},
      "hotel": {"Lavado": 1_500, "Planchado": 60}}
# 11.7 · tasa única contra tasas por departamento (mochilas)
E7 = {"q": [3_000, 1_500], "hm_corte": [0.10, 0.40], "hh_costura": [0.8, 1.0], "hh_corte": [0.05, 0.10],
      "cif_corte": 2_700_000, "cif_costura": 2_340_000, "directo": [9_800, 21_500], "recargo": 0.30,
      "competencia": 13_900}
# 11.8 · pedido especial
E8 = {"cif_variable_u": 180, "precio_pedido": 10_300, "cantidad": 600, "precio_normal": 14_000}
# 11.9 · ABC en una cervecería
E9 = {"productos": ["lata", "edición especial"], "volumen": [34_500, 5_750], "hh": [870, 290],
      "cif": 7_192_000, "directo": [835, 1_665], "precio": [1_420, 2_640],
      "actividades": {"cocción": (2_856_000, [31, 17]), "control de calidad": (1_428_000, [42, 77]),
                      "envasado": (2_184_000, [210, 154]), "despacho": (724_000, [88, 93])}}
# 11.10 · base presupuestada contra capacidad práctica
E10 = {"cif_fijo": 5_880_000, "capacidad": 2_100, "hm_A": 1_680, "hm_B": 1_260, "hm_u": 0.6,
       "cv": 4_150, "recargo": 0.25}


# ---------------------------------------------------------------- resultados

def directo(dep):
    """Método directo: cada servicio se reparte solo entre los productivos, según su base."""
    prod = list(dep["productivo"])
    total = {p: dep["productivo"][p]["cif"] for p in prod}
    reparto = {}
    for s, d in dep["servicio"].items():
        base = sum(d["base"].get(p, 0) for p in prod)
        reparto[s] = {p: d["cif"] * d["base"].get(p, 0) / base for p in prod}
        for p in prod:
            total[p] += reparto[s][p]
    return reparto, total


def abc(datos):
    n = len(datos["productos"])
    tasas, asignado = {}, [0.0] * n
    for a, (costo, uso) in datos["actividades"].items():
        tasas[a] = costo / sum(uso)
        for k in range(n):
            asignado[k] += tasas[a] * uso[k]
    return tasas, asignado


def resultados():
    R = {}
    e = EJ
    cif = sum(e["cif_items"].values())
    hm_tot = sum(q * h for q, h in zip(e["unidades"], e["hm"]))
    hh_tot = sum(q * h for q, h in zip(e["unidades"], e["hh"]))
    u_tot = sum(e["unidades"])
    t_hm, t_hh, t_u = cif / hm_tot, cif / hh_tot, cif / u_tot
    cif_hm = [t_hm * h for h in e["hm"]]
    cif_hh = [t_hh * h for h in e["hh"]]
    cif_un = [t_u] * 3
    comp = lambda cifs: [d + c for d, c in zip(e["directo"], cifs)]
    hm_real = sum(q * h for q, h in zip(e["reales"], e["hm"]))
    aplicado = hm_real * t_hm
    sub = e["cif_real"] - aplicado
    maquina = e["cif_items"]["depreciación y mantención del router CNC"] + e["cif_items"]["energía"]
    R["ej"] = dict(cif=cif, hm_tot=hm_tot, hh_tot=hh_tot, u_tot=u_tot, t_hm=t_hm, t_hh=t_hh, t_u=t_u,
                   cif_hm=cif_hm, cif_hh=cif_hh, cif_un=cif_un,
                   comp_hm=comp(cif_hm), comp_hh=comp(cif_hh), comp_un=comp(cif_un),
                   precio_hm=[c * (1 + e["recargo"]) for c in comp(cif_hm)],
                   precio_un=[c * (1 + e["recargo"]) for c in comp(cif_un)],
                   hm_real_por=[q * h for q, h in zip(e["reales"], e["hm"])], hm_real=hm_real,
                   aplicado=aplicado, sub=sub, sub_pct=sub / aplicado,
                   maquina=maquina, maquina_pct=maquina / cif,
                   dif_mesa=(cif_hm[0] - t_u) * e["unidades"][0], dif_silla=(t_u - cif_hm[1]) * e["unidades"][1],
                   dif_precio_mesa=(comp(cif_hm)[0] - comp(cif_un)[0]) * (1 + e["recargo"]),
                   dif_precio_silla=(comp(cif_un)[1] - comp(cif_hm)[1]) * (1 + e["recargo"]))
    R["ej"]["asientos"] = [
        {"n": 1, "fecha": "31-05-2026", "glosa": "CIF aplicado a la producción de mayo (797,3 HM × $7.700)",
         "lineas": [("Productos en proceso", aplicado, "D"), ("CIF", aplicado, "H")]},
        {"n": 2, "fecha": "31-05-2026", "glosa": "Cierre del CIF subaplicado contra el costo de ventas",
         "lineas": [("Costo de ventas", sub, "D"), ("CIF", sub, "H")]},
    ]

    reparto, total = directo(DEP)
    tasas_dep = {p: total[p] / DEP["productivo"][p]["base"] for p in total}
    R["dep"] = dict(reparto=reparto, total=total, tasas=tasas_dep,
                    suma=sum(total.values()),
                    entrada=sum(d["cif"] for d in DEP["servicio"].values()) + sum(d["cif"] for d in DEP["productivo"].values()))

    tasas_abc, asig = abc(ABC)
    t_trad = ABC["cif"] / sum(ABC["horas_impresion"])
    trad = [t_trad * h for h in ABC["horas_impresion"]]
    R["abc"] = dict(tasas=tasas_abc, asignado=asig, t_trad=t_trad, trad=trad,
                    por_unidad_abc=[a / v for a, v in zip(asig, ABC["volumen"])],
                    por_unidad_trad=[a / v for a, v in zip(trad, ABC["volumen"])],
                    por_act={a: [tasas_abc[a] * u for u in uso] for a, (_, uso) in ABC["actividades"].items()},
                    dif=asig[0] - trad[0])

    e = E1
    t = e["cif_anual"] / e["hh_anual"]
    R["e1"] = dict(tasa=t, ordenes={k: t * h for k, h in e["ordenes"].items()},
                   total=t * sum(e["ordenes"].values()))

    e = E2
    mod = e["hh"] * e["costo_hh"]
    cif2 = e["hh"] * e["tasa"]
    costo = e["mpd"] + mod + cif2
    R["e2"] = dict(mod=mod, cif=cif2, costo=costo, precio=costo * (1 + e["recargo"]),
                   por_invitado=costo / e["invitados"], precio_invitado=costo * (1 + e["recargo"]) / e["invitados"])

    e = E3
    apl = e["tasa"] * e["hm_real"]
    R["e3"] = dict(aplicado=apl, sobre=apl - e["cif_real"], sobre_pct=(apl - e["cif_real"]) / apl)
    R["e3"]["asiento"] = {"n": 1, "fecha": "30-06-2026", "glosa": "Cierre del CIF sobreaplicado contra el costo de ventas",
                          "lineas": [("CIF", apl - e["cif_real"], "D"), ("Costo de ventas", apl - e["cif_real"], "H")]}

    e = E4
    hm = [q * h for q, h in zip(e["unidades"], e["hm"])]
    t_hm4 = e["cif"] / sum(hm)
    t_u4 = e["cif"] / sum(e["unidades"])
    R["e4"] = dict(hm=hm, hm_tot=sum(hm), t_hm=t_hm4, t_u=t_u4, por_u_hm=[t_hm4 * h for h in e["hm"]],
                   bidones_sub=(t_hm4 * e["hm"][1] - t_u4) * e["unidades"][1])

    e = E5
    tot = sum(e["aplicado"].values())
    prorr = {k: e["sub"] * v / tot for k, v in e["aplicado"].items()}
    R["e5"] = dict(total=tot, factor=e["sub"] / tot, prorr=prorr,
                   inventario=prorr["Productos en proceso"] + prorr["Productos terminados"])
    R["e5"]["asiento"] = {"n": 1, "fecha": "31-12-2026", "glosa": "Prorrateo del CIF subaplicado del año",
                          "lineas": [(k, v, "D") for k, v in prorr.items()] + [("CIF", e["sub"], "H")]}

    e = E6
    rep6, tot6 = directo(e)
    tasas6 = {p: tot6[p] / e["productivo"][p]["base"] for p in tot6}
    hotel = {p: tasas6[p] * e["hotel"][p] for p in tasas6}
    # escalonado: primero Mantención (atiende a la caldera), después Caldera
    m = e["servicio"]["Mantención"]
    bm = sum(m["base"].values())
    esc_m = {k: m["cif"] * v / bm for k, v in m["base"].items()}
    cal_total = e["servicio"]["Caldera"]["cif"] + esc_m["Caldera"]
    c = e["servicio"]["Caldera"]
    bc = c["base"]["Lavado"] + c["base"]["Planchado"]
    esc_c = {p: cal_total * c["base"][p] / bc for p in ("Lavado", "Planchado")}
    esc_tot = {p: e["productivo"][p]["cif"] + esc_m[p] + esc_c[p] for p in ("Lavado", "Planchado")}
    R["e6"] = dict(reparto=rep6, total=tot6, tasas=tasas6, hotel=hotel, hotel_total=sum(hotel.values()),
                   esc_m=esc_m, cal_total=cal_total, esc_c=esc_c, esc_tot=esc_tot)

    e = E7
    hh_tot7 = sum(q * (a + b) for q, a, b in zip(e["q"], e["hh_costura"], e["hh_corte"]))
    hm_tot7 = sum(q * h for q, h in zip(e["q"], e["hm_corte"]))
    cos_tot7 = sum(q * h for q, h in zip(e["q"], e["hh_costura"]))
    t_planta = (e["cif_corte"] + e["cif_costura"]) / hh_tot7
    t_corte, t_cost = e["cif_corte"] / hm_tot7, e["cif_costura"] / cos_tot7
    planta = [t_planta * (a + b) for a, b in zip(e["hh_costura"], e["hh_corte"])]
    depto = [t_corte * m_ + t_cost * c_ for m_, c_ in zip(e["hm_corte"], e["hh_costura"])]
    R["e7"] = dict(hh_tot=hh_tot7, hm_tot=hm_tot7, cos_tot=cos_tot7, t_planta=t_planta, t_corte=t_corte,
                   t_cost=t_cost, planta=planta, depto=depto,
                   precio_planta=[(d + c_) * (1 + e["recargo"]) for d, c_ in zip(e["directo"], planta)],
                   precio_depto=[(d + c_) * (1 + e["recargo"]) for d, c_ in zip(e["directo"], depto)],
                   asignado_planta=sum(q * c_ for q, c_ in zip(e["q"], planta)),
                   asignado_depto=sum(q * c_ for q, c_ in zip(e["q"], depto)))

    e = E8
    completo = E7["directo"][0] + depto[0]
    cv = E7["directo"][0] + e["cif_variable_u"]
    R["e8"] = dict(completo=completo, cv=cv, perdida_aparente=(completo - e["precio_pedido"]) * e["cantidad"],
                   mc_pedido=(e["precio_pedido"] - cv) * e["cantidad"], mc_u=e["precio_pedido"] - cv,
                   mc_normal=e["precio_normal"] - cv, costo_oportunidad=(e["precio_normal"] - cv) * e["cantidad"],
                   resultado_sin_holgura=(e["precio_pedido"] - e["precio_normal"]) * e["cantidad"])

    e = E9
    tasas9, asig9 = abc(e)
    t_trad9 = e["cif"] / sum(e["hh"])
    trad9 = [t_trad9 * h for h in e["hh"]]
    u_trad = [a / v for a, v in zip(trad9, e["volumen"])]
    u_abc = [a / v for a, v in zip(asig9, e["volumen"])]
    comp_trad = [d + c_ for d, c_ in zip(e["directo"], u_trad)]
    comp_abc = [d + c_ for d, c_ in zip(e["directo"], u_abc)]
    rec_trad = e["precio"][1] / comp_trad[1] - 1
    R["e9"] = dict(tasas=tasas9, asignado=asig9, t_trad=t_trad9, trad=trad9, u_trad=u_trad, u_abc=u_abc,
                   comp_trad=comp_trad, comp_abc=comp_abc,
                   holgura_trad=[p - c_ for p, c_ in zip(e["precio"], comp_trad)],
                   holgura_abc=[p - c_ for p, c_ in zip(e["precio"], comp_abc)],
                   rec_trad=rec_trad, precio_nuevo=comp_abc[1] * (1 + rec_trad),
                   por_act={a: [tasas9[a] * u for u in uso] for a, (_, uso) in e["actividades"].items()})

    e = E10
    tA, tB = e["cif_fijo"] / e["hm_A"], e["cif_fijo"] / e["hm_B"]
    tC = e["cif_fijo"] / e["capacidad"]
    R["e10"] = dict(tA=tA, tB=tB, tC=tC, uA=tA * e["hm_u"], uB=tB * e["hm_u"], uC=tC * e["hm_u"],
                    pA=(e["cv"] + tA * e["hm_u"]) * (1 + e["recargo"]), pB=(e["cv"] + tB * e["hm_u"]) * (1 + e["recargo"]),
                    pC=(e["cv"] + tC * e["hm_u"]) * (1 + e["recargo"]),
                    ociosa_A=(e["capacidad"] - e["hm_A"]) * tC, ociosa_B=(e["capacidad"] - e["hm_B"]) * tC,
                    subida=(tB * e["hm_u"] - tA * e["hm_u"]) / (e["cv"] + tA * e["hm_u"]))
    return R


# ---------------------------------------------------------------- caso integrador

def cargar_caso():
    spec = importlib.util.spec_from_file_location("caso", AQUI / "caso_estampa_sur_contable.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m, m.calcular()


def caso(m, C):
    cif = (m.SUELDOS + sum((x["costo"] - x["residual"]) / x["vida_meses"] for x in m.MAQUINAS.values())
           + m.SEGURO_ANUAL / 12 + 580_000)
    seg = [q * s for q, s in zip(m.PRODUCIDAS, m.SEG_EST)]
    tasa = cif / sum(seg)
    cv = [p - mg for p, mg in zip(m.PRECIO, m.MARGEN)]
    cif_u = [tasa * s for s in m.SEG_EST]
    completo = [c + x for c, x in zip(cv, cif_u)]
    t_un = cif / sum(m.PRODUCIDAS)
    completo_un = [c + t_un for c in cv]
    return dict(cif=cif, seg=seg, seg_tot=sum(seg), tasa=tasa, cv=cv, cif_u=cif_u, completo=completo,
                t_un=t_un, completo_un=completo_un,
                holgura=[p - c for p, c in zip(m.PRECIO, completo)],
                holgura_pct=[(p - c) / p for p, c in zip(m.PRECIO, completo)],
                holgura_un=[p - c for p, c in zip(m.PRECIO, completo_un)],
                dif_u=[t_un - x for x in cif_u],
                dif_total=[(t_un - x) * q for x, q in zip(cif_u, m.PRODUCIDAS)],
                sueldos=m.SUELDOS, dep=C["u4"]["total_mes"], seguro=m.SEGURO_ANUAL / 12, arriendo=580_000,
                producidas=sum(m.PRODUCIDAS))


# ---------------------------------------------------------------- planilla con fórmulas

REF = re.compile(r"\b([A-Z]{1,2})(\d+)\b")


def evaluar(ws):
    """Evalúa las fórmulas de una hoja (aritmética, SUM y SUMPRODUCT de rangos de la misma hoja)."""
    cache = {}

    def celdas(a, b):
        return [c.value for fila in ws[f"{a}:{b}"] for c in fila]

    def val(coord):
        if coord not in cache:
            v = ws[coord].value
            if isinstance(v, str) and v.startswith("="):
                v = calc(v[1:])
            cache[coord] = v if v is not None else 0
        return cache[coord]

    def rango(a, b):
        return [val(c.coordinate) for fila in ws[f"{a}:{b}"] for c in fila]

    def calc(f):
        f = re.sub(r"SUMPRODUCT\((\w+):(\w+),(\w+):(\w+)\)",
                   lambda m: repr(sum(x * y for x, y in zip(rango(m[1], m[2]), rango(m[3], m[4])))), f)
        f = re.sub(r"SUM\((\w+):(\w+)\)", lambda m: repr(sum(rango(m[1], m[2]))), f)
        f = REF.sub(lambda m: "(" + repr(float(val(m[1] + m[2]))) + ")", f)
        return eval(f, {"__builtins__": {}}, {})

    return val


def planilla(R, K, m):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    neg = Font(bold=True)

    # Hoja 1 · la mueblería del ejemplo
    ws = wb.active
    ws.title = "Mueblería"
    e = EJ
    ws["A1"] = "Asignación de CIF · la mueblería de Ignacia (presupuesto de mayo)"; ws["A1"].font = neg
    ws["A3"] = "CIF presupuestado del mes"; ws["B3"] = sum(e["cif_items"].values())
    ws["A4"] = "Recargo sobre el costo"; ws["B4"] = e["recargo"]
    ws.append([])
    ws["A6"], ws["B6"], ws["C6"], ws["D6"], ws["E6"] = "Dato", *e["productos"], "Total"
    filas = [("Unidades presupuestadas", e["unidades"]), ("Horas máquina por unidad", e["hm"]),
             ("Horas hombre por unidad", e["hh"]), ("Costo directo por unidad (MPD + MOD)", e["directo"])]
    for r, (nombre, vals) in enumerate(filas, start=7):
        ws.cell(r, 1, nombre)
        for k, v in enumerate(vals):
            ws.cell(r, 2 + k, v)
    ws["E7"] = "=SUM(B7:D7)"
    ws["A11"] = "Horas máquina del mes"
    for k, col in enumerate("BCD"):
        ws[f"{col}11"] = f"={col}7*{col}8"
    ws["E11"] = "=SUM(B11:D11)"
    ws["A12"] = "Horas hombre del mes"
    for col in "BCD":
        ws[f"{col}12"] = f"={col}7*{col}9"
    ws["E12"] = "=SUM(B12:D12)"
    ws["A14"] = "Tasa por hora máquina"; ws["B14"] = "=B3/E11"
    ws["A15"] = "Tasa por hora hombre"; ws["B15"] = "=B3/E12"
    ws["A16"] = "Tasa por unidad"; ws["B16"] = "=B3/E7"
    ws["A18"], ws["B18"], ws["C18"], ws["D18"] = "Por unidad", *e["productos"]
    ws["A19"] = "CIF por hora máquina"
    ws["A20"] = "CIF por hora hombre"
    ws["A21"] = "CIF por unidad"
    ws["A22"] = "Costo completo (base horas máquina)"
    ws["A23"] = "Costo completo (base unidades)"
    ws["A24"] = "Precio con recargo (base horas máquina)"
    ws["A25"] = "Precio con recargo (base unidades)"
    for col in "BCD":
        ws[f"{col}19"] = f"=B14*{col}8"
        ws[f"{col}20"] = f"=B15*{col}9"
        ws[f"{col}21"] = "=B16"
        ws[f"{col}22"] = f"={col}10+{col}19"
        ws[f"{col}23"] = f"={col}10+{col}21"
        ws[f"{col}24"] = f"={col}22*(1+B4)"
        ws[f"{col}25"] = f"={col}23*(1+B4)"
    ws["A27"] = "Cierre de mayo"; ws["A27"].font = neg
    ws["A28"] = "Unidades producidas"
    for k, col in enumerate("BCD"):
        ws[f"{col}28"] = e["reales"][k]
        ws[f"{col}29"] = f"={col}28*{col}8"
    ws["A29"] = "Horas máquina reales"; ws["E29"] = "=SUM(B29:D29)"
    ws["A30"] = "CIF aplicado"; ws["B30"] = "=B14*E29"
    ws["A31"] = "CIF real"; ws["B31"] = e["cif_real"]
    ws["A32"] = "CIF subaplicado (real − aplicado)"; ws["B32"] = "=B31-B30"
    ws.column_dimensions["A"].width = 42
    esperado1 = {"B14": R["ej"]["t_hm"], "B15": R["ej"]["t_hh"], "B16": R["ej"]["t_u"],
                 "B22": R["ej"]["comp_hm"][0], "C22": R["ej"]["comp_hm"][1], "D22": R["ej"]["comp_hm"][2],
                 "B23": R["ej"]["comp_un"][0], "C23": R["ej"]["comp_un"][1], "D23": R["ej"]["comp_un"][2],
                 "B24": R["ej"]["precio_hm"][0], "C25": R["ej"]["precio_un"][1],
                 "B30": R["ej"]["aplicado"], "B32": R["ej"]["sub"]}

    # Hoja 2 · el caso
    ws2 = wb.create_sheet("Estampa Sur")
    ws2["A1"] = "Estampa Sur · CIF fijo de marzo por segundo de estampado"; ws2["A1"].font = neg
    ws2["A3"] = "Sueldos"; ws2["B3"] = K["sueldos"]
    ws2["A4"] = "Depreciación de las dos máquinas"; ws2["B4"] = K["dep"]
    ws2["A5"] = "Seguro devengado"; ws2["B5"] = K["seguro"]
    ws2["A6"] = "Arriendo devengado"; ws2["B6"] = K["arriendo"]
    ws2["A7"] = "CIF fijo del mes"; ws2["B7"] = "=SUM(B3:B6)"
    cols = "BCDEF"
    ws2["A9"] = "Dato"
    for k, col in enumerate(cols):
        ws2[f"{col}9"] = m.PRODUCTOS[k]
    ws2["G9"] = "Total"
    datos2 = [("Prendas producidas", m.PRODUCIDAS), ("Segundos de estampado por prenda", m.SEG_EST),
              ("Precio neto", m.PRECIO), ("Costo variable estándar", K["cv"])]
    for r, (nombre, vals) in enumerate(datos2, start=10):
        ws2.cell(r, 1, nombre)
        for k, v in enumerate(vals):
            ws2.cell(r, 2 + k, v)
    ws2["G10"] = "=SUM(B10:F10)"
    ws2["A14"] = "Segundos del mes"
    for col in cols:
        ws2[f"{col}14"] = f"={col}10*{col}11"
    ws2["G14"] = "=SUM(B14:F14)"
    ws2["A16"] = "Tasa por segundo"; ws2["B16"] = "=B7/G14"
    ws2["A17"] = "Tasa por prenda (base unidades)"; ws2["B17"] = "=B7/G10"
    ws2["A19"] = "CIF fijo por prenda (segundos)"
    ws2["A20"] = "Costo completo por prenda"
    ws2["A21"] = "Precio − costo completo"
    ws2["A22"] = "CIF fijo por prenda (unidades)"
    ws2["A23"] = "Diferencia unidades − segundos"
    ws2["A24"] = "CIF mal asignado en el mes"
    for col in cols:
        ws2[f"{col}19"] = f"=B16*{col}11"
        ws2[f"{col}20"] = f"={col}13+{col}19"
        ws2[f"{col}21"] = f"={col}12-{col}20"
        ws2[f"{col}22"] = "=B17"
        ws2[f"{col}23"] = f"={col}22-{col}19"
        ws2[f"{col}24"] = f"={col}23*{col}10"
    ws2["G24"] = "=SUM(B24:F24)"
    ws2["A25"] = "Comprobación: CIF asignado por segundos"; ws2["B25"] = "=SUMPRODUCT(B19:F19,B10:F10)"
    ws2.column_dimensions["A"].width = 42
    esperado2 = {"B7": K["cif"], "G14": K["seg_tot"], "B16": K["tasa"], "B17": K["t_un"], "G24": 0.0,
                 "B25": K["cif"]}
    for k, col in enumerate(cols):
        esperado2[f"{col}19"] = K["cif_u"][k]
        esperado2[f"{col}20"] = K["completo"][k]
        esperado2[f"{col}24"] = K["dif_total"][k]

    destino = AQUI.parent / "planillas" / "u11_asignacion_cif.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)

    from openpyxl import load_workbook
    wb2 = load_workbook(destino)
    ok = True
    for hoja, esp in (("Mueblería", esperado1), ("Estampa Sur", esperado2)):
        val = evaluar(wb2[hoja])
        for celda, v in esp.items():
            if not isinstance(wb2[hoja][celda].value, str) or not wb2[hoja][celda].value.startswith("="):
                print(f"   MAL {hoja}!{celda} no es fórmula"); ok = False; continue
            if abs(val(celda) - v) > 1e-6:
                print(f"   MAL {hoja}!{celda}: planilla {val(celda)} · Python {v}"); ok = False
    print(f"PLANILLA: {'OK' if ok else 'NO calza'}, {destino.name} (fórmulas evaluadas contra Python)")
    return ok


# ---------------------------------------------------------------- texto

def esperados(R, K):
    p = lambda x: fmt(x)
    e = []
    j = R["ej"]
    e += [p(j["cif"]), p(j["hm_tot"]), p(j["hh_tot"]), p(j["u_tot"]), p(j["t_hm"]), fmt(j["t_hh"], 2), fmt(j["t_u"], 2)]
    e += [p(x) for x in j["cif_hm"] + j["cif_hh"] + j["comp_hm"] + j["comp_hh"] + j["comp_un"] + j["precio_hm"] + j["precio_un"]]
    e += [fmt(x, 1) for x in j["hm_real_por"]] + [fmt(j["hm_real"], 1), p(j["aplicado"]), p(j["sub"]), pct(j["sub_pct"]),
                                             p(j["maquina"]), pct(j["maquina_pct"]), p(j["dif_mesa"]), p(j["dif_silla"]),
                                             p(j["dif_precio_mesa"]), p(j["dif_precio_silla"])]
    d = R["dep"]
    for s in d["reparto"].values():
        e += [p(x) for x in s.values()]
    e += [p(x) for x in d["total"].values()] + [p(x) for x in d["tasas"].values()] + [p(d["suma"])]
    a = R["abc"]
    e += [p(x) for x in a["tasas"].values()] + [p(x) for x in a["asignado"]] + [p(a["t_trad"])] + [p(x) for x in a["trad"]]
    e += [fmt(x, 1) for x in a["por_unidad_abc"] + a["por_unidad_trad"]] + [p(a["dif"])]
    for v in a["por_act"].values():
        e += [p(x) for x in v]
    j = R["e1"]; e += [p(j["tasa"]), p(j["total"])] + [p(x) for x in j["ordenes"].values()]
    j = R["e2"]; e += [p(j["mod"]), p(j["cif"]), p(j["costo"]), p(j["precio"]), p(j["por_invitado"]), p(j["precio_invitado"])]
    j = R["e3"]; e += [p(j["aplicado"]), p(j["sobre"]), pct(j["sobre_pct"])]
    j = R["e4"]; e += [p(x) for x in j["hm"]] + [p(j["hm_tot"]), p(j["t_hm"]), p(j["t_u"]), p(j["bidones_sub"])] + [p(x) for x in j["por_u_hm"]]
    j = R["e5"]; e += [p(j["total"]), pct(j["factor"]), p(j["inventario"])] + [p(x) for x in j["prorr"].values()]
    j = R["e6"]
    for s in j["reparto"].values():
        e += [p(x) for x in s.values()]
    e += [p(x) for x in j["total"].values()] + [p(x) for x in j["tasas"].values()] + [p(x) for x in j["hotel"].values()]
    e += [p(j["hotel_total"]), p(j["cal_total"])] + [p(x) for x in j["esc_m"].values()] + [p(x) for x in j["esc_c"].values()]
    e += [p(x) for x in j["esc_tot"].values()]
    j = R["e7"]
    e += [p(j["hh_tot"]), p(j["hm_tot"]), p(j["cos_tot"]), p(j["t_planta"]), p(j["t_corte"]), p(j["t_cost"]),
          p(j["asignado_planta"])]
    e += [p(x) for x in j["planta"] + j["depto"] + j["precio_planta"] + j["precio_depto"]]
    j = R["e8"]; e += [p(j["completo"]), p(j["cv"]), p(j["perdida_aparente"]), p(j["mc_pedido"]), p(j["mc_u"]),
                       p(j["mc_normal"]), p(j["costo_oportunidad"]), p(abs(j["resultado_sin_holgura"]))]
    j = R["e9"]
    e += [p(x) for x in j["tasas"].values()] + [p(x) for x in j["asignado"]] + [p(j["t_trad"])] + [p(x) for x in j["trad"]]
    e += [fmt(x, 2) for x in j["u_abc"] + j["u_trad"] + j["comp_abc"] + j["comp_trad"] + j["holgura_trad"] + j["holgura_abc"]]
    e += [pct(j["rec_trad"]), fmt(j["precio_nuevo"], 2)]
    for v in j["por_act"].values():
        e += [p(x) for x in v]
    j = R["e10"]
    e += [p(j[k]) for k in ("tA", "tB", "tC", "uA", "uB", "uC", "ociosa_A", "ociosa_B")]
    e += [fmt(j["pA"], 2), p(j["pB"]), p(j["pC"]), pct(j["subida"])]
    # caso
    e += [fmt(K["cif"], 2), p(K["seg_tot"]), fmt(K["tasa"], 4), fmt(K["t_un"], 2), p(K["producidas"])]
    e += [p(x) for x in K["seg"]]
    e += [fmt(x, 2) for x in K["cif_u"] + K["completo"] + K["holgura"] + K["completo_un"] + K["holgura_un"]]
    e += [fmt(abs(x), 2) for x in K["dif_u"]] + [pct(x) for x in K["holgura_pct"]] + [p(abs(x)) for x in K["dif_total"]]
    e += [fmt(K["dep"], 2), p(K["seguro"]), p(K["sueldos"])]
    return e


def imprimir(R, K):
    def show(v):
        if isinstance(v, float):
            return f"{v:,.4f}"
        if isinstance(v, list):
            return "[" + ", ".join(show(x) for x in v) + "]"
        if isinstance(v, dict):
            return "{" + ", ".join(f"{k}: {show(x)}" for k, x in v.items()) + "}"
        return str(v)
    for k, d in list(R.items()) + [("caso", K)]:
        print(f"== {k}")
        for kk, v in d.items():
            if kk in ("asientos", "asiento"):
                continue
            print(f"   {kk}: {show(v)}")


if __name__ == "__main__":
    R = resultados()
    m, C = cargar_caso()
    K = caso(m, C)
    imprimir(R, K)

    asientos = R["ej"]["asientos"] + [R["e3"]["asiento"], R["e5"]["asiento"]]
    chequeos = {f"asiento {a['glosa']}: cuadra": cuadra(a)[0] for a in asientos}
    chequeos.update({
        "ejemplo: el CIF asignado por horas máquina suma el presupuesto":
            abs(sum(q * c for q, c in zip(EJ["unidades"], R["ej"]["cif_hm"])) - R["ej"]["cif"]) < 1e-6,
        "ejemplo: lo mismo por horas hombre":
            abs(sum(q * c for q, c in zip(EJ["unidades"], R["ej"]["cif_hh"])) - R["ej"]["cif"]) < 1e-6,
        "departamentos: el método directo no pierde ni crea costo": abs(R["dep"]["suma"] - R["dep"]["entrada"]) < 1e-6,
        "ABC: suma el CIF total": abs(sum(R["abc"]["asignado"]) - ABC["cif"]) < 1e-6,
        "11.5: el prorrateo suma la subaplicación": abs(sum(R["e5"]["prorr"].values()) - E5["sub"]) < 1e-6,
        "11.6: directo y escalonado reparten lo mismo en total":
            abs(sum(R["e6"]["total"].values()) - sum(R["e6"]["esc_tot"].values())) < 1e-6,
        "11.7: tasa única y por departamento asignan el mismo total":
            abs(R["e7"]["asignado_planta"] - R["e7"]["asignado_depto"]) < 1e-6
            and abs(R["e7"]["asignado_planta"] - E7["cif_corte"] - E7["cif_costura"]) < 1e-6,
        "11.9: ABC suma el CIF total": abs(sum(R["e9"]["asignado"]) - E9["cif"]) < 1e-6,
        "caso: tasa propia = script": abs(K["tasa"] - C["u11"]["tasa_por_segundo"]) < 1e-9,
        "caso: segundos propios = script": K["seg_tot"] == C["u11"]["segundos_estampado_mes"],
        "caso: CIF por prenda propio = script": all(abs(a - b) < 1e-9 for a, b in zip(K["cif_u"], C["u11"]["cif_por_prenda"])),
        "caso: costo completo propio = script":
            all(abs(a - b) < 1e-9 for a, b in zip(K["completo"], C["u11"]["costo_absorcion_por_prenda"])),
        "caso: CIF fijo propio = u9 del script": abs(K["cif"] - C["u9"]["cif_fijos_mes"]) < 1e-6,
        "caso: lo que una base quita, la otra lo pone (suma cero)": abs(sum(K["dif_total"])) < 1e-6,
    })
    print("\nCHEQUEOS:")
    for k, v in chequeos.items():
        print(f"   {'OK ' if v else 'MAL'} {k}")
    ok_planilla = planilla(R, K, m)
    if not all(chequeos.values()) or not ok_planilla:
        sys.exit(1)
    print("CADENA: OK, la pieza del caso calza con caso_estampa_sur_contable.py")
    if "--revisar" in sys.argv:
        md = (AQUI.parent / "unidades" / "11-asignacion-de-costos-indirectos.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %")
        faltan = [x for x in esperados(R, K) if x not in md]
        print("\nREVISIÓN:", "OK, todas las cifras están en el texto" if not faltan else f"faltan {len(faltan)}: {faltan}")
        sys.exit(1 if faltan else 0)
