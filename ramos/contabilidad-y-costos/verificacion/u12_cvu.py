"""Unidad 12 · Costo-volumen-utilidad y punto de equilibrio.

Recalcula todas las cifras de la unidad desde los datos, genera la planilla del punto de equilibrio
multiproducto `planillas/u12_punto_de_equilibrio.xlsx` (con fórmulas), evalúa esas fórmulas con un
evaluador propio y las compara con Python. Las mezclas con un recurso escaso se comprueban contra
fuerza bruta. Con --revisar busca cada resultado en el .md.

    py -3 -P u12_cvu.py            # imprime los resultados y genera la planilla
    py -3 -P u12_cvu.py --revisar  # además los busca en el .md

Todos los datos son originales de esta base. Ninguna empresa nombrada existe. La tasa de impuesto del
ejercicio 12.2 es un supuesto del ejercicio, no la tasa legal vigente.
"""
from math import floor
from pathlib import Path
import importlib.util, re, sys

sys.stdout.reconfigure(encoding="utf-8")
AQUI = Path(__file__).resolve().parent


def fmt(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, dec=1):
    return fmt(100 * x, dec) + " %"


# ---------------------------------------------------------------- CVU

def multiproducto(p, cv, q, cf):
    tot = sum(q)
    mezcla = [x / tot for x in q]
    mc = [a - b for a, b in zip(p, cv)]
    mcp = sum(m * c for m, c in zip(mezcla, mc))
    pp = sum(m * a for m, a in zip(mezcla, p))
    pe = cf / mcp
    uo = tot * mcp - cf
    return dict(mezcla=mezcla, mc=mc, mcp=mcp, pp=pp, razon=mcp / pp, pe=pe, pe_prod=[pe * m for m in mezcla],
                pe_pesos=pe * pp, pe_ing=[pe * m * a for m, a in zip(mezcla, p)],
                ventas=tot * pp, mc_total=tot * mcp, uo=uo, ms=(tot - pe) / tot, gao=tot * mcp / uo,
                comprobacion=sum(pe * m * (a - b) for m, a, b in zip(mezcla, p, cv)) - cf)


def uno(p, cv, cf, q=None):
    mc = p - cv
    r = dict(mc=mc, razon=mc / p, pe=cf / mc, pe_pesos=cf / mc * p)
    if q is not None:
        r.update(uo=q * mc - cf, mc_total=q * mc, ms=(q - cf / mc) / q, ms_u=q - cf / mc,
                 gao=q * mc / (q * mc - cf), ventas=q * p)
    return r


def mejor_mezcla(mc, uso, dem, cap):
    """Llena el recurso escaso por margen por unidad de recurso (enteros)."""
    orden = sorted(range(len(mc)), key=lambda k: -mc[k] / uso[k])
    q, libre = [0] * len(mc), cap
    for k in orden:
        q[k] = min(dem[k], floor(libre / uso[k] + 1e-9))
        libre -= q[k] * uso[k]
    return q, sum(a * b for a, b in zip(q, mc)), orden


def por_margen_unitario(mc, uso, dem, cap):
    orden = sorted(range(len(mc)), key=lambda k: -mc[k])
    q, libre = [0] * len(mc), cap
    for k in orden:
        q[k] = min(dem[k], floor(libre / uso[k] + 1e-9))
        libre -= q[k] * uso[k]
    return q, sum(a * b for a, b in zip(q, mc))


def fuerza_bruta(mc, uso_ent, dem, cap_ent):
    """Óptimo entero con tres productos; el uso va en unidades enteras del recurso."""
    mejor = (-1, None)
    for a in range(dem[0] + 1):
        ra = cap_ent - a * uso_ent[0]
        if ra < 0:
            break
        for b in range(dem[1] + 1):
            rb = ra - b * uso_ent[1]
            if rb < 0:
                break
            c = min(dem[2], rb // uso_ent[2])
            v = a * mc[0] + b * mc[1] + c * mc[2]
            if v > mejor[0]:
                mejor = (v, (a, b, c))
    return mejor


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · la heladería de Martín (enero y febrero)
EJ = {"productos": ["cono simple", "vaso doble", "pote de 1 litro"], "p": [2_250, 3_450, 8_900],
      "cv": [880, 1_430, 3_640], "q": [3_150, 1_890, 1_260], "cf": 4_986_000,
      "objetivo": 6_150_000, "precio_cono_nuevo": 2_475, "vendedora": 640_000,
      "litros": [0.12, 0.22, 1.0], "alza_febrero": 0.20, "capacidad_l": 2_050}
# 12.1 · lavado de autos
E1 = {"p": 12_400, "cv": 3_350, "cf": 3_529_500, "q": 515}
# 12.2 · academia de reforzamiento (t es un supuesto del ejercicio)
E2 = {"p": 64_500, "cv": 13_800, "cf": 4_563_000, "u": 2_028_000, "u_pct": 0.15, "neta": 1_512_000, "t": 0.25}
# 12.3 · carro de comida
E3 = {"p": 6_450, "cv": 2_580, "cf": 2_709_000, "q": 1_120, "baja": 0.10}
# 12.4 · tienda con estado de resultados por contribución
E4 = {"ventas": 18_450_000, "cv": 11_070_000, "cf": 5_535_000}
# 12.5 y 12.6 · panadería
E5 = {"productos": ["pan (kg)", "empanadas", "tortas"], "p": [1_940, 2_350, 18_700], "cv": [1_060, 1_170, 8_950],
      "q": [6_150, 2_870, 205], "cf": 7_348_000, "q_nueva": [6_250, 2_870, 105]}
# 12.7 · mermeladas: sensibilidad
E7 = {"p": 3_240, "cv": 1_710, "cf": 4_192_000, "q": 4_560,
      "escenarios": {"precio −8 %": ("p", -0.08), "costo variable +10 %": ("cv", 0.10),
                     "costo fijo +15 %": ("cf", 0.15), "volumen −10 %": ("q", -0.10)}}
# 12.8 · automatizar una envasadora de café
E8 = {"p": 5_450, "actual": {"cv": 2_930, "cf": 2_115_000}, "maquina": {"cv": 2_310, "cf": 3_627_000}, "q": 2_240}
# 12.9 · pastelería con el horno como recurso escaso (horas)
E9 = {"productos": ["kuchen", "pie de limón", "tarta de frambuesa"], "mc": [7_250, 5_420, 9_830],
      "horas": [0.5, 0.25, 1.0], "dem": [148, 204, 62], "cap": 164, "cf": 1_915_000}
# 12.10 · gimnasio
E10 = {"p": 31_900, "cv": 4_150, "cf": 8_436_000, "q": 418, "cap": 500, "p_baja": 26_900, "publicidad": 1_480_000}


# ---------------------------------------------------------------- resultados

def resultados():
    R = {}
    e = EJ
    b = multiproducto(e["p"], e["cv"], e["q"], e["cf"])
    q_obj = (e["cf"] + e["objetivo"]) / b["mcp"]
    p2 = [e["precio_cono_nuevo"]] + e["p"][1:]
    s1 = multiproducto(p2, e["cv"], e["q"], e["cf"])
    s2 = multiproducto(e["p"], e["cv"], e["q"], e["cf"] + e["vendedora"])
    dem = [round(x * (1 + e["alza_febrero"])) for x in e["q"]]
    uso_cl = [12, 22, 100]                           # centilitros
    mc = b["mc"]
    q_ok, mc_ok, orden = mejor_mezcla(mc, uso_cl, dem, e["capacidad_l"] * 100)
    q_mal, mc_mal = por_margen_unitario(mc, uso_cl, dem, e["capacidad_l"] * 100)
    bf = fuerza_bruta(mc, uso_cl, dem, e["capacidad_l"] * 100)
    R["ej"] = dict(base=b, q_obj=q_obj, u_obj_ventas=q_obj * b["pp"], s1=s1, s2=s2,
                   uo_s1=sum(e["q"]) * s1["mcp"] - e["cf"], uo_s2=sum(e["q"]) * s2["mcp"] - e["cf"] - e["vendedora"],
                   dem=dem, litros_dem=sum(d * l for d, l in zip(dem, e["litros"])),
                   mc_litro=[m / l for m, l in zip(mc, e["litros"])], q_ok=q_ok, mc_ok=mc_ok, q_mal=q_mal,
                   mc_mal=mc_mal, dif=mc_ok - mc_mal, bf=bf, uo_feb=mc_ok - e["cf"],
                   litros_ok=[q * l for q, l in zip(q_ok, e["litros"])],
                   litros_mal=[q * l for q, l in zip(q_mal, e["litros"])])
    R["ej"]["uo_s2"] = sum(e["q"]) * s2["mcp"] - (e["cf"] + e["vendedora"])

    e = E1
    R["e1"] = uno(e["p"], e["cv"], e["cf"], e["q"])

    e = E2
    r = uno(e["p"], e["cv"], e["cf"])
    q_u = (e["cf"] + e["u"]) / r["mc"]
    q_pct = e["cf"] / (e["p"] * (1 - e["u_pct"]) - e["cv"])
    uai = e["neta"] / (1 - e["t"])
    q_t = (e["cf"] + uai) / r["mc"]
    R["e2"] = dict(**r, q_u=q_u, y_u=q_u * e["p"], q_pct=q_pct, den_pct=e["p"] * (1 - e["u_pct"]) - e["cv"],
                   q_pct_ent=int(-(-q_pct // 1)), u_pct_real=None, uai=uai, q_t=q_t, q_t_ent=int(-(-q_t // 1)))
    qe = R["e2"]["q_pct_ent"]
    R["e2"]["u_pct_real"] = (qe * r["mc"] - e["cf"]) / (qe * e["p"])

    e = E3
    r = uno(e["p"], e["cv"], e["cf"], e["q"])
    q2 = e["q"] * (1 - e["baja"])
    uo2 = q2 * r["mc"] - e["cf"]
    R["e3"] = dict(**r, uo_baja=uo2, caida=(r["uo"] - uo2) / r["uo"], q_baja=q2)

    e = E4
    razon = (e["ventas"] - e["cv"]) / e["ventas"]
    uo = e["ventas"] - e["cv"] - e["cf"]
    R["e4"] = dict(mc=e["ventas"] - e["cv"], razon=razon, pe=e["cf"] / razon, uo=uo,
                   ms=(e["ventas"] - e["cf"] / razon) / e["ventas"], gao=(e["ventas"] - e["cv"]) / uo,
                   cv_pct=e["cv"] / e["ventas"])

    e = E5
    R["e5"] = multiproducto(e["p"], e["cv"], e["q"], e["cf"])
    R["e6"] = multiproducto(e["p"], e["cv"], e["q_nueva"], e["cf"])

    e = E7
    base = uno(e["p"], e["cv"], e["cf"], e["q"])
    esc = {}
    for nombre, (k, d) in e["escenarios"].items():
        x = {"p": e["p"], "cv": e["cv"], "cf": e["cf"], "q": e["q"]}
        x[k] = x[k] * (1 + d)
        r = uno(x["p"], x["cv"], x["cf"], x["q"])
        esc[nombre] = dict(valor=x[k], pe=r["pe"], uo=r["uo"], cambio=(r["uo"] - base["uo"]) / base["uo"])
    uno_pct = {"precio": e["q"] * e["p"] * 0.01, "costo variable": e["q"] * e["cv"] * 0.01,
               "costo fijo": e["cf"] * 0.01, "volumen": e["q"] * base["mc"] * 0.01}
    R["e7"] = dict(base=base, esc=esc, uno_pct=uno_pct, uno_pct_rel={k: v / base["uo"] for k, v in uno_pct.items()})

    e = E8
    a = uno(e["p"], e["actual"]["cv"], e["actual"]["cf"], e["q"])
    m = uno(e["p"], e["maquina"]["cv"], e["maquina"]["cf"], e["q"])
    indif = (e["maquina"]["cf"] - e["actual"]["cf"]) / (e["actual"]["cv"] - e["maquina"]["cv"])
    R["e8"] = dict(a=a, m=m, indif=indif, uo_indif=indif * a["mc"] - e["actual"]["cf"],
                   dif=a["uo"] - m["uo"])

    e = E9
    uso_q = [int(h * 4) for h in e["horas"]]          # cuartos de hora
    q_ok, mc_ok, orden = mejor_mezcla(e["mc"], uso_q, e["dem"], e["cap"] * 4)
    q_mal, mc_mal = por_margen_unitario(e["mc"], uso_q, e["dem"], e["cap"] * 4)
    bf = fuerza_bruta(e["mc"], uso_q, e["dem"], e["cap"] * 4)
    q_mas, mc_mas, _ = mejor_mezcla(e["mc"], uso_q, e["dem"], (e["cap"] + 1) * 4)
    R["e9"] = dict(mc_h=[m_ / h for m_, h in zip(e["mc"], e["horas"])], q_ok=q_ok, mc_ok=mc_ok, q_mal=q_mal,
                   mc_mal=mc_mal, dif=mc_ok - mc_mal, uo_ok=mc_ok - e["cf"], uo_mal=mc_mal - e["cf"], bf=bf,
                   horas_ok=[q * h for q, h in zip(q_ok, e["horas"])], horas_mal=[q * h for q, h in zip(q_mal, e["horas"])],
                   hora_extra=mc_mas - mc_ok, horas_dem=sum(d * h for d, h in zip(e["dem"], e["horas"])))

    e = E10
    base = uno(e["p"], e["cv"], e["cf"], e["q"])
    q_baja = (e["cf"] + base["uo"]) / (e["p_baja"] - e["cv"])
    q_pub = (e["cf"] + e["publicidad"] + base["uo"]) / base["mc"]
    p_min = e["cv"] + (e["cf"] + base["uo"]) / e["cap"]
    uo_cap_baja = e["cap"] * (e["p_baja"] - e["cv"]) - e["cf"]
    R["e10"] = dict(base=base, q_baja=q_baja, q_baja_ent=int(-(-q_baja // 1)), q_pub=q_pub,
                    q_pub_ent=int(-(-q_pub // 1)), alza_baja=q_baja / e["q"] - 1, alza_pub=q_pub / e["q"] - 1,
                    p_min=p_min, uo_cap_baja=uo_cap_baja, pe_baja=e["cf"] / (e["p_baja"] - e["cv"]),
                    pe_pub=(e["cf"] + e["publicidad"]) / base["mc"],
                    ms_pub=(int(-(-q_pub // 1)) - (e["cf"] + e["publicidad"]) / base["mc"]) / int(-(-q_pub // 1)))
    return R


# ---------------------------------------------------------------- caso integrador

def cargar_caso():
    spec = importlib.util.spec_from_file_location("caso", AQUI / "caso_estampa_sur_contable.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m, m.calcular()


def caso(m, C):
    cv = [p - mg for p, mg in zip(m.PRECIO, m.MARGEN)]
    fijos = C["u9"]["cif_fijos_mes"] + m.OTROS_FIJOS
    b = multiproducto(m.PRECIO, cv, m.VENDIDAS, fijos)
    interes = C["u6"]["interes_mes_1"]
    q_tot = sum(m.VENDIDAS)
    q_baja = q_tot * 0.8
    uo_baja = q_baja * b["mcp"] - fijos
    mc_seg = [mg / s for mg, s in zip(m.MARGEN, m.SEG_EST)]
    return dict(b=b, cv=cv, fijos=fijos, cif=C["u9"]["cif_fijos_mes"], admin=m.OTROS_FIJOS, q_tot=q_tot,
                vendidas=m.VENDIDAS, margen=m.MARGEN, interes=interes, pe_interes=(fijos + interes) / b["mcp"],
                uo_baja=uo_baja, caida_baja=(b["uo"] - uo_baja) / b["uo"], mc_seg=mc_seg,
                ventas=C["u2"]["ventas_netas"], uai=b["uo"] - interes)


# ---------------------------------------------------------------- planilla con fórmulas

REF = re.compile(r"\b([A-Z]{1,2})(\d+)\b")


def evaluar(ws):
    """Evalúa las fórmulas de una hoja (aritmética, SUM y SUMPRODUCT de rangos de la misma hoja)."""
    cache = {}

    def val(coord):
        if coord not in cache:
            v = ws[coord].value
            if isinstance(v, str) and v.startswith("="):
                v = calc(v[1:])
            cache[coord] = v if v is not None else 0
        return cache[coord]

    def rango(a, b):
        return [val(c.coordinate) for fila in ws[f"{a}:{b}"] for c in fila]

    def calc(f):
        f = re.sub(r"SUMPRODUCT\((\w+):(\w+),(\w+):(\w+)\)",
                   lambda m: repr(sum(x * y for x, y in zip(rango(m[1], m[2]), rango(m[3], m[4])))), f)
        f = re.sub(r"SUM\((\w+):(\w+)\)", lambda m: repr(sum(rango(m[1], m[2]))), f)
        f = REF.sub(lambda m: "(" + repr(float(val(m[1] + m[2]))) + ")", f)
        return eval(f, {"__builtins__": {}}, {})

    return val


def hoja_equilibrio(ws, titulo, productos, p, cv, q, cf, cf_nombre, objetivo):
    """Arma la hoja del punto de equilibrio multiproducto. Devuelve las celdas clave."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    n = len(productos)
    cols = [get_column_letter(2 + k) for k in range(n)]
    ult, tot = cols[-1], get_column_letter(2 + n)
    ws["A1"] = titulo; ws["A1"].font = Font(bold=True)
    ws["A3"] = "Dato"
    for k, c in enumerate(cols):
        ws[f"{c}3"] = productos[k]
    ws[f"{tot}3"] = "Total"
    ws["A4"] = "Unidades vendidas"; ws["A5"] = "Precio neto"; ws["A6"] = "Costo variable unitario"
    for k, c in enumerate(cols):
        ws[f"{c}4"], ws[f"{c}5"], ws[f"{c}6"] = q[k], p[k], cv[k]
    ws[f"{tot}4"] = f"=SUM(B4:{ult}4)"
    ws["A7"] = cf_nombre; ws["B7"] = cf
    ws["A8"] = "Utilidad objetivo"; ws["B8"] = objetivo
    ws["A10"] = "Mezcla (participación)"
    ws["A11"] = "Margen de contribución unitario"
    ws["A12"] = "Margen ponderado por producto"
    ws["A13"] = "Precio ponderado por producto"
    for c in cols:
        ws[f"{c}10"] = f"={c}4/{tot}4"
        ws[f"{c}11"] = f"={c}5-{c}6"
        ws[f"{c}12"] = f"={c}11*{c}10"
        ws[f"{c}13"] = f"={c}5*{c}10"
    ws[f"{tot}10"] = f"=SUM(B10:{ult}10)"
    ws[f"{tot}12"] = f"=SUM(B12:{ult}12)"
    ws[f"{tot}13"] = f"=SUM(B13:{ult}13)"
    ws["A15"] = "Punto de equilibrio (unidades)"; ws["B15"] = f"=B7/{tot}12"
    ws["A16"] = "Punto de equilibrio (pesos)"; ws["B16"] = f"=B15*{tot}13"
    ws["A17"] = "Razón de margen de contribución"; ws["B17"] = f"={tot}12/{tot}13"
    ws["A18"] = "Equilibrio por producto (unidades)"
    ws["A19"] = "Ingreso en equilibrio por producto"
    for c in cols:
        ws[f"{c}18"] = f"=B15*{c}10"
        ws[f"{c}19"] = f"={c}18*{c}5"
    ws[f"{tot}18"] = f"=SUM(B18:{ult}18)"
    ws[f"{tot}19"] = f"=SUM(B19:{ult}19)"
    ws["A20"] = "Comprobación: utilidad en equilibrio"
    ws["B20"] = f"=SUMPRODUCT(B18:{ult}18,B5:{ult}5)-SUMPRODUCT(B18:{ult}18,B6:{ult}6)-B7"
    ws["A22"] = "Ventas del período (pesos)"; ws["B22"] = f"=SUMPRODUCT(B4:{ult}4,B5:{ult}5)"
    ws["A23"] = "Margen de contribución total"; ws["B23"] = f"=SUMPRODUCT(B4:{ult}4,B11:{ult}11)"
    ws["A24"] = "Utilidad operacional"; ws["B24"] = "=B23-B7"
    ws["A25"] = "Margen de seguridad"; ws["B25"] = f"=({tot}4-B15)/{tot}4"
    ws["A26"] = "Grado de apalancamiento operativo"; ws["B26"] = "=B23/B24"
    ws["A27"] = "Unidades para la utilidad objetivo"; ws["B27"] = f"=(B7+B8)/{tot}12"
    ws.column_dimensions["A"].width = 40
    return {"mcp": f"{tot}12", "pp": f"{tot}13", "pe": "B15", "pe_pesos": "B16", "razon": "B17",
            "comprobacion": "B20", "uo": "B24", "ms": "B25", "gao": "B26", "q_obj": "B27",
            "pe_prod": [f"{c}18" for c in cols], "suma_pe": f"{tot}18"}


def planilla(R, K, m):
    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    ws = wb.active; ws.title = "Heladería"
    e = EJ
    c1 = hoja_equilibrio(ws, "Punto de equilibrio multiproducto · la heladería de Martín (enero)", e["productos"],
                         e["p"], e["cv"], e["q"], e["cf"], "Costos fijos del mes", e["objetivo"])
    ws2 = wb.create_sheet("Estampa Sur")
    c2 = hoja_equilibrio(ws2, "Punto de equilibrio multiproducto · Estampa Sur (marzo)", m.PRODUCTOS, m.PRECIO,
                         K["cv"], m.VENDIDAS, K["fijos"], "Costos fijos operacionales", 0)
    destino = AQUI.parent / "planillas" / "u12_punto_de_equilibrio.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)
    wb = load_workbook(destino)
    ok = True
    for hoja, celdas, b, q_obj in (("Heladería", c1, R["ej"]["base"], R["ej"]["q_obj"]),
                                   ("Estampa Sur", c2, K["b"], K["fijos"] / K["b"]["mcp"])):
        val = evaluar(wb[hoja])
        esp = {"mcp": b["mcp"], "pp": b["pp"], "pe": b["pe"], "pe_pesos": b["pe_pesos"], "razon": b["razon"],
               "comprobacion": 0.0, "uo": b["uo"], "ms": b["ms"], "gao": b["gao"], "q_obj": q_obj,
               "suma_pe": b["pe"]}
        for k, v in esp.items():
            if abs(val(celdas[k]) - v) > 1e-6:
                print(f"   MAL {hoja}!{celdas[k]} ({k}): planilla {val(celdas[k])} · Python {v}"); ok = False
        for celda, v in zip(celdas["pe_prod"], b["pe_prod"]):
            if abs(val(celda) - v) > 1e-6:
                print(f"   MAL {hoja}!{celda}"); ok = False
        formulas = sum(1 for fila in wb[hoja].iter_rows() for c in fila if isinstance(c.value, str) and c.value.startswith("="))
        print(f"   {hoja}: {formulas} fórmulas")
    print(f"PLANILLA: {'OK' if ok else 'NO calza'}, {destino.name} (fórmulas evaluadas contra Python)")
    return ok


# ---------------------------------------------------------------- texto

def esperados(R, K):
    p = lambda x: fmt(x)
    e = []
    j = R["ej"]; b = j["base"]
    e += [pct(x, 0) for x in b["mezcla"]] + [p(x) for x in b["mc"]] + [p(b["mcp"]), p(b["pp"]), pct(b["razon"], 2),
                                                                      fmt(b["pe"], 2), p(b["pe_pesos"])]
    e += [fmt(x, 1) for x in b["pe_prod"]] + [p(x) for x in b["pe_ing"]]
    e += [p(b["ventas"]), p(b["mc_total"]), p(b["uo"]), pct(b["ms"]), fmt(b["gao"], 2)]
    e += [fmt(j["q_obj"], 1), p(j["u_obj_ventas"]), fmt(j["s1"]["mcp"], 1), fmt(j["s1"]["pe"], 1), p(j["uo_s1"]),
          fmt(j["s2"]["pe"], 1), p(j["uo_s2"])]
    e += [p(x) for x in j["dem"]] + [fmt(j["litros_dem"], 2)] + [p(x) for x in j["mc_litro"]]
    e += [p(x) for x in j["q_ok"]] + [p(x) for x in j["q_mal"]] + [p(j["mc_ok"]), p(j["mc_mal"]), p(j["dif"]), p(j["uo_feb"])]
    e += [fmt(x, 2) for x in j["litros_ok"]] + [fmt(x, 2) for x in j["litros_mal"]]
    j = R["e1"]; e += [p(j[k]) for k in ("mc", "pe", "pe_pesos", "uo", "mc_total")] + [pct(j["razon"]), pct(j["ms"]), fmt(j["gao"], 2)]
    j = R["e2"]; e += [p(j["mc"]), p(j["pe"]), p(j["q_u"]), p(j["y_u"]), fmt(j["q_pct"], 1), p(j["den_pct"]),
                       str(j["q_pct_ent"]), pct(j["u_pct_real"]), p(j["uai"]), fmt(j["q_t"], 1), str(j["q_t_ent"])]
    j = R["e3"]; e += [p(j[k]) for k in ("mc", "pe", "mc_total", "uo", "ms_u", "uo_baja", "q_baja")] + [pct(j["ms"]), fmt(j["gao"], 2), pct(j["caida"])]
    j = R["e4"]; e += [p(j["mc"]), pct(j["razon"], 0), p(j["pe"]), p(j["uo"]), pct(j["ms"], 0), fmt(j["gao"], 1)]
    for k in ("e5", "e6"):
        j = R[k]
        e += [pct(x) for x in j["mezcla"]] + [p(x) for x in j["mc"]] + [fmt(j["mcp"], 2), fmt(j["pp"], 2), pct(j["razon"], 2),
                                                                       fmt(j["pe"], 1), p(j["pe_pesos"]), p(j["uo"]), pct(j["ms"])]
        e += [fmt(x, 1) for x in j["pe_prod"]]
    j = R["e7"]; bb = j["base"]
    e += [p(bb["mc"]), p(bb["pe"]), p(bb["uo"])]
    for v in j["esc"].values():
        e += [fmt(v["valor"], 0), fmt(v["pe"], 1), p(v["uo"]), pct(abs(v["cambio"]))]
    e += [p(x) for x in j["uno_pct"].values()] + [pct(x, 2) for x in j["uno_pct_rel"].values()]
    j = R["e8"]
    for k in ("a", "m"):
        e += [p(j[k]["mc"]), fmt(j[k]["pe"], 1), p(j[k]["uo"]), fmt(j[k]["gao"], 2), pct(j[k]["ms"])]
    e += [p(j["indif"]), p(j["uo_indif"]), p(j["dif"])]
    j = R["e9"]
    e += [p(x) for x in j["mc_h"]] + [str(x) for x in j["q_ok"]] + [str(x) for x in j["q_mal"]]
    e += [p(j["mc_ok"]), p(j["mc_mal"]), p(j["dif"]), p(j["uo_ok"]), p(j["uo_mal"]), p(j["hora_extra"]), fmt(j["horas_dem"], 0)]
    e += [fmt(x, 0) for x in j["horas_ok"]] + [fmt(x, 0) for x in j["horas_mal"]]
    j = R["e10"]; bb = j["base"]
    e += [p(bb["mc"]), p(bb["uo"]), p(bb["pe"]), pct(bb["ms"]), fmt(j["q_baja"], 1), str(j["q_baja_ent"]),
          fmt(j["q_pub"], 1), str(j["q_pub_ent"]), pct(j["alza_baja"]), pct(j["alza_pub"]), p(j["p_min"]),
          p(j["uo_cap_baja"]), fmt(j["pe_baja"], 1), fmt(j["pe_pub"], 1)]
    # caso
    b = K["b"]
    e += [pct(x, 0) for x in b["mezcla"]] + [fmt(b["mcp"], 1), fmt(K["fijos"], 2), fmt(K["cif"], 2), p(K["admin"]),
                                             fmt(b["pe"], 2), p(b["pe"]), p(b["pe_pesos"]), p(b["pp"]), pct(b["razon"]),
                                             pct(b["ms"], 0), pct(b["ms"]), fmt(b["gao"], 2)]
    e += [fmt(x, 1) for x in b["pe_prod"]] + [p(x) for x in b["pe_ing"]]
    e += [p(K["ventas"]), p(b["mc_total"]), fmt(b["uo"], 2), p(K["interes"]), fmt(K["pe_interes"], 2), fmt(K["uai"], 2)]
    e += [p(K["q_tot"] * 0.8), fmt(K["uo_baja"], 2), pct(K["caida_baja"])] + [fmt(x, 2) for x in K["mc_seg"]]
    return e


def imprimir(R, K):
    def show(v):
        if isinstance(v, float):
            return f"{v:,.4f}"
        if isinstance(v, (list, tuple)):
            return "[" + ", ".join(show(x) for x in v) + "]"
        if isinstance(v, dict):
            return "{" + ", ".join(f"{k}: {show(x)}" for k, x in v.items()) + "}"
        return str(v)
    for k, d in list(R.items()) + [("caso", K)]:
        print(f"== {k}")
        for kk, v in d.items():
            print(f"   {kk}: {show(v)}")


if __name__ == "__main__":
    R = resultados()
    m, C = cargar_caso()
    K = caso(m, C)
    imprimir(R, K)
    u12 = C["u12"]
    b = K["b"]
    chequeos = {
        "ejemplo: en equilibrio la utilidad es cero": abs(R["ej"]["base"]["comprobacion"]) < 1e-6,
        "ejemplo: GAO = 1 / margen de seguridad": abs(R["ej"]["base"]["gao"] - 1 / R["ej"]["base"]["ms"]) < 1e-9,
        "ejemplo: la mezcla por margen por litro es el óptimo entero (fuerza bruta)": R["ej"]["bf"][0] == R["ej"]["mc_ok"],
        "12.4: PE en pesos = CF / razón de MC": abs(R["e4"]["pe"] * R["e4"]["razon"] - E4["cf"]) < 1e-6,
        "12.5: en equilibrio la utilidad es cero": abs(R["e5"]["comprobacion"]) < 1e-6,
        "12.5: PE en pesos por dos caminos": abs(R["e5"]["pe_pesos"] - E5["cf"] / R["e5"]["razon"]) < 1e-6,
        "12.8: en la indiferencia las dos utilidades son iguales":
            abs(R["e8"]["indif"] * R["e8"]["m"]["mc"] - E8["maquina"]["cf"] - R["e8"]["uo_indif"]) < 1e-6,
        "12.9: la mezcla por margen por hora es el óptimo entero (fuerza bruta)": R["e9"]["bf"][0] == R["e9"]["mc_ok"],
        "caso: margen ponderado = script": abs(b["mcp"] - u12["margen_ponderado"]) < 1e-9,
        "caso: fijos operacionales = script": abs(K["fijos"] - u12["costos_fijos_operacionales"]) < 1e-6,
        "caso: equilibrio en unidades = script": abs(b["pe"] - u12["equilibrio_unidades"]) < 1e-9,
        "caso: equilibrio en pesos = script": abs(b["pe_pesos"] - u12["equilibrio_pesos"]) < 1e-6,
        "caso: equilibrio por producto = script": all(abs(a - c) < 1e-9 for a, c in zip(b["pe_prod"], u12["equilibrio_por_producto"])),
        "caso: margen de seguridad = script": abs(b["ms"] - u12["margen_seguridad"]) < 1e-12,
        "caso: apalancamiento = script": abs(b["gao"] - u12["apalancamiento_operativo"]) < 1e-12,
        "caso: utilidad operacional = resultado por costeo variable de la unidad 10":
            abs(b["uo"] - C["u10"]["resultado_variable"]) < 1e-6,
        "caso: ventas propias = asiento 6 del script": abs(b["ventas"] - K["ventas"]) < 1e-6,
        "caso: en equilibrio la utilidad es cero": abs(b["comprobacion"]) < 1e-6,
    }
    print("\nCHEQUEOS:")
    for k, v in chequeos.items():
        print(f"   {'OK ' if v else 'MAL'} {k}")
    ok_planilla = planilla(R, K, m)
    if not all(chequeos.values()) or not ok_planilla:
        sys.exit(1)
    print("CADENA: OK, la pieza del caso calza con caso_estampa_sur_contable.py")
    if "--revisar" in sys.argv:
        md = (AQUI.parent / "unidades" / "12-costo-volumen-utilidad.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %")
        faltan = [x for x in esperados(R, K) if x not in md]
        print("\nREVISIÓN:", "OK, todas las cifras están en el texto" if not faltan else f"faltan {len(faltan)}: {faltan}")
        sys.exit(1 if faltan else 0)
