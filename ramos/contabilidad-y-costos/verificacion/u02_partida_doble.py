"""Unidad 2 · Partida doble: asientos, libro diario, mayor y balance de comprobación.

Recalcula todas las cifras de la unidad desde los datos, comprueba que cada asiento cuadra, arma el
mayor y el balance de comprobación de sumas y saldos, genera la planilla
`planillas/u02_partida_doble.xlsx` (libro diario, mayor y comprobación con fórmulas) y comprueba sus
resultados evaluando esas fórmulas en Python. Con --revisar busca cada resultado, y cada fila de las
tablas generadas, escrito tal cual en el .md.

    py -3 -P u02_partida_doble.py            # imprime los resultados y genera la planilla
    py -3 -P u02_partida_doble.py --tablas   # además imprime las tablas en Markdown
    py -3 -P u02_partida_doble.py --revisar  # además las busca en el .md

Todos los datos son originales de esta base (pasados por herramientas/colisiones.py). No vienen de
ningún libro ni curso. La única cifra real es la tasa de IVA (19 %).
"""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import importlib.util, sys

sys.stdout.reconfigure(encoding="utf-8")
AQUI = Path(__file__).resolve().parent
RAMO = AQUI.parent
IVA = 0.19


# ---------------------------------------------------------------- herramientas

def fmt(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, dec=1):
    return fmt(100 * x, dec) + " %"


def peso(x):
    """Redondeo al peso, mitad hacia arriba (como una factura)."""
    return int(Decimal(str(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def iva(neto):
    return peso(neto * IVA)


def cargar(nombre, archivo):
    spec = importlib.util.spec_from_file_location(nombre, AQUI / archivo)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def asiento(n, fecha, glosa, lineas):
    """lineas = [(cuenta, debe, haber)]. Falla si no cuadra."""
    d = sum(x[1] for x in lineas); h = sum(x[2] for x in lineas)
    assert abs(d - h) < 1e-6, (n, d, h)
    return dict(n=n, fecha=fecha, glosa=glosa, lineas=lineas, total=d)


def mayor(asientos):
    """{cuenta: [(n, debe, haber), ...]} en el orden de aparición."""
    M = {}
    for a in asientos:
        for c, d, h in a["lineas"]:
            M.setdefault(c, []).append((a["n"], d, h))
    return M


def comprobacion(asientos, orden=None):
    M = mayor(asientos)
    filas = []
    for c in (orden or list(M)):
        d = sum(x[1] for x in M[c]); h = sum(x[2] for x in M[c])
        filas.append(dict(cuenta=c, debe=d, haber=h, sd=max(d - h, 0), sa=max(h - d, 0)))
    tot = {k: sum(f[k] for f in filas) for k in ("debe", "haber", "sd", "sa")}
    return filas, tot


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · Capa a Capa, impresión 3D por encargo, mayo de 2026
EJ = dict(aporte=7_400_000, impresoras=3_260_000, pie_impresoras=0.5, filamento=870_000, arriendo_mes=365_000,
          venta=2_540_000, contado=0.30, costo_venta=510_000, sueldos=1_150_000, pago_prov=640_000,
          cobro=1_050_000, luz=64_000)

# 2.2 · un carro de comida: asientos simples sin IVA
E2 = dict(aporte=3_500_000, deposito=2_800_000, insumos=420_000, venta=690_000, permiso=85_000)
# 2.3 · compra con IVA, 40 % al contado
E3 = dict(neto=1_260_000, contado=0.40)
# 2.4 · venta con margen sobre el costo
E4 = dict(costo=740_000, margen=0.35)
# 2.5 · la lavandería: diario, mayor y comprobación sin IVA
E5 = dict(aporte=7_250_000, deposito=5_900_000, lavadoras=3_840_000, pie=0.30, detergente=515_000,
          servicios=1_410_000, efectivo=0.60, consumo=298_000, sueldos=1_180_000)
# 2.6 · la tienda de bicicletas: compra, devolución y venta con IVA
E6 = dict(unid=12, costo=178_000, devueltas=2, vendidas=7, precio=279_000, contado=0.5)
# 2.7 · cinco errores y el balance de comprobación
E7 = dict(venta=470_000, venta_mal=407_000, omitida=135_000, invertido=310_000, mal_cuenta=820_000, doble=265_000)
# 2.8 · del mayor al asiento
E8 = dict(prestamo=4_600_000, mercaderia=1_380_000, venta=760_000, amortizacion=470_000, interes=61_000, costo=475_000)
# 2.9 · el taller de muebles Raulí, julio de 2026
E9 = dict(aporte_banco=9_000_000, camioneta=7_200_000, madera=2_450_000, pie_madera=0.5, venta=3_900_000, contado=0.40,
          costo_venta=1_170_000, devol_neto=300_000, devol_costo=90_000, anticipo=950_000, herramientas=680_000,
          sueldos=1_600_000, arriendo=520_000)
# 2.10 · la agencia de viajes con tres errores que el balance de comprobación no ve
E10 = dict(aporte=5_200_000, computadores=1_800_000, seguro_anual=1_260_000, venta=3_450_000, sueldos=1_450_000)


# ---------------------------------------------------------------- el ejemplo

def ejemplo():
    e = EJ
    i_imp = iva(e["impresoras"]); t_imp = e["impresoras"] + i_imp
    i_fil = iva(e["filamento"])
    i_ven = iva(e["venta"]); t_ven = e["venta"] + i_ven
    ven_c = peso(t_ven * e["contado"]); ven_cr = t_ven - ven_c
    i_luz = iva(e["luz"])
    A = [
        asiento(1, "02-05-2026", "Aporte inicial de Fernanda", [("Banco", e["aporte"], 0), ("Capital", 0, e["aporte"])]),
        asiento(2, "03-05-2026", "Compra de dos impresoras, mitad al contado y mitad a 60 días",
                [("Maquinaria", e["impresoras"], 0), ("IVA crédito fiscal", i_imp, 0),
                 ("Banco", 0, peso(t_imp * e["pie_impresoras"])), ("Cuentas por pagar", 0, t_imp - peso(t_imp * e["pie_impresoras"]))]),
        asiento(3, "05-05-2026", "Compra de filamento a crédito del proveedor",
                [("Materiales", e["filamento"], 0), ("IVA crédito fiscal", i_fil, 0), ("Proveedores", 0, e["filamento"] + i_fil)]),
        asiento(4, "06-05-2026", "Arriendo de mayo y junio pagado por transferencia",
                [("Arriendo", e["arriendo_mes"], 0), ("Arriendo pagado por anticipado", e["arriendo_mes"], 0),
                 ("Banco", 0, 2 * e["arriendo_mes"])]),
        asiento(5, "20-05-2026", "Venta de piezas, 30 % por transferencia y el resto a crédito",
                [("Banco", ven_c, 0), ("Clientes", ven_cr, 0), ("Ventas", 0, e["venta"]), ("IVA débito fiscal", 0, i_ven)]),
        asiento(6, "20-05-2026", "Costo del filamento de las piezas vendidas",
                [("Costo de ventas", e["costo_venta"], 0), ("Materiales", 0, e["costo_venta"])]),
        asiento(7, "31-05-2026", "Sueldos de mayo pagados", [("Sueldos", e["sueldos"], 0), ("Banco", 0, e["sueldos"])]),
        asiento(8, "31-05-2026", "Abono a proveedores", [("Proveedores", e["pago_prov"], 0), ("Banco", 0, e["pago_prov"])]),
        asiento(9, "31-05-2026", "Cobro a clientes", [("Banco", e["cobro"], 0), ("Clientes", 0, e["cobro"])]),
        asiento(10, "31-05-2026", "Cuenta de luz de mayo, se paga en junio",
                [("Electricidad", e["luz"], 0), ("IVA crédito fiscal", i_luz, 0), ("Cuentas por pagar", 0, e["luz"] + i_luz)]),
    ]
    orden = ["Banco", "Clientes", "Materiales", "IVA crédito fiscal", "Arriendo pagado por anticipado", "Maquinaria",
             "Proveedores", "Cuentas por pagar", "IVA débito fiscal", "Capital", "Ventas", "Costo de ventas",
             "Arriendo", "Sueldos", "Electricidad"]
    filas, tot = comprobacion(A, orden)
    s = {f["cuenta"]: f["sd"] - f["sa"] for f in filas}
    util = e["venta"] - e["costo_venta"] - e["arriendo_mes"] - e["sueldos"] - e["luz"]
    cf = s["IVA crédito fiscal"]; df = -s["IVA débito fiscal"]
    R = dict(asientos=A, filas=filas, tot=tot, saldo=s, util=util, iva_cf=cf, iva_df=df, remanente=cf - df,
             i_imp=i_imp, t_imp=t_imp, i_ven=i_ven, t_ven=t_ven, ven_c=ven_c, ven_cr=ven_cr)
    return R


# ---------------------------------------------------------------- ejercicios

def ejercicios():
    R = {}
    e = E2
    caja = e["aporte"] - e["deposito"] + e["venta"]
    banco = e["deposito"] - e["insumos"] - e["permiso"]
    R["e2"] = dict(caja=caja, banco=banco, sumas=e["aporte"] + e["deposito"] + e["insumos"] + e["venta"] + e["permiso"])

    e = E3
    i = iva(e["neto"]); t = e["neto"] + i; c = peso(t * e["contado"])
    R["e3"] = dict(iva=i, total=t, contado=c, credito=t - c)

    e = E4
    precio = peso(e["costo"] * (1 + e["margen"]))
    i = iva(precio)
    precio_sp = peso(e["costo"] / (1 - e["margen"]))
    R["e4"] = dict(precio=precio, iva=i, total=precio + i, margen_pesos=precio - e["costo"],
                   margen_sobre_precio=(precio - e["costo"]) / precio, precio_sp=precio_sp)

    e = E5
    pie = peso(e["lavadoras"] * e["pie"])
    efectivo = peso(e["servicios"] * e["efectivo"])
    A = [asiento(1, "01-06-2026", "Aporte", [("Caja", e["aporte"], 0), ("Capital", 0, e["aporte"])]),
         asiento(2, "02-06-2026", "Depósito", [("Banco", e["deposito"], 0), ("Caja", 0, e["deposito"])]),
         asiento(3, "04-06-2026", "Lavadoras", [("Maquinaria", e["lavadoras"], 0), ("Banco", 0, pie),
                                                ("Documentos por pagar", 0, e["lavadoras"] - pie)]),
         asiento(4, "08-06-2026", "Detergente", [("Insumos", e["detergente"], 0), ("Banco", 0, e["detergente"])]),
         asiento(5, "15-06-2026", "Servicios", [("Caja", efectivo, 0), ("Clientes", e["servicios"] - efectivo, 0),
                                                ("Ingresos por servicios", 0, e["servicios"])]),
         asiento(6, "30-06-2026", "Consumo", [("Gasto en insumos", e["consumo"], 0), ("Insumos", 0, e["consumo"])]),
         asiento(7, "30-06-2026", "Sueldos", [("Sueldos", e["sueldos"], 0), ("Banco", 0, e["sueldos"])])]
    orden = ["Caja", "Banco", "Clientes", "Insumos", "Maquinaria", "Documentos por pagar", "Capital",
             "Ingresos por servicios", "Gasto en insumos", "Sueldos"]
    filas, tot = comprobacion(A, orden)
    R["e5"] = dict(asientos=A, filas=filas, tot=tot, pie=pie, doc=e["lavadoras"] - pie, efectivo=efectivo,
                   credito=e["servicios"] - efectivo,
                   util=e["servicios"] - e["consumo"] - e["sueldos"])

    e = E6
    neto_c = e["unid"] * e["costo"]; i_c = iva(neto_c)
    neto_d = e["devueltas"] * e["costo"]; i_d = iva(neto_d)
    neto_v = e["vendidas"] * e["precio"]; i_v = iva(neto_v); t_v = neto_v + i_v
    cont = peso(t_v * e["contado"])
    costo_v = e["vendidas"] * e["costo"]
    R["e6"] = dict(neto_c=neto_c, i_c=i_c, t_c=neto_c + i_c, neto_d=neto_d, i_d=i_d, t_d=neto_d + i_d,
                   neto_v=neto_v, i_v=i_v, t_v=t_v, contado=cont, credito=t_v - cont, costo_v=costo_v,
                   cf=i_c - i_d, df=i_v, a_pagar=i_v - (i_c - i_d),
                   saldo_merc=neto_c - neto_d - costo_v, proveedores=neto_c + i_c - neto_d - i_d,
                   margen_bruto=neto_v - costo_v)
    assert R["e6"]["saldo_merc"] == (e["unid"] - e["devueltas"] - e["vendidas"]) * e["costo"]

    e = E7
    dif_transp = e["venta"] - e["venta_mal"]
    exceso_haber = e["doble"] - dif_transp
    R["e7"] = dict(dif_transp=dif_transp, nueve=dif_transp / 9, exceso_haber=exceso_haber,
                   banco_sobre=2 * e["invertido"], prov_sobre=2 * e["invertido"])

    e = E8
    i_m = iva(e["mercaderia"]); i_v = iva(e["venta"])
    t_m = e["mercaderia"] + i_m; t_v = e["venta"] + i_v; cuota = e["amortizacion"] + e["interes"]
    R["e8"] = dict(i_m=i_m, t_m=t_m, i_v=i_v, t_v=t_v, cuota=cuota,
                   sumas=e["prestamo"] + t_m + t_v + cuota + e["costo"],
                   banco=e["prestamo"] - t_m + t_v - cuota, prestamo=e["prestamo"] - e["amortizacion"],
                   merc=e["mercaderia"] - e["costo"])

    e = E9
    i_mad = iva(e["madera"]); t_mad = e["madera"] + i_mad; pie_mad = peso(t_mad * e["pie_madera"])
    i_ven = iva(e["venta"]); t_ven = e["venta"] + i_ven; ven_c = peso(t_ven * e["contado"])
    i_dev = iva(e["devol_neto"])
    i_her = iva(e["herramientas"])
    A = [asiento(1, "01-07-2026", "Aportes de los socios", [("Banco", e["aporte_banco"], 0), ("Vehículos", e["camioneta"], 0),
                                                            ("Capital", 0, e["aporte_banco"] + e["camioneta"])]),
         asiento(2, "03-07-2026", "Compra de madera", [("Madera", e["madera"], 0), ("IVA crédito fiscal", i_mad, 0),
                                                       ("Banco", 0, pie_mad), ("Proveedores", 0, t_mad - pie_mad)]),
         asiento(3, "08-07-2026", "Venta de muebles", [("Banco", ven_c, 0), ("Clientes", t_ven - ven_c, 0),
                                                       ("Ventas", 0, e["venta"]), ("IVA débito fiscal", 0, i_ven)]),
         asiento(4, "08-07-2026", "Costo de la venta", [("Costo de ventas", e["costo_venta"], 0), ("Madera", 0, e["costo_venta"])]),
         asiento(5, "10-07-2026", "Devolución de una mesa", [("Devoluciones sobre ventas", e["devol_neto"], 0),
                                                             ("IVA débito fiscal", i_dev, 0), ("Clientes", 0, e["devol_neto"] + i_dev)]),
         asiento(6, "10-07-2026", "La mesa vuelve a bodega", [("Madera", e["devol_costo"], 0), ("Costo de ventas", 0, e["devol_costo"])]),
         asiento(7, "15-07-2026", "Anticipo de un cliente", [("Banco", e["anticipo"], 0),
                                                             ("Ingreso percibido por adelantado", 0, e["anticipo"])]),
         asiento(8, "20-07-2026", "Compra de herramientas", [("Herramientas", e["herramientas"], 0), ("IVA crédito fiscal", i_her, 0),
                                                             ("Banco", 0, e["herramientas"] + i_her)]),
         asiento(9, "31-07-2026", "Sueldos de julio", [("Sueldos", e["sueldos"], 0), ("Banco", 0, e["sueldos"])]),
         asiento(10, "31-07-2026", "Arriendo de julio", [("Arriendo", e["arriendo"], 0), ("Banco", 0, e["arriendo"])])]
    orden = ["Banco", "Clientes", "Madera", "IVA crédito fiscal", "Vehículos", "Herramientas", "Proveedores",
             "IVA débito fiscal", "Ingreso percibido por adelantado", "Capital", "Ventas", "Devoluciones sobre ventas",
             "Costo de ventas", "Sueldos", "Arriendo"]
    filas, tot = comprobacion(A, orden)
    ventas_netas = e["venta"] - e["devol_neto"]
    costo_neto = e["costo_venta"] - e["devol_costo"]
    util = ventas_netas - costo_neto - e["sueldos"] - e["arriendo"]
    df = i_ven - i_dev; cf = i_mad + i_her
    R["e9"] = dict(asientos=A, filas=filas, tot=tot, i_mad=i_mad, t_mad=t_mad, pie_mad=pie_mad, i_ven=i_ven, t_ven=t_ven,
                   ven_c=ven_c, ven_cr=t_ven - ven_c, i_dev=i_dev, t_dev=e["devol_neto"] + i_dev, i_her=i_her,
                   t_her=e["herramientas"] + i_her, ventas_netas=ventas_netas, costo_neto=costo_neto, util=util,
                   df=df, cf=cf, iva_pagar=df - cf, util_sin_devol=util + e["devol_neto"] - e["devol_costo"],
                   sobre_util=e["devol_neto"] - e["devol_costo"], sobre_iva=i_dev)

    e = E10
    i_comp = iva(e["computadores"]); i_ven = iva(e["venta"])
    seguro_mes = e["seguro_anual"] / 12
    util_mal = -e["venta"] - e["seguro_anual"] - e["sueldos"]
    util_bien = e["venta"] - seguro_mes - e["sueldos"]
    sumas = e["aporte"] + e["computadores"] + i_comp + e["seguro_anual"] + e["venta"] + i_ven + e["sueldos"]
    R["e10"] = dict(i_comp=i_comp, t_comp=e["computadores"] + i_comp, i_ven=i_ven, t_ven=e["venta"] + i_ven,
                    seguro_mes=seguro_mes, anticipado=e["seguro_anual"] - seguro_mes, util_mal=util_mal,
                    util_bien=util_bien, sumas=sumas, iva_pagar=i_ven - i_comp,
                    banco=e["aporte"] - e["computadores"] - i_comp - e["seguro_anual"] - e["sueldos"],
                    df_mal=i_comp + i_ven, dif_util=util_bien - util_mal)
    return R


# ---------------------------------------------------------------- caso integrador

ORDEN_CASO = ["Banco", "Clientes", "Poleras lisas básicas", "Otras prendas lisas", "Tinta", "IVA crédito fiscal",
              "Arriendo pagado por anticipado", "Seguro pagado por anticipado", "Maquinaria", "Depreciación acumulada",
              "Proveedores", "IVA débito fiscal", "Documentos por pagar a largo plazo", "Préstamo bancario", "Capital",
              "Ventas", "Sueldos", "Gastos de administración"]
FECHAS_CASO = {0: "01-03", 1: "01-03", 2: "01-03", 3: "02-03", 4: "04-03", 5: "04-03", 6: "31-03", 7: "31-03",
               8: "31-03", 9: "31-03", 10: "31-03", 11: "31-03", 12: "31-03"}


def caso():
    m = cargar("caso", "caso_estampa_sur_contable.py")
    C = m.calcular()
    a, p = m.BALANCE_INICIAL["activo"], m.BALANCE_INICIAL["pasivo"]
    cap = sum(a.values()) - sum(p.values())
    # asiento 0 · apertura con el balance de la unidad 1 (la depreciación acumulada va al Haber)
    ap = [(k, v, 0) for k, v in a.items() if v > 0] + [(k, 0, -v) for k, v in a.items() if v < 0]
    ap += [(k, 0, v) for k, v in p.items()] + [("Capital", 0, cap)]
    A = [asiento(0, FECHAS_CASO[0], "Apertura con el balance al 1 de marzo", ap)]
    ajustes = []
    for f in C["u2"]["asientos"]:
        lin = [(k, peso(v), 0) for k, v in f["debe"].items()] + [(k, 0, peso(v)) for k, v in f["haber"].items()]
        dif = sum(x[1] for x in lin) - sum(x[2] for x in lin)
        if dif:
            # la factura va en pesos enteros: el neto absorbe el medio peso para que neto + IVA = total
            i = max((j for j, x in enumerate(lin) if x[1] and "IVA" not in x[0]), key=lambda j: lin[j][1])
            ajustes.append((f["n"], lin[i][0], lin[i][1], lin[i][1] - dif, f["debe"][lin[i][0]]))
            lin[i] = (lin[i][0], lin[i][1] - dif, 0)
        A.append(asiento(f["n"], FECHAS_CASO[f["n"]], f["glosa"], lin))
    filas, tot = comprobacion(A, ORDEN_CASO)
    s = {f["cuenta"]: f["sd"] - f["sa"] for f in filas}
    # chequeo contra el script: cada línea a menos de un peso
    lineas_ok = True
    for f in C["u2"]["asientos"]:
        mio = next(x for x in A if x["n"] == f["n"])
        for c, d, h in mio["lineas"]:
            ref = f["debe"].get(c) if d else f["haber"].get(c)
            lineas_ok &= ref is not None and abs((d or h) - ref) < 1e-6
    # banco paso a paso (asientos 0 a 5) para mostrar el día de la máquina
    banco, paso = 0, {}
    for x in A:
        for c, d, h in x["lineas"]:
            if c == "Banco":
                banco += d - h
        paso[x["n"]] = banco
    sin_prestamo = paso[3] - next(h for c, d, h in A[4]["lineas"] if c == "Banco")
    activo_31 = sum(s[c] for c in ORDEN_CASO[:10])
    R = dict(asientos=A, filas=filas, tot=tot, saldo=s, ajustes=ajustes, lineas_ok=lineas_ok, paso=paso,
             sin_prestamo=sin_prestamo, iva_cf=s["IVA crédito fiscal"], iva_df=-s["IVA débito fiscal"],
             ventas=-s["Ventas"], n_asientos=len(C["u2"]["asientos"]), activo_31=activo_31,
             total_12=sum(x["total"] for x in A if x["n"] > 0), apertura=A[0]["total"],
             gastos=s["Sueldos"] + s["Gastos de administración"])
    return C, R


# ---------------------------------------------------------------- tablas en Markdown

def md_diario(A, con_fecha=True):
    out = ["| N.º | Fecha | Cuenta | Debe | Haber |", "|---|---|---|---|---|"]
    for a in A:
        primero = True
        for c, d, h in a["lineas"]:
            n = str(a["n"]) if primero else ""
            fe = a["fecha"] if primero else ""
            out.append(f"| {n} | {fe} | {'&emsp;' if not d else ''}{c} | {fmt(d) if d else ''} | {fmt(h) if h else ''} |")
            primero = False
        out.append(f"| | | *{a['glosa']}* | | |")
    return out


def md_comprobacion(filas, tot):
    out = ["| Cuenta | Sumas Debe | Sumas Haber | Saldo deudor | Saldo acreedor |", "|---|---|---|---|---|"]
    for f in filas:
        out.append(f"| {f['cuenta']} | {fmt(f['debe']) if f['debe'] else '0'} | {fmt(f['haber']) if f['haber'] else '0'} | "
                   f"{fmt(f['sd']) if f['sd'] else ''} | {fmt(f['sa']) if f['sa'] else ''} |")
    out.append(f"| **Totales** | **{fmt(tot['debe'])}** | **{fmt(tot['haber'])}** | **{fmt(tot['sd'])}** | **{fmt(tot['sa'])}** |")
    return out


def md_t(cuenta, movs, etiqueta=lambda n: f"({n})"):
    d = [x for x in movs if x[1]]; h = [x for x in movs if x[2]]
    out = [f"| Debe · {cuenta} | | Haber · {cuenta} | |", "|---|---|---|---|"]
    for i in range(max(len(d), len(h))):
        a = d[i] if i < len(d) else None; b = h[i] if i < len(h) else None
        out.append(f"| {etiqueta(a[0]) if a else ''} | {fmt(a[1]) if a else ''} | {etiqueta(b[0]) if b else ''} | {fmt(b[2]) if b else ''} |")
    sd = sum(x[1] for x in movs); sh = sum(x[2] for x in movs)
    out.append(f"| Sumas | {fmt(sd)} | | {fmt(sh)} |")
    if sd >= sh:
        out.append(f"| **Saldo deudor** | **{fmt(sd - sh)}** | | |")
    else:
        out.append(f"| | | **Saldo acreedor** | **{fmt(sh - sd)}** |")
    return out


def tablas(EJR, X, CR):
    T = {}
    T["ej_diario"] = md_diario(EJR["asientos"])
    M = mayor(EJR["asientos"])
    for c in ("Banco", "Clientes", "IVA crédito fiscal", "Proveedores"):
        T[f"ej_t_{c}"] = md_t(c, M[c])
    T["ej_comp"] = md_comprobacion(EJR["filas"], EJR["tot"])
    T["e5_comp"] = md_comprobacion(X["e5"]["filas"], X["e5"]["tot"])
    T["e9_diario"] = md_diario(X["e9"]["asientos"])
    T["e9_comp"] = md_comprobacion(X["e9"]["filas"], X["e9"]["tot"])
    T["caso_diario"] = md_diario(CR["asientos"])
    M = mayor(CR["asientos"])
    et = lambda n: "Ap." if n == 0 else f"({n})"
    for c in ("Banco", "Proveedores", "Clientes", "IVA crédito fiscal", "IVA débito fiscal"):
        T[f"caso_t_{c}"] = md_t(c, M[c], et)
    T["caso_comp"] = md_comprobacion(CR["filas"], CR["tot"])
    return T


# ---------------------------------------------------------------- planilla (con fórmulas)

def hoja_diario(wb, nombre, A):
    from openpyxl.styles import Font
    ws = wb.create_sheet(nombre)
    ws.append(["Asiento", "Fecha", "Cuenta", "Debe", "Haber", "Glosa"])
    for c in ws[1]:
        c.font = Font(bold=True)
    filas = {}
    for a in A:
        for k, (c, d, h) in enumerate(a["lineas"]):
            ws.append([a["n"], a["fecha"], c, d or None, h or None, a["glosa"] if k == 0 else None])
            filas.setdefault(c, []).append(ws.max_row)
    ultima = ws.max_row
    ws.append([]); ws.append(["", "", "Totales", f"=SUM(D2:D{ultima})", f"=SUM(E2:E{ultima})"])
    tot = ws.max_row
    ws.append(["", "", "Diferencia (debe dar 0)", f"=D{tot}-E{tot}"])
    # control por asiento
    ws["H1"] = "Asiento"; ws["I1"] = "Debe"; ws["J1"] = "Haber"; ws["K1"] = "Cuadra (0)"
    for i, a in enumerate(A, start=2):
        ws[f"H{i}"] = a["n"]
        ws[f"I{i}"] = f"=SUMIF($A$2:$A${ultima},H{i},$D$2:$D${ultima})"
        ws[f"J{i}"] = f"=SUMIF($A$2:$A${ultima},H{i},$E$2:$E${ultima})"
        ws[f"K{i}"] = f"=I{i}-J{i}"
    for col, w in zip("ABCDEFGHIJK", (8, 8, 34, 14, 14, 52, 3, 8, 14, 14, 11)):
        ws.column_dimensions[col].width = w
    return ws, ultima, filas, dict(debe=f"D{tot}", haber=f"E{tot}", dif=f"D{tot + 1}",
                                   control=[f"K{i}" for i in range(2, len(A) + 2)])


def hoja_mayor(wb, nombre, diario, filas, orden):
    ws = wb.create_sheet(nombre)
    r = 1
    celdas = {}
    for c in orden:
        ws.cell(r, 1, c); r += 1
        ws.cell(r, 1, "Asiento"); ws.cell(r, 2, "Debe"); ws.cell(r, 3, "Haber"); ws.cell(r, 4, "Saldo (Debe − Haber)"); r += 1
        ini = r
        for fd in filas[c]:
            ws.cell(r, 1, f"='{diario}'!A{fd}")
            ws.cell(r, 2, f"='{diario}'!D{fd}")
            ws.cell(r, 3, f"='{diario}'!E{fd}")
            ws.cell(r, 4, f"=B{r}-C{r}" if r == ini else f"=D{r-1}+B{r}-C{r}")
            r += 1
        ws.cell(r, 1, "Sumas"); ws.cell(r, 2, f"=SUM(B{ini}:B{r-1})"); ws.cell(r, 3, f"=SUM(C{ini}:C{r-1})")
        ws.cell(r, 4, f"=D{r-1}")
        celdas[c] = dict(debe=f"B{r}", haber=f"C{r}", saldo=f"D{r}")
        r += 2
    for col, w in zip("ABCD", (34, 14, 14, 20)):
        ws.column_dimensions[col].width = w
    return celdas


def hoja_comprobacion(wb, nombre, diario, ultima, orden):
    ws = wb.create_sheet(nombre)
    ws.append(["Cuenta", "Sumas Debe", "Sumas Haber", "Saldo deudor", "Saldo acreedor"])
    celdas = {}
    for i, c in enumerate(orden, start=2):
        ws[f"A{i}"] = c
        ws[f"B{i}"] = f"=SUMIF('{diario}'!$C$2:$C${ultima},A{i},'{diario}'!$D$2:$D${ultima})"
        ws[f"C{i}"] = f"=SUMIF('{diario}'!$C$2:$C${ultima},A{i},'{diario}'!$E$2:$E${ultima})"
        ws[f"D{i}"] = f"=MAX(B{i}-C{i},0)"
        ws[f"E{i}"] = f"=MAX(C{i}-B{i},0)"
        celdas[c] = dict(debe=f"B{i}", haber=f"C{i}", sd=f"D{i}", sa=f"E{i}")
    n = len(orden) + 1
    ws.append(["Totales"] + [f"=SUM({x}2:{x}{n})" for x in "BCDE"])
    t = n + 1
    ws.append(["Sumas iguales (0)", f"=B{t}-C{t}", "", f"=D{t}-E{t}"])
    celdas["_tot"] = dict(debe=f"B{t}", haber=f"C{t}", sd=f"D{t}", sa=f"E{t}", dif_sumas=f"B{t+1}", dif_saldos=f"D{t+1}")
    for col, w in zip("ABCDE", (34, 15, 15, 15, 15)):
        ws.column_dimensions[col].width = w
    return celdas


def generar_planilla(CR, EJR):
    from openpyxl import Workbook
    wb = Workbook()
    portada = wb.active
    portada.title = "Léeme"
    for fila in [["Unidad 2 · Partida doble (base abierta de ingeniería)"],
                 ["Hojas: diario, mayor y balance de comprobación del caso Estampa Sur (marzo) y del ejemplo Capa a Capa (mayo de 2026)."],
                 ["Todo se calcula con fórmulas desde la hoja del diario: si cambias un monto, el mayor y la comprobación se actualizan."],
                 ["Montos en pesos, IVA 19 % redondeado al peso."]]:
        portada.append(fila)
    portada.column_dimensions["A"].width = 110
    out = {}
    for pref, R, orden in (("Caso", CR, ORDEN_CASO), ("Ejemplo", EJR, [f["cuenta"] for f in EJR["filas"]])):
        dn = f"{pref} diario"
        _, ultima, filas, cd = hoja_diario(wb, dn, R["asientos"])
        cm = hoja_mayor(wb, f"{pref} mayor", dn, filas, orden)
        cc = hoja_comprobacion(wb, f"{pref} comprobación", dn, ultima, orden)
        out[pref] = (dn, cd, cm, cc)
    destino = RAMO / "planillas" / "u02_partida_doble.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)
    return destino, out


def revisar_planilla(ruta, out, CR, EJR):
    u01 = cargar("u01", "u01_estados_financieros.py")
    ev = u01.Evaluador(ruta)
    ok = {}
    for pref, R in (("Caso", CR), ("Ejemplo", EJR)):
        dn, cd, cm, cc = out[pref]
        ok[f"{pref}: totales del diario"] = (abs(ev.valor(dn, cd["debe"]) - R["tot"]["debe"]) < 1e-6
                                            and abs(ev.valor(dn, cd["dif"])) < 1e-6)
        ok[f"{pref}: cada asiento cuadra en la planilla"] = all(abs(ev.valor(dn, c)) < 1e-6 for c in cd["control"])
        ok[f"{pref}: mayor = saldos calculados"] = all(abs(ev.valor(f"{pref} mayor", cm[f["cuenta"]]["saldo"]) - (f["sd"] - f["sa"])) < 1e-6
                                                      for f in R["filas"])
        hc = f"{pref} comprobación"
        ok[f"{pref}: comprobación fila por fila"] = all(abs(ev.valor(hc, cc[f["cuenta"]][k]) - f[k]) < 1e-6
                                                       for f in R["filas"] for k in ("debe", "haber", "sd", "sa"))
        t = cc["_tot"]
        ok[f"{pref}: sumas y saldos iguales"] = (abs(ev.valor(hc, t["dif_sumas"])) < 1e-6 and abs(ev.valor(hc, t["dif_saldos"])) < 1e-6
                                                 and abs(ev.valor(hc, t["sd"]) - R["tot"]["sd"]) < 1e-6)
    return ok


# ---------------------------------------------------------------- cifras esperadas en el texto

def esperados(EJR, X, CR, C):
    p = fmt
    e = []
    j = EJR
    e += [p(j["i_imp"]), p(j["t_imp"]), p(j["i_ven"]), p(j["t_ven"]), p(j["ven_c"]), p(j["ven_cr"]), p(j["util"]),
          p(j["iva_cf"]), p(j["iva_df"]), p(j["remanente"]), p(j["saldo"]["Banco"]), p(j["saldo"]["Clientes"]),
          p(j["tot"]["debe"]), p(j["tot"]["sd"])]
    j = X["e2"]; e += [p(j["caja"]), p(j["banco"]), p(j["sumas"])]
    j = X["e3"]; e += [p(j["iva"]), p(j["total"]), p(j["contado"]), p(j["credito"])]
    j = X["e4"]; e += [p(j["precio"]), p(j["iva"]), p(j["total"]), p(j["margen_pesos"]), pct(j["margen_sobre_precio"]),
                       p(j["precio_sp"])]
    j = X["e5"]; e += [p(j["pie"]), p(j["doc"]), p(j["efectivo"]), p(j["credito"]), p(j["util"]), p(j["tot"]["debe"]),
                       p(j["tot"]["sd"])]
    j = X["e6"]; e += [p(j[k]) for k in ("neto_c", "i_c", "t_c", "neto_d", "i_d", "t_d", "neto_v", "i_v", "t_v", "contado",
                                         "costo_v", "cf", "df", "a_pagar", "saldo_merc", "proveedores", "margen_bruto")]
    j = X["e7"]; e += [p(j["dif_transp"]), p(j["nueve"]), p(j["exceso_haber"]), p(j["banco_sobre"])]
    j = X["e8"]; e += [p(j[k]) for k in ("i_m", "t_m", "i_v", "t_v", "cuota", "sumas", "banco", "prestamo", "merc")]
    j = X["e9"]; e += [p(j[k]) for k in ("i_mad", "t_mad", "pie_mad", "i_ven", "t_ven", "ven_c", "ven_cr", "i_dev", "t_dev",
                                         "i_her", "t_her", "ventas_netas", "costo_neto", "util", "df", "cf", "iva_pagar",
                                         "util_sin_devol", "sobre_util", "sobre_iva")]
    e += [p(X["e9"]["tot"]["debe"]), p(X["e9"]["tot"]["sd"])]
    j = X["e10"]; e += [p(j[k]) for k in ("i_comp", "t_comp", "i_ven", "t_ven", "seguro_mes", "anticipado", "util_mal",
                                          "util_bien", "sumas", "iva_pagar", "banco", "df_mal", "dif_util")]
    j = CR
    e += [p(j["iva_cf"]), p(j["iva_df"]), p(j["ventas"]), p(j["saldo"]["Banco"]), p(j["saldo"]["Proveedores"] * -1),
          p(j["saldo"]["Clientes"]), p(j["tot"]["debe"]), p(j["tot"]["sd"]), p(j["sin_prestamo"]), p(j["apertura"]),
          p(j["total_12"]), p(j["paso"][3]), p(j["activo_31"]), p(j["gastos"])]
    m = cargar("caso_e", "caso_estampa_sur_contable.py")
    exacto = C["u2"]["compra_tinta_kg"] * m.TINTA_KG
    tinta = next(d for c, d, h in CR["asientos"][3]["lineas"] if c == "Tinta")
    e += [fmt(C["u2"]["compra_tinta_kg"], 3), fmt(exacto, 1), p(tinta)]
    return e


if __name__ == "__main__":
    EJR = ejemplo()
    X = ejercicios()
    C, CR = caso()
    print("== ejemplo", {k: v for k, v in EJR.items() if k not in ("asientos", "filas")})
    for k, v in X.items():
        print(f"== {k}", {kk: vv for kk, vv in v.items() if kk not in ("asientos", "filas")})
    print("== caso", {k: v for k, v in CR.items() if k not in ("asientos", "filas")})

    ruta, out = generar_planilla(CR, EJR)
    chequeos = {
        "caso: 12 asientos, todos cuadran": CR["n_asientos"] == 12,
        "caso: cada línea igual a la del script": CR["lineas_ok"],
        "caso: el script ya trae cada asiento en pesos enteros": not CR["ajustes"],
        "caso: ventas netas = script": CR["ventas"] == C["u2"]["ventas_netas"],
        "caso: IVA crédito = script (unidad 3)": abs(CR["iva_cf"] - C["u3"]["credito"]) < 1e-6,
        "caso: IVA débito = script (unidad 3)": abs(CR["iva_df"] - C["u3"]["debito"]) < 1e-6,
        "caso: sumas iguales": CR["tot"]["debe"] == CR["tot"]["haber"],
        "caso: saldos iguales": CR["tot"]["sd"] == CR["tot"]["sa"],
        "caso: sumas = apertura + los 12 asientos": CR["tot"]["debe"] == CR["apertura"] + CR["total_12"],
        "ejemplo: sumas y saldos iguales": EJR["tot"]["debe"] == EJR["tot"]["haber"] and EJR["tot"]["sd"] == EJR["tot"]["sa"],
        "2.5: sumas y saldos iguales": X["e5"]["tot"]["debe"] == X["e5"]["tot"]["haber"] and X["e5"]["tot"]["sd"] == X["e5"]["tot"]["sa"],
        "2.9: sumas y saldos iguales": X["e9"]["tot"]["debe"] == X["e9"]["tot"]["haber"] and X["e9"]["tot"]["sd"] == X["e9"]["tot"]["sa"],
    }
    chequeos.update({"planilla · " + k: v for k, v in revisar_planilla(ruta, out, CR, EJR).items()})
    print(f"\nPLANILLA: {ruta}")
    print("CHEQUEOS:")
    for k, v in chequeos.items():
        print(f"   {'OK ' if v else 'MAL'} {k}")
    if not all(chequeos.values()):
        sys.exit(1)
    print("CADENA: OK, la pieza del caso calza con caso_estampa_sur_contable.py")
    T = tablas(EJR, X, CR)
    if "--tablas" in sys.argv:
        for k, v in T.items():
            print(f"\n<!-- {k} -->")
            print("\n".join(v))
    if "--revisar" in sys.argv:
        md = (RAMO / "unidades" / "02-partida-doble.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %").replace("−", "-")
        faltan = [x for x in esperados(EJR, X, CR, C) if x not in md]
        lineas = set(l.strip() for l in md.splitlines())
        faltan_t = [f"{k}: {l}" for k, v in T.items() for l in v if l.strip() not in lineas]
        print("\nREVISIÓN:", "OK, todas las cifras y las filas de las tablas están en el texto" if not (faltan or faltan_t)
              else f"faltan {len(faltan)} cifras: {faltan}\n           y {len(faltan_t)} filas: {faltan_t[:15]}")
        sys.exit(1 if (faltan or faltan_t) else 0)
