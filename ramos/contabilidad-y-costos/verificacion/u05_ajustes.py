"""Unidad 5 · Ajustes de cierre: anticipados, devengados y provisiones.

Recalcula todas las cifras de la unidad desde los datos, comprueba que cada asiento cuadre, genera la
planilla `planillas/u05_ajustes.xlsx` (con fórmulas, no valores pegados) y evalúa esas fórmulas para
comprobar que dan lo mismo que el cálculo en Python. Con --revisar busca cada resultado en el .md.

    py -3 -P u05_ajustes.py            # imprime los resultados y genera la planilla
    py -3 -P u05_ajustes.py --revisar  # además los busca en el .md

Todos los datos son originales de esta base (pasados por herramientas/colisiones.py). La única cifra
legal que se usa es la tasa de IVA (19 %).
"""
from pathlib import Path
import importlib.util, re, sys

sys.stdout.reconfigure(encoding="utf-8")
IVA = 0.19
AQUI = Path(__file__).resolve().parent
RAMO = AQUI.parent


# ---------------------------------------------------------------- herramientas

def fmt(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def p(x):
    return fmt(round(x), 0)


ASIENTOS = []


def asiento(nombre, fecha, debe, haber, glosa):
    """Registra un asiento y comprueba que cuadre."""
    d, h = sum(debe.values()), sum(haber.values())
    assert abs(d - h) < 1e-6, f"{nombre} no cuadra: {d} contra {h}"
    a = dict(nombre=nombre, fecha=fecha, debe=debe, haber=haber, glosa=glosa, total=d)
    ASIENTOS.append(a)
    return a


def devengo(monto, meses_cubre, meses_consumidos):
    por_mes = monto / meses_cubre
    dev = por_mes * meses_consumidos
    return dict(por_mes=por_mes, devengado=dev, saldo=monto - dev)


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · Hostal Las Gaviotas, cierre anual al 31-12-2026
EJ = {
    "seguro": {"monto": 1_236_000, "cubre": 12, "consumidos": 5},       # pagado 01-08-2026, neto
    "estacionamiento": {"monto": 930_000, "cubre": 6, "consumidos": 2},  # APPA pagado 01-11-2026
    "local": {"monto": 2_580_000, "cubre": 6, "consumidos": 4},          # ARA recibido 01-09-2026
    "sueldos_dic": 3_150_000,
    "luz_dic": 184_000,                                                  # neto, factura del 31-12
    "deposito": {"capital": 5_000_000, "tasa": 0.004, "meses": 2},       # interés simple, desde 01-11
    "clientes": 6_420_000, "pct_incobrable": 0.03, "estimacion_previa": 45_000,
    "resultado_previo": 3_900_000,
}
# 5.1 · base caja contra base devengada (peluquería canina, abril, montos netos sin IVA)
E1 = {"servicios_abril": 2_460_000, "cobrado_de_abril": 1_980_000, "cobro_marzo": 310_000,
      "anticipo_mayo": 240_000, "sueldos": 1_150_000, "arriendo_mayo_pagado": 360_000,
      "arriendo_abril_pagado_en_marzo": 360_000, "insumos_consumidos": 185_000, "agua_abril": 42_000}
# 5.2 · seguro de una flota pagado el 01-05-2026 por 12 meses, cierre al 30-09-2026
E2 = {"monto": 1_548_000, "cubre": 12, "consumidos": 5}
# 5.3 · taller de música cobrado por adelantado el 01-03-2026, 6 meses; estados al 30-06-2026
E3 = {"monto": 1_380_000, "cubre": 6, "consumidos": 4}
# 5.4 · las cuatro familias en una mueblería al 31-05-2026
E4 = {"sueldos": 2_240_000, "intereses_por_cobrar": 18_500, "anticipo_cliente": 650_000, "gas_neto": 96_000}
# 5.5 · arriendo registrado como gasto al pagar (método del resultado)
E5 = {"monto": 2_720_000, "cubre": 8, "consumidos": 3}
# 5.6 · remuneraciones de mayo con descuentos (montos dados como dato)
E6 = {"bruto": 4_860_000, "desc_previsionales": 1_003_590, "impuesto_unico": 41_200, "aportes_empleador": 145_800}
# 5.7 · estimación de incobrables por antigüedad
E7 = {"tramos": [("al día", 8_400_000, 0.01), ("1 a 30 días", 2_150_000, 0.04),
                 ("31 a 90 días", 960_000, 0.15), ("más de 90 días", 380_000, 0.50)],
      "saldo_previo": 210_000, "castigo": 130_000}
# 5.8 · provisión de vacaciones
E8 = {"trabajadores": [("Ana", 26_000, 11), ("Benjamín", 31_000, 7.5), ("Carla", 22_500, 15), ("Diego", 40_000, 4)],
      "saldo_previo": 742_000}
# 5.9 · efecto de los ajustes omitidos
E9 = {"seguro_consumido": 437_000, "ara_devengado": 915_000, "sueldos": 1_285_000, "intereses_cobrar": 74_000,
      "incobrables": 263_000, "resultado_informado": 5_240_000, "bono": 0.08}
# 5.10 · reconstruir los ajustes desde dos balances de comprobación (pólizas y contrato de 12 meses)
E10 = {"sppa": (2_040_000, 1_530_000), "ara": (3_780_000, 2_520_000), "sueldos_por_pagar": (0, 1_910_000),
       "estimacion": (128_000, 317_000), "intereses_por_cobrar": (0, 57_500), "dep_acumulada": (4_920_000, 5_535_000),
       "meses_poliza": 12, "meses_contrato": 12}


# ---------------------------------------------------------------- resultados

def resultados():
    R = {}
    e = EJ
    s = devengo(**{"monto": e["seguro"]["monto"], "meses_cubre": e["seguro"]["cubre"],
                   "meses_consumidos": e["seguro"]["consumidos"]})
    est = devengo(e["estacionamiento"]["monto"], e["estacionamiento"]["cubre"], e["estacionamiento"]["consumidos"])
    loc = devengo(e["local"]["monto"], e["local"]["cubre"], e["local"]["consumidos"])
    luz_iva = e["luz_dic"] * IVA
    interes = e["deposito"]["capital"] * e["deposito"]["tasa"] * e["deposito"]["meses"]
    requerida = e["clientes"] * e["pct_incobrable"]
    aj_inc = requerida - e["estimacion_previa"]
    pago_seguro = asiento("ej-pago", "01-08-2026", {"Seguro pagado por anticipado": e["seguro"]["monto"],
                                                    "IVA crédito fiscal": e["seguro"]["monto"] * IVA},
                          {"Banco": e["seguro"]["monto"] * (1 + IVA)}, "pago de la póliza anual")
    aj = [
        asiento("A1", "31-12-2026", {"Gasto por seguro": s["devengado"]}, {"Seguro pagado por anticipado": s["devengado"]}, "5 meses de seguro"),
        asiento("A2", "31-12-2026", {"Gasto por arriendo": est["devengado"]}, {"Arriendo pagado por anticipado": est["devengado"]}, "2 meses de estacionamiento"),
        asiento("A3", "31-12-2026", {"Arriendo recibido por anticipado": loc["devengado"]}, {"Arriendo ganado": loc["devengado"]}, "4 meses del local"),
        asiento("A4", "31-12-2026", {"Gasto en sueldos": e["sueldos_dic"]}, {"Sueldos por pagar": e["sueldos_dic"]}, "sueldos de diciembre"),
        asiento("A5", "31-12-2026", {"Gasto en electricidad": e["luz_dic"], "IVA crédito fiscal": luz_iva},
                {"Cuentas por pagar": e["luz_dic"] + luz_iva}, "luz de diciembre"),
        asiento("A6", "31-12-2026", {"Intereses por cobrar": interes}, {"Intereses ganados": interes}, "2 meses del depósito"),
        asiento("A7", "31-12-2026", {"Gasto por incobrables": aj_inc}, {"Estimación de deudores incobrables": aj_inc}, "estimación al 3 %"),
    ]
    gastos = s["devengado"] + est["devengado"] + e["sueldos_dic"] + e["luz_dic"] + aj_inc
    ingresos = loc["devengado"] + interes
    R["ej"] = dict(seguro=s, est=est, local=loc, iva_seguro=e["seguro"]["monto"] * IVA,
                   banco_seguro=pago_seguro["total"], luz_iva=luz_iva, luz_total=e["luz_dic"] + luz_iva,
                   interes=interes, requerida=requerida, aj_inc=aj_inc, gastos=gastos, ingresos=ingresos,
                   efecto=ingresos - gastos, resultado=e["resultado_previo"] + ingresos - gastos,
                   total_diario=sum(a["total"] for a in aj), neto_clientes=e["clientes"] - requerida)

    e = E1
    caja_in = e["cobrado_de_abril"] + e["cobro_marzo"] + e["anticipo_mayo"]
    caja_out = e["sueldos"] + e["arriendo_mayo_pagado"]
    dev_in = e["servicios_abril"]
    dev_out = e["sueldos"] + e["arriendo_abril_pagado_en_marzo"] + e["insumos_consumidos"] + e["agua_abril"]
    R["e1"] = dict(caja_in=caja_in, caja_out=caja_out, caja=caja_in - caja_out, dev_in=dev_in, dev_out=dev_out,
                   dev=dev_in - dev_out, por_cobrar=e["servicios_abril"] - e["cobrado_de_abril"],
                   dif=(caja_in - caja_out) - (dev_in - dev_out),
                   cobros_ajenos=e["cobro_marzo"] + e["anticipo_mayo"], dif_ingresos=caja_in - dev_in,
                   dif_gastos=dev_out - caja_out)
    assert abs(R["e1"]["dif_ingresos"] + R["e1"]["dif_gastos"] - R["e1"]["dif"]) < 1e-6

    d = devengo(E2["monto"], E2["cubre"], E2["consumidos"])
    asiento("5.2-pago", "01-05-2026", {"Seguro pagado por anticipado": E2["monto"], "IVA crédito fiscal": E2["monto"] * IVA},
            {"Banco": E2["monto"] * (1 + IVA)}, "pago de la póliza")
    asiento("5.2-aj", "30-09-2026", {"Gasto por seguro": d["devengado"]}, {"Seguro pagado por anticipado": d["devengado"]}, "5 meses")
    R["e2"] = dict(d, iva=E2["monto"] * IVA, banco=E2["monto"] * (1 + IVA))

    d = devengo(E3["monto"], E3["cubre"], E3["consumidos"])
    asiento("5.3-aj", "30-06-2026", {"Ingreso recibido por anticipado": d["devengado"]}, {"Ingresos por servicios": d["devengado"]}, "4 meses")
    R["e3"] = d

    e = E4
    gas_iva = e["gas_neto"] * IVA
    asiento("5.4-d", "31-05-2026", {"Gasto en gas": e["gas_neto"], "IVA crédito fiscal": gas_iva},
            {"Cuentas por pagar": e["gas_neto"] + gas_iva}, "gas de mayo")
    R["e4"] = dict(gas_iva=gas_iva, gas_total=e["gas_neto"] + gas_iva,
                   efecto=-e["sueldos"] + e["intereses_por_cobrar"] - e["gas_neto"])

    d = devengo(E5["monto"], E5["cubre"], E5["consumidos"])
    asiento("5.5-aj", "31-12-2026", {"Arriendo pagado por anticipado": d["saldo"]}, {"Gasto por arriendo": d["saldo"]}, "parte no consumida")
    R["e5"] = d

    e = E6
    liquido = e["bruto"] - e["desc_previsionales"] - e["impuesto_unico"]
    a = asiento("5.6-dev", "31-05-2026", {"Gasto en remuneraciones": e["bruto"], "Aportes del empleador": e["aportes_empleador"]},
                {"Retenciones previsionales por pagar": e["desc_previsionales"] + e["aportes_empleador"],
                 "Impuesto único por pagar": e["impuesto_unico"], "Remuneraciones por pagar": liquido}, "remuneraciones de mayo")
    R["e6"] = dict(liquido=liquido, retenciones=e["desc_previsionales"] + e["aportes_empleador"], gasto=a["total"],
                   pct_liquido=liquido / a["total"])

    e = E7
    req = [m * t for _, m, t in e["tramos"]]
    total_req = sum(req)
    clientes = sum(m for _, m, _ in e["tramos"])
    aj = total_req - e["saldo_previo"]
    asiento("5.7-aj", "31-12-2026", {"Gasto por incobrables": aj}, {"Estimación de deudores incobrables": aj}, "ajuste")
    asiento("5.7-castigo", "15-01-2027", {"Estimación de deudores incobrables": e["castigo"]}, {"Clientes": e["castigo"]}, "castigo")
    R["e7"] = dict(req=req, total=total_req, aj=aj, clientes=clientes, neto=clientes - total_req,
                   neto_despues=(clientes - e["castigo"]) - (total_req - e["castigo"]),
                   clientes_despues=clientes - e["castigo"], estimacion_despues=total_req - e["castigo"])

    e = E8
    por = [c * dd for _, c, dd in e["trabajadores"]]
    tot = sum(por)
    R["e8"] = dict(por=por, total=tot, aj=tot - e["saldo_previo"], uso_diego=por[3], queda=tot - por[3])

    e = E9
    efecto = -e["seguro_consumido"] + e["ara_devengado"] - e["sueldos"] + e["intereses_cobrar"] - e["incobrables"]
    activo = e["seguro_consumido"] + e["incobrables"] - e["intereses_cobrar"]      # sobrevaloración
    pasivo = e["sueldos"] - e["ara_devengado"]                                      # subvaloración
    corregido = e["resultado_informado"] + efecto
    assert abs((activo + pasivo) - (-efecto)) < 1e-6
    R["e9"] = dict(efecto=efecto, activo=activo, pasivo=pasivo, corregido=corregido,
                   bono_inf=e["resultado_informado"] * e["bono"], bono_ok=corregido * e["bono"],
                   bono_de_mas=(e["resultado_informado"] - corregido) * e["bono"])

    e = E10
    dif = {k: v[1] - v[0] for k, v in e.items() if isinstance(v, tuple)}
    seguro_mes = e["sppa"][0] / e["meses_poliza"]
    ara_mes = e["ara"][0] / e["meses_contrato"]
    # un anticipado que baja es gasto (SPPA) o ingreso (ARA); un pasivo o una estimación que sube es gasto
    efecto = (dif["sppa"] - dif["ara"] - dif["sueldos_por_pagar"] - dif["estimacion"]
              + dif["intereses_por_cobrar"] - dif["dep_acumulada"])
    R["e10"] = dict(gasto_seguro=-dif["sppa"], ingreso_ara=-dif["ara"], sueldos=dif["sueldos_por_pagar"],
                    incobrables=dif["estimacion"], intereses=dif["intereses_por_cobrar"], depreciacion=dif["dep_acumulada"],
                    seguro_mes=seguro_mes, meses_seguro=-dif["sppa"] / seguro_mes, ara_mes=ara_mes,
                    meses_ara=-dif["ara"] / ara_mes, efecto=efecto)
    return R


# ---------------------------------------------------------------- caso integrador

def cargar_caso():
    spec = importlib.util.spec_from_file_location("caso", AQUI / "caso_estampa_sur_contable.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m, m.calcular()


def caso(m, C):
    """Recálculo propio de la pieza u5 desde los datos declarados del script."""
    arriendo_inicial = m.BALANCE_INICIAL["activo"]["Arriendo pagado por anticipado"]   # marzo, pagado en febrero
    arr = devengo(arriendo_inicial, 1, 1)
    seg = devengo(m.SEGURO_ANUAL, 12, 1)
    asiento("C-5.1", "31-03", {"Costo indirecto: arriendo del taller": arr["devengado"]},
            {"Arriendo pagado por anticipado": arr["devengado"]}, "arriendo de marzo")
    asiento("C-5.2", "31-03", {"Costo indirecto: seguro": seg["devengado"]},
            {"Seguro pagado por anticipado": seg["devengado"]}, "un mes de seguro")
    appa_final = arriendo_inicial + m.ARRIENDO_TRIMESTRE - arr["devengado"]
    sppa_final = m.SEGURO_ANUAL - seg["devengado"]
    propio = dict(arriendo_devengado=arr["devengado"], seguro_devengado=seg["devengado"],
                  arriendo_anticipado_al_31=appa_final, seguro_anticipado_al_31=sppa_final)
    extra = dict(arriendo_inicial=arriendo_inicial, trimestre=m.ARRIENDO_TRIMESTRE, seguro_anual=m.SEGURO_ANUAL,
                 iva_seguro=m.SEGURO_ANUAL * IVA, total_ajustes=arr["devengado"] + seg["devengado"],
                 cif_fijos=C["u9"]["cif_fijos_mes"], sueldos=m.SUELDOS, depreciacion=C["u4"]["total_mes"],
                 appa_antes=arriendo_inicial + m.ARRIENDO_TRIMESTRE, anticipados=appa_final + sppa_final)
    return propio, extra


def cadena(C, propio):
    ok = all(abs(C["u5"][k] - v) < 1e-6 for k, v in propio.items())
    print("CADENA:", "OK, la pieza u5 calza con caso_estampa_sur_contable.py" if ok else f"NO calza: {C['u5']} contra {propio}")
    return ok


# ---------------------------------------------------------------- planilla con fórmulas

def evaluar(wb):
    """Evalúa las fórmulas de la planilla (referencias, + − × ÷ ^ y SUMA) sin Excel."""
    memo = {}

    def valor(hoja, ref):
        clave = (hoja, ref.replace("$", ""))
        if clave in memo:
            return memo[clave]
        v = wb[hoja][clave[1]].value
        if isinstance(v, str) and v.startswith("="):
            v = calcular(hoja, v[1:])
        memo[clave] = v if v is not None else 0
        return memo[clave]

    def rango(hoja, a, b):
        from openpyxl.utils import range_boundaries
        c1, f1, c2, f2 = range_boundaries(f"{a}:{b}".replace("$", ""))
        from openpyxl.utils import get_column_letter as L
        return [valor(hoja, f"{L(c)}{f}") for f in range(f1, f2 + 1) for c in range(c1, c2 + 1)]

    hoja_re = r"(?:'([^']+)'|([A-Za-z_]\w*))!"

    def calcular(hoja, f):
        def sub_rango(mt):
            h = mt.group(1) or mt.group(2) or hoja
            return f"sum(_r({h!r},{mt.group(3)!r},{mt.group(4)!r}))"
        f = re.sub(rf"SUM\((?:{hoja_re})?(\$?[A-Z]+\$?\d+):(\$?[A-Z]+\$?\d+)\)", sub_rango, f)

        def sub_ref(mt):
            h = mt.group(1) or mt.group(2) or hoja
            return f"_v({h!r},{mt.group(3)!r})"
        f = re.sub(rf"(?<![\w'])(?:{hoja_re})?(\$?[A-Z]{{1,2}}\$?\d+)(?![\w(])", sub_ref, f)
        f = f.replace("^", "**")
        return eval(f, {"__builtins__": {}, "sum": sum, "_v": valor, "_r": rango})

    return lambda hoja, ref: valor(hoja, ref)


def planilla(R, propio, extra):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    neg = Font(bold=True)
    formato = '#,##0'

    # Hoja 1 · anticipados: devengado = monto / meses que cubre × meses consumidos
    ws = wb.active
    ws.title = "Anticipados"
    ws.append(["Partida", "Tipo", "Monto", "Meses que cubre", "Meses consumidos", "Por mes", "Devengado", "Saldo al cierre"])
    filas = [("Hostal · seguro (01-08 al 31-12)", "gasto pagado por anticipado", EJ["seguro"]),
             ("Hostal · estacionamiento (01-11 al 31-12)", "gasto pagado por anticipado", EJ["estacionamiento"]),
             ("Hostal · local arrendado (01-09 al 31-12)", "ingreso recibido por anticipado", EJ["local"]),
             ("Estampa Sur · arriendo de marzo", "gasto pagado por anticipado", {"monto": extra["arriendo_inicial"], "cubre": 1, "consumidos": 1}),
             ("Estampa Sur · seguro anual", "gasto pagado por anticipado", {"monto": extra["seguro_anual"], "cubre": 12, "consumidos": 1})]
    for k, (nom, tipo, d) in enumerate(filas, start=2):
        ws.append([nom, tipo, d["monto"], d["cubre"], d["consumidos"], f"=C{k}/D{k}", f"=F{k}*E{k}", f"=C{k}-G{k}"])
    for c in "A1 B1 C1 D1 E1 F1 G1 H1".split():
        ws[c].font = neg
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 42, 32
    for col in "CFGH":
        ws.column_dimensions[col].width = 14
        for k in range(2, 7):
            ws[f"{col}{k}"].number_format = formato

    # Hoja 2 · incobrables del hostal
    wi = wb.create_sheet("Incobrables")
    wi.append(["Clientes al 31-12", EJ["clientes"]])
    wi.append(["Porcentaje estimado incobrable", EJ["pct_incobrable"]])
    wi.append(["Estimación requerida", "=B1*B2"])
    wi.append(["Saldo previo de la estimación", EJ["estimacion_previa"]])
    wi.append(["Ajuste del cierre", "=B3-B4"])
    wi.append(["Clientes, valor neto", "=B1-B3"])
    wi.column_dimensions["A"].width = 34
    for k in (1, 3, 4, 5, 6):
        wi[f"B{k}"].number_format = formato
    wi["B2"].number_format = "0%"

    # Hoja 3 · libro diario de los ajustes del hostal
    wd = wb.create_sheet("Diario")
    wd.append(["Asiento", "Fecha", "Cuenta", "Debe", "Haber", "Glosa"])
    lineas = [
        ("A1", "Gasto por seguro", "=Anticipados!G2", None, "5 meses de seguro"),
        ("A1", "Seguro pagado por anticipado", None, "=Anticipados!G2", ""),
        ("A2", "Gasto por arriendo", "=Anticipados!G3", None, "2 meses de estacionamiento"),
        ("A2", "Arriendo pagado por anticipado", None, "=Anticipados!G3", ""),
        ("A3", "Arriendo recibido por anticipado", "=Anticipados!G4", None, "4 meses del local"),
        ("A3", "Arriendo ganado", None, "=Anticipados!G4", ""),
        ("A4", "Gasto en sueldos", EJ["sueldos_dic"], None, "sueldos de diciembre, se pagan el 05-01"),
        ("A4", "Sueldos por pagar", None, "=D8", ""),
        ("A5", "Gasto en electricidad", EJ["luz_dic"], None, "luz de diciembre, vence el 15-01"),
        ("A5", "IVA crédito fiscal", f"=D10*{IVA}", None, ""),
        ("A5", "Cuentas por pagar", None, "=D10+D11", ""),
        ("A6", "Intereses por cobrar", f"={EJ['deposito']['capital']}*{EJ['deposito']['tasa']}*{EJ['deposito']['meses']}", None, "2 meses del depósito a plazo"),
        ("A6", "Intereses ganados", None, "=D13", ""),
        ("A7", "Gasto por incobrables", "=Incobrables!B5", None, "estimación al 3 % de clientes"),
        ("A7", "Estimación de deudores incobrables", None, "=Incobrables!B5", ""),
    ]
    for n, cta, de, ha, g in lineas:
        wd.append([n, "31-12-2026", cta, de, ha, g])
    fin = 1 + len(lineas)
    wd.append(["", "", "Totales", f"=SUM(D2:D{fin})", f"=SUM(E2:E{fin})", "deben ser iguales"])
    wd.append(["", "", "Diferencia (debe dar 0)", f"=D{fin + 1}-E{fin + 1}", None, ""])
    for c in "A1 B1 C1 D1 E1 F1".split():
        wd[c].font = neg
    wd.column_dimensions["C"].width, wd.column_dimensions["F"].width = 36, 40
    for col in "DE":
        wd.column_dimensions[col].width = 14
        for k in range(2, fin + 3):
            wd[f"{col}{k}"].number_format = formato

    destino = RAMO / "planillas" / "u05_ajustes.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)

    # Comprobación: las fórmulas dan lo mismo que el cálculo en Python
    v = evaluar(wb)
    j = R["ej"]
    pares = [(v("Anticipados", "G2"), j["seguro"]["devengado"]), (v("Anticipados", "H2"), j["seguro"]["saldo"]),
             (v("Anticipados", "G3"), j["est"]["devengado"]), (v("Anticipados", "H3"), j["est"]["saldo"]),
             (v("Anticipados", "G4"), j["local"]["devengado"]), (v("Anticipados", "H4"), j["local"]["saldo"]),
             (v("Anticipados", "G5"), propio["arriendo_devengado"]), (v("Anticipados", "G6"), propio["seguro_devengado"]),
             (v("Anticipados", "H6"), propio["seguro_anticipado_al_31"]),
             (v("Incobrables", "B5"), j["aj_inc"]), (v("Incobrables", "B6"), j["neto_clientes"]),
             (v("Diario", f"D{fin + 1}"), j["total_diario"]), (v("Diario", f"E{fin + 1}"), j["total_diario"]),
             (v("Diario", f"D{fin + 2}"), 0)]
    ok = all(abs(a - b) < 1e-6 for a, b in pares)
    print(f"PLANILLA: {'OK' if ok else 'NO calza'}, {destino.relative_to(RAMO.parents[1])} ({len(pares)} celdas comprobadas)")
    return ok


# ---------------------------------------------------------------- cifras que el .md debe contener

def esperados(R, propio, extra):
    e = []
    j = R["ej"]
    for k in ("seguro", "est", "local"):
        e += [p(j[k]["por_mes"]), p(j[k]["devengado"]), p(j[k]["saldo"])]
    e += [p(j[x]) for x in ("iva_seguro", "banco_seguro", "luz_iva", "luz_total", "interes", "requerida", "aj_inc",
                            "gastos", "ingresos", "resultado", "total_diario", "neto_clientes")]
    e += [p(-j["efecto"])]
    j = R["e1"]; e += [p(j[x]) for x in ("caja_in", "caja_out", "caja", "dev_in", "dev_out", "dev", "por_cobrar", "dif",
                                         "cobros_ajenos", "dif_ingresos", "dif_gastos")]
    j = R["e2"]; e += [p(j[x]) for x in ("por_mes", "devengado", "saldo", "iva", "banco")]
    j = R["e3"]; e += [p(j[x]) for x in ("por_mes", "devengado", "saldo")]
    j = R["e4"]; e += [p(j["gas_iva"]), p(j["gas_total"]), p(-j["efecto"])]
    j = R["e5"]; e += [p(j[x]) for x in ("por_mes", "devengado", "saldo")]
    j = R["e6"]; e += [p(j["liquido"]), p(j["retenciones"]), p(j["gasto"]), fmt(100 * j["pct_liquido"], 1) + " %"]
    j = R["e7"]; e += [p(x) for x in j["req"]] + [p(j[x]) for x in ("total", "aj", "clientes", "neto", "clientes_despues", "estimacion_despues")]
    j = R["e8"]; e += [p(x) for x in j["por"]] + [p(j["total"]), p(j["aj"]), p(j["queda"])]
    j = R["e9"]; e += [p(-j["efecto"]), p(j["activo"]), p(j["pasivo"]), p(j["corregido"]), p(j["bono_inf"]), p(j["bono_ok"]), p(j["bono_de_mas"])]
    j = R["e10"]; e += [p(j[x]) for x in ("gasto_seguro", "ingreso_ara", "sueldos", "incobrables", "intereses", "depreciacion",
                                          "seguro_mes", "ara_mes")] + [p(-j["efecto"])]
    e += [p(v) for v in propio.values()]
    e += [p(extra[x]) for x in ("iva_seguro", "total_ajustes", "cif_fijos", "appa_antes", "anticipados")]
    return e


if __name__ == "__main__":
    R = resultados()
    m, C = cargar_caso()
    propio, extra = caso(m, C)
    for k, d in R.items():
        print(f"== {k}")
        for kk, v in d.items():
            if isinstance(v, dict):
                v = {a: round(b, 2) for a, b in v.items()}
            elif isinstance(v, list):
                v = [round(x, 2) for x in v]
            elif isinstance(v, float):
                v = round(v, 4)
            print(f"   {kk}: {v}")
    print("== caso (propio)", propio)
    print("== caso (extra)", {k: round(v, 2) for k, v in extra.items()})
    print(f"ASIENTOS: {len(ASIENTOS)} registrados, todos cuadran")
    ok_cadena = cadena(C, propio)
    ok_planilla = planilla(R, propio, extra)
    if not (ok_cadena and ok_planilla):
        sys.exit(1)
    if "--revisar" in sys.argv:
        md = (RAMO / "unidades" / "05-ajustes-de-cierre.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %")
        faltan = [x for x in esperados(R, propio, extra) if x not in md]
        print("REVISIÓN:", "OK, todas las cifras están en el texto" if not faltan else f"faltan {len(faltan)}: {faltan}")
        sys.exit(1 if faltan else 0)
