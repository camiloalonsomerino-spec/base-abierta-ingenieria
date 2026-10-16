"""Unidad 9 · Costos: conceptos y clasificación.

Recalcula todas las cifras de la unidad desde los datos, genera la planilla
`planillas/u09_clasificacion_costos.xlsx` (con fórmulas, no valores pegados) y comprueba sus
resultados evaluando esas fórmulas en Python. Con --revisar busca cada resultado en el .md.

    py -3 -P u09_clasificacion_costos.py            # imprime los resultados y arma la planilla
    py -3 -P u09_clasificacion_costos.py --revisar  # además los busca en el .md

Todos los datos son originales de esta base (pasados por herramientas/colisiones.py). El caso es
el de caso_estampa_sur_contable.py, que se importa sin tocarlo.
"""
from pathlib import Path
import importlib.util, re, sys

sys.stdout.reconfigure(encoding="utf-8")
AQUI = Path(__file__).resolve().parent


# ---------------------------------------------------------------- herramientas

def fmt(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def recta(pts):
    """Mínimos cuadrados: costo = a + b · actividad."""
    n = len(pts)
    mx = sum(x for x, _ in pts) / n
    my = sum(y for _, y in pts) / n
    b = sum((x - mx) * (y - my) for x, y in pts) / sum((x - mx) ** 2 for x, _ in pts)
    return my - b * mx, b


def alto_bajo(pts):
    """Punto alto y punto bajo según la ACTIVIDAD (no según el costo)."""
    alto = max(pts, key=lambda p: p[0])
    bajo = min(pts, key=lambda p: p[0])
    b = (alto[1] - bajo[1]) / (alto[0] - bajo[0])
    return alto[1] - b * alto[0], b, alto, bajo


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · taller de tazas de cerámica, mayo
EJ = dict(mp_ini=176_500, compras=1_341_900, mp_fin=239_980,
          mod=662 * 2_860, insumos=17 * 2_860, supervisora=684_000, gas=297_300,
          dep_horno=118_500, arriendo_taller=532_000, mantencion=36_400,
          pp_ini=206_300, pp_fin=204_260, terminadas=2_860,
          pt_ini_u=230, vendidas=2_470, precio=3_390, comision=0.04,
          sueldo_admin=615_000, publicidad=174_000, arriendo_tienda=163_000,
          # gas del horno: (tazas, costo) de enero a mayo
          gas_meses=[("enero", 2_140, 253_700), ("febrero", 1_730, 227_340), ("marzo", 2_580, 279_900),
                     ("abril", 3_310, 318_980), ("mayo", 2_860, 297_300)],
          # estándar variable por taza y pedido especial
          std=dict(arcilla=447, mano_obra=662, insumos=17, gas=58),
          pedido=dict(q=380, precio=1_560, logo=115), capacidad=3_450)

# 9.1 · fábrica de mermeladas: (partida, monto, elemento, comportamiento)
E1 = [("Fruta", 1_860_000, "MPD", "variable"), ("Azúcar", 420_000, "MPD", "variable"),
      ("Frascos y tapas", 538_000, "MPD", "variable"), ("Pectina", 74_000, "CIF", "variable"),
      ("Sueldo de las tres operarias de cocina", 1_950_000, "MOD", "fijo"),
      ("Sueldo de la jefa de producción", 980_000, "CIF", "fijo"),
      ("Gas de las ollas", 212_000, "CIF", "variable"),
      ("Depreciación de ollas y envasadora", 145_000, "CIF", "fijo"),
      ("Arriendo de la planta", 700_000, "CIF", "fijo"),
      ("Detergente y sanitizante de la planta", 58_000, "CIF", "variable"),
      ("Comisión de la vendedora", 336_000, "período", "variable"),
      ("Contador externo", 250_000, "período", "fijo"),
      ("Despacho a clientes", 184_000, "período", "variable")]
# 9.2 · cervecería artesanal
E2 = dict(compras=2_140_000, mp_ini=310_000, mp_fin=265_000, mod=1_480_000, cif=1_395_000,
          pp_ini=420_000, pp_fin=515_000, pt_ini=690_000, pt_fin=540_000, ventas=9_300_000, gav=2_240_000)
# 9.3 · cocina de almuerzos
E3 = dict(fijo=1_200_000, var=1_850, niveles=[800, 1_600, 2_400])
# 9.4 · plotter: hundido y oportunidad
E4 = dict(plotter=2_600_000, vida=48, q=40, material=4_800, tinta=1_300, precio=11_500, alternativa=95_000)
# 9.5 · energía de una lavandería: (mes, kg, costo)
E5 = dict(meses=[("enero", 8_400, 1_052_000), ("febrero", 7_100, 948_000), ("marzo", 9_800, 1_146_000),
                 ("abril", 6_300, 891_200), ("mayo", 10_600, 1_209_400), ("junio", 8_900, 1_088_000)],
          pronostico=9_000)
# 9.6 · mueblería
E6 = dict(mp_ini=640_000, compras=3_920_000, mp_fin=710_000, mpi=330_000, mod=2_760_000, moi=1_150_000,
          otros_cif=1_290_000, pp_ini=380_000, pp_fin=530_000, terminadas=200, pt_ini_u=20, vendidas=190,
          precio=72_000, gav=2_380_000)
# 9.7 · pastelería, pedido especial con capacidad ociosa
E7 = dict(q_normal=1_730, capacidad=2_260, ingredientes=4_270, caja=640, mano_obra=1_090, energia=335,
          fijos=5_882_000, precio=14_790, pedido=260, precio_pedido=8_870, decoracion=385)
# 9.8 · hacer o comprar
E8 = dict(q=1_180, material=5_230, mano_obra=2_270, cif_var=910, cif_fijos=3_658_000, evitables=1_416_000,
          proveedor=10_380, arriendo_espacio=1_090_000)
# 9.9 · pedido especial sin capacidad ociosa
E9 = dict(q_normal=2_140, capacidad=2_330, pedido=470, precio_pedido=9_460, precio=12_180, cv=7_390, hilo=290)
# 9.10 · mantención de un frigorífico: (mes, horas máquina, costo); el mes 5 tuvo una avería
E10 = dict(meses=[(1, 1_090, 1_268_000), (2, 1_240, 1_321_500), (3, 1_410, 1_402_800), (4, 1_170, 1_293_900),
                  (5, 1_630, 1_947_000), (6, 1_330, 1_369_200), (7, 1_470, 1_417_600), (8, 1_060, 1_249_300)],
           atipico=5, presupuesto=1_500)


# ---------------------------------------------------------------- resultados

def resultados():
    R = {}
    e = EJ
    mpd = e["mp_ini"] + e["compras"] - e["mp_fin"]
    cif = e["insumos"] + e["supervisora"] + e["gas"] + e["dep_horno"] + e["arriendo_taller"] + e["mantencion"]
    fab = mpd + e["mod"] + cif
    prod = fab + e["pp_ini"] - e["pp_fin"]
    cu = prod / e["terminadas"]
    pt_ini = e["pt_ini_u"] * cu
    pt_fin_u = e["pt_ini_u"] + e["terminadas"] - e["vendidas"]
    pt_fin = pt_fin_u * cu
    cventas = prod + pt_ini - pt_fin
    ventas = e["vendidas"] * e["precio"]
    comision = ventas * e["comision"]
    gav = comision + e["sueldo_admin"] + e["publicidad"] + e["arriendo_tienda"]
    pts = [(x, y) for _, x, y in e["gas_meses"]]
    a, b, alto, bajo = alto_bajo(pts)
    ar, br = recta(pts)
    s = e["std"]
    cv_taza = sum(s.values())
    p = e["pedido"]
    cv_ped = cv_taza + p["logo"]
    R["ej"] = dict(mpd=mpd, cif=cif, fab=fab, prod=prod, cu=cu, pt_ini=pt_ini, pt_fin_u=pt_fin_u, pt_fin=pt_fin,
                   cventas=cventas, ventas=ventas, comision=comision, gav=gav, ro=ventas - cventas - gav,
                   primo=mpd + e["mod"], conversion=e["mod"] + cif, mod=e["mod"], insumos=e["insumos"],
                   margen_bruto=ventas - cventas,
                   gas_a=a, gas_b=b, gas_alto=alto, gas_bajo=bajo, gas_mayo_pred=a + b * 2_860,
                   gas_mayo_dif=e["gas"] - (a + b * 2_860), gas_var_mayo=b * 2_860,
                   reg_a=ar, reg_b=br, cv_taza=cv_taza, cv_ped=cv_ped, contrib_u=p["precio"] - cv_ped,
                   contrib=(p["precio"] - cv_ped) * p["q"], completo=cu + p["logo"],
                   perdida_aparente=(p["precio"] - cu - p["logo"]) * p["q"],
                   holgura=e["capacidad"] - e["terminadas"])
    assert abs(mpd - 447 * 2_860) < 1e-9 and abs(cventas - e["vendidas"] * cu) < 1e-6

    tot = lambda clase: sum(m for _, m, c, _ in E1 if c == clase)
    mpd1, mod1, cif1, per1 = tot("MPD"), tot("MOD"), tot("CIF"), tot("período")
    fijos1 = sum(m for _, m, c, k in E1 if c != "período" and k == "fijo")
    var1 = sum(m for _, m, c, k in E1 if c != "período" and k == "variable")
    R["e1"] = dict(mpd=mpd1, mod=mod1, cif=cif1, periodo=per1, fab=mpd1 + mod1 + cif1, primo=mpd1 + mod1,
                   conversion=mod1 + cif1, fijos=fijos1, variables=var1)
    assert fijos1 + var1 == mpd1 + mod1 + cif1

    e = E2
    mp = e["mp_ini"] + e["compras"] - e["mp_fin"]
    fab = mp + e["mod"] + e["cif"]
    prod = fab + e["pp_ini"] - e["pp_fin"]
    cv = prod + e["pt_ini"] - e["pt_fin"]
    R["e2"] = dict(mp=mp, primo=mp + e["mod"], conversion=e["mod"] + e["cif"], fab=fab, prod=prod, cventas=cv,
                   mb=e["ventas"] - cv, ro=e["ventas"] - cv - e["gav"])

    e = E3
    R["e3"] = [dict(q=q, total=e["fijo"] + e["var"] * q, fijo_u=e["fijo"] / q, total_u=e["fijo"] / q + e["var"])
               for q in e["niveles"]]

    e = E4
    cv4 = e["material"] + e["tinta"]
    R["e4"] = dict(dep=e["plotter"] / e["vida"], cv=cv4, contrib_u=e["precio"] - cv4,
                   contrib=(e["precio"] - cv4) * e["q"], ventaja=(e["precio"] - cv4) * e["q"] - e["alternativa"],
                   ingreso=e["precio"] * e["q"])

    pts = [(x, y) for _, x, y in E5["meses"]]
    a, b, alto, bajo = alto_bajo(pts)
    ar, br = recta(pts)
    q = E5["pronostico"]
    R["e5"] = dict(a=a, b=b, alto=alto, bajo=bajo, pred=a + b * q, reg_a=ar, reg_b=br, reg_pred=ar + br * q)

    e = E6
    mp = e["mp_ini"] + e["compras"] - e["mp_fin"]
    mpd = mp - e["mpi"]
    cif = e["mpi"] + e["moi"] + e["otros_cif"]
    fab = mpd + e["mod"] + cif
    prod = fab + e["pp_ini"] - e["pp_fin"]
    cu = prod / e["terminadas"]
    pt_fin_u = e["pt_ini_u"] + e["terminadas"] - e["vendidas"]
    cvent = prod + e["pt_ini_u"] * cu - pt_fin_u * cu
    ventas = e["vendidas"] * e["precio"]
    R["e6"] = dict(mp=mp, mpd=mpd, cif=cif, fab=fab, prod=prod, cu=cu, pt_ini=e["pt_ini_u"] * cu,
                   pt_fin_u=pt_fin_u, pt_fin=pt_fin_u * cu, cventas=cvent, ventas=ventas, mb=ventas - cvent,
                   ro=ventas - cvent - e["gav"], primo=mpd + e["mod"], conversion=e["mod"] + cif)

    e = E7
    cv7 = e["ingredientes"] + e["caja"] + e["mano_obra"] + e["energia"]
    fu = e["fijos"] / e["q_normal"]
    cvp = cv7 + e["decoracion"]
    R["e7"] = dict(cv=cv7, fijo_u=fu, completo=cv7 + fu, cv_ped=cvp, contrib_u=e["precio_pedido"] - cvp,
                   contrib=(e["precio_pedido"] - cvp) * e["pedido"], holgura=e["capacidad"] - e["q_normal"],
                   aparente=(e["precio_pedido"] - cv7 - fu) * e["pedido"], margen_normal=e["precio"] - cv7)

    e = E8
    cv8 = e["material"] + e["mano_obra"] + e["cif_var"]
    hacer = cv8 * e["q"] + e["evitables"]
    comprar = e["proveedor"] * e["q"]
    R["e8"] = dict(cv=cv8, fijo_u=e["cif_fijos"] / e["q"], completo=cv8 + e["cif_fijos"] / e["q"],
                   hacer=hacer, hacer_u=hacer / e["q"], comprar=comprar, ventaja_hacer=comprar - hacer,
                   comprar_neto=comprar - e["arriendo_espacio"],
                   ventaja_comprar=hacer - (comprar - e["arriendo_espacio"]),
                   inevitables=e["cif_fijos"] - e["evitables"], cv_total=cv8 * e["q"])

    e = E9
    cvp = e["cv"] + e["hilo"]
    holg = e["capacidad"] - e["q_normal"]
    desplaza = e["pedido"] - holg
    mc_normal = e["precio"] - e["cv"]
    contrib = (e["precio_pedido"] - cvp) * e["pedido"]
    perdida = desplaza * mc_normal
    R["e9"] = dict(cv_ped=cvp, holgura=holg, desplaza=desplaza, mc_normal=mc_normal, contrib_u=e["precio_pedido"] - cvp,
                   contrib=contrib, oportunidad=perdida, neto=contrib - perdida,
                   minimo=cvp + perdida / e["pedido"], parcial=holg * (e["precio_pedido"] - cvp))

    pts = [(x, y) for _, x, y in E10["meses"]]
    a, b, alto, bajo = alto_bajo(pts)
    ar, br = recta(pts)
    limpios = [(x, y) for m, x, y in E10["meses"] if m != E10["atipico"]]
    al2, bl2, alto2, bajo2 = alto_bajo(limpios)
    ac, bc = recta(limpios)
    h = E10["presupuesto"]
    R["e10"] = dict(a=a, b=b, alto=alto, bajo=bajo, pred=a + b * h, reg_a=ar, reg_b=br, reg_pred=ar + br * h,
                    lim_a=ac, lim_b=bc, lim_pred=ac + bc * h, ab2_a=al2, ab2_b=bl2, ab2_pred=al2 + bl2 * h,
                    alto2=alto2)
    return R


# ---------------------------------------------------------------- caso

def cargar_caso():
    spec = importlib.util.spec_from_file_location("caso", AQUI / "caso_estampa_sur_contable.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m, m.calcular()


def caso(m, C):
    """Recálculo propio de la pieza u9, desde los datos declarados del caso."""
    cv = [p - mg for p, mg in zip(m.PRECIO, m.MARGEN)]
    tinta = [g * m.TINTA_KG / 1000 for g in m.TINTA_G]
    otros = [c - pl - t for c, pl, t in zip(cv, m.PRENDA_LISA, tinta)]
    dep = sum((q["costo"] - q["residual"]) / q["vida_meses"] for q in m.MAQUINAS.values())
    seguro = m.SEGURO_ANUAL / 12
    arriendo = m.BALANCE_INICIAL["activo"]["Arriendo pagado por anticipado"]
    cif = m.SUELDOS + dep + seguro + arriendo
    P = m.PRODUCIDAS
    lisas = sum(q * c for q, c in zip(P, m.PRENDA_LISA))
    tinta_tot = sum(q * t for q, t in zip(P, tinta))
    otros_tot = sum(q * o for q, o in zip(P, otros))
    var_fab = sum(q * c for q, c in zip(P, cv))
    assert abs(lisas + tinta_tot + otros_tot - var_fab) < 1e-6
    return dict(cv=cv, tinta=tinta, otros=otros, dep=dep, seguro=seguro, arriendo=arriendo, cif=cif,
                sueldos=m.SUELDOS, admin=m.OTROS_FIJOS, interes=C["u6"]["interes_mes_1"],
                lisas=lisas, tinta_tot=tinta_tot, tinta_kg=sum(q * g for q, g in zip(P, m.TINTA_G)) / 1000,
                otros_tot=otros_tot, var_fab=var_fab, producido=var_fab + cif, unidades=sum(P),
                precio=m.PRECIO, margen=m.MARGEN, lisa=m.PRENDA_LISA, gramos=m.TINTA_G, producidas=P,
                maquina=m.MAQUINA["neto"], arriendo_trim=m.ARRIENDO_TRIMESTRE, seguro_anual=m.SEGURO_ANUAL,
                dep_pulpo=C["u4"]["depreciacion_mensual"]["pulpo 6 (inicial)"],
                dep_segunda=C["u4"]["depreciacion_mensual"]["segunda máquina"])


# ---------------------------------------------------------------- planilla (con fórmulas) y su evaluación

def evaluar(ws):
    """Evalúa las fórmulas de una hoja (+ − × ÷, SUM, SUMPRODUCT, SLOPE, INTERCEPT, MAX, MIN)."""
    from openpyxl.utils import range_boundaries, get_column_letter
    memo = {}

    def v(c):
        if c not in memo:
            x = ws[c].value
            memo[c] = calc(x[1:]) if isinstance(x, str) and x.startswith("=") else (x or 0)
        return memo[c]

    def r(a, b):
        c1, r1, c2, r2 = range_boundaries(f"{a}:{b}")
        return [v(f"{get_column_letter(c)}{f}") for f in range(r1, r2 + 1) for c in range(c1, c2 + 1)]

    def slope(ys, xs):
        return recta(list(zip(xs, ys)))[1]

    def icpt(ys, xs):
        return recta(list(zip(xs, ys)))[0]

    env = dict(_v=v, _r=r, SUM=lambda *a: sum(sum(x) if isinstance(x, list) else x for x in a),
               SUMPRODUCT=lambda a, b: sum(x * y for x, y in zip(a, b)), SLOPE=slope, INTERCEPT=icpt,
               MAX=lambda a: max(a), MIN=lambda a: min(a))

    def calc(expr):
        e = re.sub(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)", r'_r("\1\2","\3\4")', expr)
        e = re.sub(r'(?<![A-Z"_])\$?([A-Z]{1,3})\$?(\d+)(?![\d"])', r'_v("\1\2")', e)
        return eval(e, {"__builtins__": {}}, env)

    return v


def planilla(R, K):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    neg = Font(bold=True)
    num = "#,##0"

    # Hoja 1 · ejemplo
    ws = wb.active
    ws.title = "Ejemplo tazas"
    e = EJ
    filas = [("Inventario inicial de materias primas", e["mp_ini"]), ("Compras de arcilla y esmalte (netas)", e["compras"]),
             ("Inventario final de materias primas", e["mp_fin"]),
             ("Materia prima directa usada", "=B2+B3-B4"),
             ("Mano de obra directa (a trato)", e["mod"]),
             ("Insumos de taller (MPI)", e["insumos"]), ("Supervisora (MOI)", e["supervisora"]),
             ("Gas del horno", e["gas"]), ("Depreciación del horno", e["dep_horno"]),
             ("Arriendo del taller", e["arriendo_taller"]), ("Mantención", e["mantencion"]),
             ("Costos indirectos de fabricación", "=SUM(B7:B12)"),
             ("Costo de fabricación", "=B5+B6+B13"),
             ("Productos en proceso iniciales", e["pp_ini"]), ("Productos en proceso finales", e["pp_fin"]),
             ("Costo de lo producido", "=B14+B15-B16"),
             ("Tazas terminadas", e["terminadas"]), ("Costo por taza terminada", "=B17/B18"),
             ("Tazas en bodega al inicio", e["pt_ini_u"]), ("Productos terminados iniciales", "=B20*B19"),
             ("Tazas vendidas", e["vendidas"]), ("Tazas en bodega al final", "=B20+B18-B22"),
             ("Productos terminados finales", "=B23*B19"),
             ("Costo de ventas", "=B17+B21-B24"),
             ("Precio de venta", e["precio"]), ("Ventas", "=B22*B26"), ("Margen bruto", "=B27-B25"),
             ("Comisión de venta (4 %)", "=0.04*B27"), ("Sueldo administrativa", e["sueldo_admin"]),
             ("Publicidad", e["publicidad"]), ("Arriendo de la tienda", e["arriendo_tienda"]),
             ("Gastos de administración y ventas", "=SUM(B29:B32)"),
             ("Resultado operacional", "=B28-B33"),
             ("Costo primo", "=B5+B6"), ("Costo de conversión", "=B6+B13")]
    ws["A1"], ws["B1"] = "Partida", "Pesos"
    for i, (t, x) in enumerate(filas, start=2):
        ws.cell(i, 1, t)
        ws.cell(i, 2, x).number_format = num
    for c in ("A1", "B1", "A14", "A17", "A25", "A34"):
        ws[c].font = neg
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14

    # Hoja 2 · gas del horno
    wg = wb.create_sheet("Gas del horno")
    wg.append(["Mes", "Tazas", "Costo del gas"])
    for mes, x, y in e["gas_meses"]:
        wg.append([mes, x, y])
    wg["E1"], wg["F1"] = "Punto alto y punto bajo", ""
    wg["E2"], wg["F2"] = "Variable por taza (b)", "=(C5-C3)/(B5-B3)"
    wg["E3"], wg["F3"] = "Fijo del mes (a)", "=C5-F2*B5"
    wg["E4"], wg["F4"] = "Regresión: pendiente", "=SLOPE(C2:C6,B2:B6)"
    wg["E5"], wg["F5"] = "Regresión: intercepto", "=INTERCEPT(C2:C6,B2:B6)"
    wg["E6"], wg["F6"] = "Gas esperado en mayo (alto-bajo)", "=F3+F2*B6"
    wg.column_dimensions["E"].width = 34

    # Hoja 3 · caso
    wc = wb.create_sheet("Caso Estampa Sur")
    wc.append(["Producto", "Precio", "Margen", "Costo variable", "Prenda lisa", "Tinta (g)", "Tinta ($)",
               "Otros variables", "Producidas", "Costo variable del mes"])
    nombres = ["Poleras básicas", "Polerones", "Poleras estampado doble", "Bolsas", "Gorros"]
    for i, n in enumerate(nombres):
        f = i + 2
        wc.append([n, K["precio"][i], K["margen"][i], f"=B{f}-C{f}", K["lisa"][i], K["gramos"][i],
                   f"=F{f}*16500/1000", f"=D{f}-E{f}-G{f}", K["producidas"][i], f"=I{f}*D{f}"])
    wc["L1"], wc["M1"] = "CIF fijos de marzo", ""
    wc["L2"], wc["M2"] = "Sueldos", K["sueldos"]
    wc["L3"], wc["M3"] = "Depreciación de las dos máquinas", "=9840000/96+(14300000-1800000)/60"
    wc["L4"], wc["M4"] = "Seguro (un mes)", "=948000/12"
    wc["L5"], wc["M5"] = "Arriendo de marzo", K["arriendo"]
    wc["L6"], wc["M6"] = "Total CIF fijos", "=SUM(M2:M5)"
    wc["L7"], wc["M7"] = "Costo variable de fabricación", "=SUM(J2:J6)"
    wc["L8"], wc["M8"] = "Costo de lo producido en marzo", "=M6+M7"
    wc["L9"], wc["M9"] = "Gastos de administración (período)", K["admin"]
    wc["L10"], wc["M10"] = "Prendas lisas consumidas", "=SUMPRODUCT(E2:E6,I2:I6)"
    wc["L11"], wc["M11"] = "Tinta consumida", "=SUMPRODUCT(G2:G6,I2:I6)"
    wc["L12"], wc["M12"] = "Otros variables consumidos", "=SUMPRODUCT(H2:H6,I2:I6)"
    wc.column_dimensions["A"].width = 24
    wc.column_dimensions["L"].width = 36
    for row in wc.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("=")):
                c.number_format = "#,##0.0"

    destino = AQUI.parent / "planillas" / "u09_clasificacion_costos.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)

    # comprobación: las fórmulas evaluadas dan lo mismo que Python
    v1, v2, v3 = evaluar(ws), evaluar(wg), evaluar(wc)
    j = R["ej"]
    pares = [(v1("B5"), j["mpd"]), (v1("B13"), j["cif"]), (v1("B14"), j["fab"]), (v1("B17"), j["prod"]),
             (v1("B19"), j["cu"]), (v1("B25"), j["cventas"]), (v1("B34"), j["ro"]), (v1("B35"), j["primo"]),
             (v1("B36"), j["conversion"]), (v2("F2"), j["gas_b"]), (v2("F3"), j["gas_a"]),
             (v2("F4"), j["reg_b"]), (v2("F5"), j["reg_a"]), (v2("F6"), j["gas_mayo_pred"]),
             (v3("M6"), K["cif"]), (v3("M7"), K["var_fab"]), (v3("M8"), K["producido"]),
             (v3("M10"), K["lisas"]), (v3("M11"), K["tinta_tot"]), (v3("M12"), K["otros_tot"])]
    pares += [(v3(f"D{i + 2}"), K["cv"][i]) for i in range(5)] + [(v3(f"H{i + 2}"), K["otros"][i]) for i in range(5)]
    ok = all(abs(a - b) < 1e-6 for a, b in pares)
    print(f"PLANILLA: {'OK' if ok else 'NO calza'}, {len(pares)} celdas con fórmula comprobadas · {destino.name}")
    return ok


# ---------------------------------------------------------------- textos esperados

def esperados(R, K):
    p = fmt
    e = []
    j = R["ej"]
    e += [p(j[k]) for k in ("mpd", "cif", "fab", "prod", "cu", "pt_ini", "pt_fin", "cventas", "ventas", "comision",
                            "gav", "ro", "primo", "conversion", "mod", "insumos", "margen_bruto", "gas_a",
                            "gas_mayo_pred", "gas_mayo_dif", "gas_var_mayo", "cv_taza", "cv_ped", "contrib_u",
                            "contrib", "completo", "holgura")]
    e += [p(-j["perdida_aparente"]), str(j["pt_fin_u"]), p(j["gas_b"]), fmt(j["reg_b"], 2), p(j["reg_a"])]
    j = R["e1"]; e += [p(x) for x in j.values()]
    j = R["e2"]; e += [p(x) for x in j.values()]
    for f in R["e3"]:
        e += [p(f["total"]), p(f["fijo_u"]), p(f["total_u"])]
    j = R["e4"]; e += [p(j["dep"]), p(j["cv"]), p(j["contrib_u"]), p(j["contrib"]), p(j["ventaja"]), p(j["ingreso"])]
    j = R["e5"]; e += [p(j["a"]), p(j["b"]), p(j["pred"]), fmt(j["reg_b"], 2), p(j["reg_a"]), p(j["reg_pred"])]
    j = R["e6"]; e += [p(j[k]) for k in ("mp", "mpd", "cif", "fab", "prod", "cu", "pt_ini", "pt_fin", "cventas",
                                         "ventas", "mb", "ro", "primo", "conversion")] + [str(j["pt_fin_u"])]
    j = R["e7"]; e += [p(j[k]) for k in ("cv", "fijo_u", "completo", "cv_ped", "contrib_u", "contrib", "holgura",
                                         "margen_normal")] + [p(-j["aparente"])]
    j = R["e8"]; e += [p(j[k]) for k in ("cv", "fijo_u", "completo", "hacer", "hacer_u", "comprar", "ventaja_hacer",
                                         "comprar_neto", "ventaja_comprar", "inevitables", "cv_total")]
    j = R["e9"]; e += [p(j[k]) for k in ("cv_ped", "holgura", "desplaza", "mc_normal", "contrib_u", "contrib",
                                         "oportunidad", "parcial")] + [p(-j["neto"]), fmt(j["minimo"], 0)]
    j = R["e10"]; e += [fmt(j["b"], 2), p(-j["a"]), p(j["pred"]), fmt(j["reg_b"], 2), p(j["reg_a"]), p(j["reg_pred"]),
                        fmt(j["lim_b"], 2), p(j["lim_a"]), p(j["lim_pred"]), fmt(j["ab2_b"], 2), p(j["ab2_a"]),
                        p(j["ab2_pred"])]
    # caso
    e += [p(x) for x in K["cv"]] + [fmt(x, 1) for x in K["tinta"]] + [fmt(x, 1) for x in K["otros"]]
    e += [p(K[k]) for k in ("dep", "seguro", "arriendo", "cif", "sueldos", "admin", "interes", "lisas", "var_fab",
                            "producido", "unidades", "maquina", "arriendo_trim", "seguro_anual", "dep_pulpo",
                            "dep_segunda")]
    e += [fmt(K["cif"], 2), fmt(K["tinta_tot"], 1), fmt(K["otros_tot"], 1), fmt(K["tinta_kg"], 3)]
    e += [p(pr - c) for pr, c in zip(K["precio"], K["cv"])]
    return e


def imprimir(R, K):
    for k, d in R.items():
        print(f"== {k}")
        if isinstance(d, list):
            for f in d:
                print("  ", {a: round(b, 4) for a, b in f.items()})
            continue
        for kk, v in d.items():
            print(f"   {kk}: {round(v, 4) if isinstance(v, float) else v}")
    print("== caso (propio)")
    for kk, v in K.items():
        if isinstance(v, list):
            v = [round(x, 4) for x in v]
        elif isinstance(v, float):
            v = round(v, 4)
        print(f"   {kk}: {v}")


if __name__ == "__main__":
    R = resultados()
    m, C = cargar_caso()
    K = caso(m, C)
    imprimir(R, K)
    u9 = C["u9"]
    ok = (all(abs(a - b) < 1e-9 for a, b in zip(K["cv"], u9["costo_variable"]))
          and all(abs(a - b) < 1e-9 for a, b in zip(K["otros"], u9["otros_variables"]))
          and abs(K["cif"] - u9["cif_fijos_mes"]) < 1e-6 and K["admin"] == u9["gastos_admin"]
          and abs(K["interes"] - u9["gasto_financiero"]) < 1e-9)
    print("CADENA:", "OK, la pieza u9 calza con caso_estampa_sur_contable.py" if ok else "NO calza")
    if not ok:
        sys.exit(1)
    if not planilla(R, K):
        sys.exit(1)
    if "--revisar" in sys.argv:
        md = (AQUI.parent / "unidades" / "09-clasificacion-de-costos.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %")
        faltan = [x for x in esperados(R, K) if x not in md]
        print("REVISIÓN:", "OK, todas las cifras están en el texto" if not faltan else f"faltan {len(faltan)}: {faltan}")
        sys.exit(1 if faltan else 0)
