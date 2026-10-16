"""Unidad 6 · Préstamos e intereses en la contabilidad.

Recalcula todas las cifras de la unidad desde los datos, comprueba que cada asiento cuadre y que cada
tabla de amortización termine en cero (y calce con la fórmula del saldo), genera la planilla
`planillas/u06_prestamos.xlsx` (con fórmulas) y evalúa esas fórmulas para comprobar que dan lo mismo
que el cálculo en Python. Con --revisar busca cada resultado, y cada fila de las dos tablas largas,
en el .md.

    py -3 -P u06_prestamos.py            # imprime los resultados y genera la planilla
    py -3 -P u06_prestamos.py --tablas   # además imprime las tablas en Markdown
    py -3 -P u06_prestamos.py --revisar  # además los busca en el .md

Todos los datos son originales de esta base (pasados por herramientas/colisiones.py). Las tasas no
describen el mercado chileno de ninguna fecha. La única cifra legal que se usa es el IVA (19 %).
"""
from pathlib import Path
import importlib.util, re, sys

sys.stdout.reconfigure(encoding="utf-8")
IVA = 0.19
AQUI = Path(__file__).resolve().parent
RAMO = AQUI.parent


# ---------------------------------------------------------------- herramientas

def fmt(x, dec=0):
    if abs(x) < 0.5 * 10 ** -dec:
        x = 0.0
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def p(x):
    return fmt(x, 0)


def p2(x):
    return fmt(x, 2)


def pct(x, dec=2):
    return fmt(100 * x, dec) + " %"


ASIENTOS = []


def asiento(nombre, fecha, debe, haber):
    d, h = sum(debe.values()), sum(haber.values())
    assert abs(d - h) < 1e-6, f"{nombre} no cuadra: {d} contra {h}"
    ASIENTOS.append(dict(nombre=nombre, fecha=fecha, debe=debe, haber=haber, total=d))
    return d


def factor(i, n):
    return (1 - (1 + i) ** -n) / i


def cuota(P, i, n):
    return P / factor(i, n)


def tabla(P, i, n):
    """Tabla francesa: fila k con saldo inicial, interés, amortización, cuota y saldo final."""
    c, s, t = cuota(P, i, n), P, []
    for k in range(1, n + 1):
        it = s * i
        a = c - it
        t.append(dict(k=k, ini=s, interes=it, amort=a, cuota=c, fin=s - a))
        s -= a
    assert abs(t[-1]["fin"]) < 1e-6, "la tabla no termina en cero"
    for k in (1, n // 2, n - 1):   # segundo camino: saldo por fórmula = saldo de la tabla
        assert abs(t[k - 1]["fin"] - c * factor(i, n - k)) < 1e-6
    return c, t


def saldo(t, k):
    """Saldo después de pagar k cuotas (k = 0 es el monto prestado)."""
    return t[0]["ini"] if k == 0 else t[k - 1]["fin"]


def corriente(t, pagadas):
    """Capital que se paga en las 12 cuotas siguientes y el que queda para después."""
    hasta = min(pagadas + 12, len(t))
    return saldo(t, pagadas) - saldo(t, hasta), saldo(t, hasta)


def tabla_md(t, dec):
    f = p2 if dec == 2 else p
    return [f"| {r['k']} | {f(r['ini'])} | {f(r['interes'])} | {f(r['amort'])} | {f(r['cuota'])} | {f(r['fin'])} |" for r in t]


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · panadería Masa Madre: recibe el 01-10-2026, cuotas el día 1 desde el 01-11-2026
EJ = {"P": 7_350_000, "i": 0.0135, "n": 18, "pagadas_al_cierre": 2}
# 6.1 y 6.6 · préstamo para una cámara de frío
E1 = {"P": 4_650_000, "i": 0.0145, "n": 15}
# 6.2 · el error de llevar la cuota entera a gasto (con E1, cuotas 1 a 3)
E2 = {"cuotas": 3}
# 6.3 · recibido el 15-01-2026, cuotas el día 15 desde el 15-02-2026, cierre semestral al 30-06-2026
E3 = {"P": 5_200_000, "i": 0.0125, "n": 24, "pagadas": 5, "dias": 15, "dias_mes": 30}
# 6.4 · camioneta comprada el 01-03-2026 con pie y dos letras
E4 = {"neto": 18_900_000, "pie": 6_900_000, "letra": 6_000_000, "letras_meses": (10, 20)}
# 6.5 · recibido el 01-01-2026, cuotas a fin de cada mes desde el 31-01-2026, cierre al 31-12-2026
E5 = {"P": 12_600_000, "i": 0.011, "n": 36, "pagadas": 12}
# 6.7 · pagaré con interés simple al vencimiento: firmado el 01-08-2026, vence el 01-05-2027
E7 = {"P": 9_600_000, "i": 0.0085, "meses": 9, "meses_2026": 5}
# 6.8 · sistema alemán
E8 = {"P": 6_240_000, "i": 0.015, "n": 8}
# 6.9 · cuotas anuales: recibido el 01-07-2026, primera cuota el 01-07-2027, cierre al 31-12
E9 = {"P": 18_500_000, "i": 0.09, "n": 4}
# 6.10 · la porción corriente y el compromiso de razón corriente con el banco
E10 = {"P": 15_000_000, "i": 0.0105, "n": 48, "pagadas": 12, "ac": 18_400_000, "pc_sin": 9_300_000,
       "minimo": 1.5, "r_operacional": 4_100_000}


# ---------------------------------------------------------------- resultados

def resultados():
    R = {}
    e = EJ
    c, t = tabla(e["P"], e["i"], e["n"])
    k = e["pagadas_al_cierre"]
    cor, nocor = corriente(t, k)
    i3 = t[k]["interes"]
    asiento("ej-recibe", "01-10-2026", {"Banco": e["P"]}, {"Préstamo bancario": e["P"]})
    for r in t[:k]:
        asiento(f"ej-cuota{r['k']}", "", {"Préstamo bancario": r["amort"], "Gasto financiero": r["interes"]}, {"Banco": c})
    asiento("ej-ajuste", "31-12-2026", {"Gasto financiero": i3}, {"Intereses por pagar": i3})
    asiento("ej-cuota3", "01-01-2027", {"Préstamo bancario": t[k]["amort"], "Intereses por pagar": i3}, {"Banco": c})
    gasto_2026 = sum(r["interes"] for r in t[:k + 1])
    total_int = c * e["n"] - e["P"]
    R["ej"] = dict(cuota=c, t=t, interes_total=total_int, total_pagado=c * e["n"], i3=i3, gasto_2026=gasto_2026,
                   saldo_cierre=saldo(t, k), corriente=cor, no_corriente=nocor, pasivo_total=saldo(t, k) + i3,
                   doce_cuotas=12 * c, error_corriente=12 * c - cor,
                   bruto=c * e["n"], diferidos=total_int, amort_crece=t[1]["amort"] / t[0]["amort"],
                   pagado_2026=k * c, interes_pagado=sum(r["interes"] for r in t[:k]),
                   pct_2027=cor / saldo(t, k))
    assert abs(R["ej"]["pasivo_total"] - saldo(t, k) * (1 + e["i"])) < 1e-6

    e = E1
    c, t = tabla(e["P"], e["i"], e["n"])
    asiento("6.1-cuota1", "", {"Préstamo bancario": t[0]["amort"], "Gasto financiero": t[0]["interes"]}, {"Banco": c})
    R["e1"] = dict(cuota=c, int1=t[0]["interes"], amort1=t[0]["amort"], saldo1=t[0]["fin"], interes_total=c * e["n"] - e["P"])
    q = E2["cuotas"]
    R["e2"] = dict(pagado=q * c, interes=sum(r["interes"] for r in t[:q]), amort=sum(r["amort"] for r in t[:q]),
                   saldo_ok=saldo(t, q), filas=t[:q])
    # 6.6 · método bruto con el mismo préstamo
    bruto, dif = c * e["n"], c * e["n"] - e["P"]
    asiento("6.6-recibe", "", {"Banco": e["P"], "Intereses diferidos": dif}, {"Préstamo bancario": bruto})
    asiento("6.6-devenga", "", {"Gasto financiero": t[0]["interes"]}, {"Intereses diferidos": t[0]["interes"]})
    asiento("6.6-paga", "", {"Préstamo bancario": c}, {"Banco": c})
    neto = (bruto - c) - (dif - t[0]["interes"])
    assert abs(neto - t[0]["fin"]) < 1e-6
    R["e6"] = dict(bruto=bruto, dif=dif, bruto_1=bruto - c, dif_1=dif - t[0]["interes"], neto_1=neto)

    e = E3
    c, t = tabla(e["P"], e["i"], e["n"])
    s5 = saldo(t, e["pagadas"])
    i6 = t[e["pagadas"]]["interes"]
    dev = s5 * e["i"] * e["dias"] / e["dias_mes"]
    asiento("6.3-ajuste", "30-06-2026", {"Gasto financiero": dev}, {"Intereses por pagar": dev})
    asiento("6.3-cuota6", "15-07-2026", {"Préstamo bancario": t[e["pagadas"]]["amort"], "Intereses por pagar": dev,
                                         "Gasto financiero": i6 - dev}, {"Banco": c})
    cor, nocor = corriente(t, e["pagadas"])
    R["e3"] = dict(cuota=c, saldo5=s5, i6=i6, dev=dev, julio=i6 - dev, amort6=t[e["pagadas"]]["amort"],
                   gasto_sem=sum(r["interes"] for r in t[:e["pagadas"]]) - 0 + dev,
                   int_ene=e["P"] * e["i"] * e["dias"] / e["dias_mes"], corriente=cor, no_corriente=nocor)
    # el gasto del semestre: del 15-01 al 30-06 = cuotas 1 a 5 (que cubren del 15-01 al 15-06) + 15 días de junio
    R["e3"]["gasto_sem"] = sum(r["interes"] for r in t[:e["pagadas"]]) + dev

    e = E4
    iva = e["neto"] * IVA
    banco = e["pie"] + iva
    asiento("6.4-compra", "01-03-2026", {"Vehículos": e["neto"], "IVA crédito fiscal": iva},
            {"Banco": banco, "Documentos por pagar a corto plazo": e["letra"], "Documentos por pagar a largo plazo": e["letra"]})
    assert abs(e["pie"] + 2 * e["letra"] - e["neto"]) < 1e-6
    R["e4"] = dict(iva=iva, banco=banco, bruto=e["neto"] + iva, letras=2 * e["letra"])

    e = E5
    c, t = tabla(e["P"], e["i"], e["n"])
    cor, nocor = corriente(t, e["pagadas"])
    R["e5"] = dict(cuota=c, saldo12=saldo(t, 12), corriente=cor, no_corriente=nocor, doce=12 * c,
                   int_futuro=12 * c - cor, gasto_2026=sum(r["interes"] for r in t[:12]), saldo24=saldo(t, 24))

    e = E7
    mensual = e["P"] * e["i"]
    i26 = mensual * e["meses_2026"]
    i27 = mensual * (e["meses"] - e["meses_2026"])
    asiento("6.7-ajuste", "31-12-2026", {"Gasto financiero": i26}, {"Intereses por pagar": i26})
    pago = asiento("6.7-pago", "01-05-2027", {"Documentos por pagar": e["P"], "Intereses por pagar": i26, "Gasto financiero": i27},
                   {"Banco": e["P"] + mensual * e["meses"]})
    R["e7"] = dict(mensual=mensual, i26=i26, i27=i27, total=mensual * e["meses"], pago=pago)

    e = E8
    amort = e["P"] / e["n"]
    filas, s = [], e["P"]
    for k in range(1, e["n"] + 1):
        it = s * e["i"]
        filas.append(dict(k=k, interes=it, cuota=amort + it))
        s -= amort
    tot = sum(f["interes"] for f in filas)
    assert abs(tot - e["i"] * e["P"] * (e["n"] + 1) / 2) < 1e-6
    R["e8"] = dict(amort=amort, int1=filas[0]["interes"], c1=filas[0]["cuota"], int2=filas[1]["interes"],
                   c2=filas[1]["cuota"], c8=filas[-1]["cuota"], total=tot, saldo2=e["P"] - 2 * amort)

    e = E9
    c, t = tabla(e["P"], e["i"], e["n"])
    i1 = t[0]["interes"]
    dev26 = i1 / 2
    i2 = t[1]["interes"]
    asiento("6.9-ajuste26", "31-12-2026", {"Gasto financiero": dev26}, {"Intereses por pagar": dev26})
    asiento("6.9-cuota1", "01-07-2027", {"Préstamo bancario": t[0]["amort"], "Intereses por pagar": dev26,
                                         "Gasto financiero": i1 - dev26}, {"Banco": c})
    R["e9"] = dict(cuota=c, i1=i1, dev26=dev26, amort1=t[0]["amort"], saldo1=t[0]["fin"], i2=i2,
                   gasto27=(i1 - dev26) + i2 / 2, dev27=i2 / 2, interes_total=c * e["n"] - e["P"])

    e = E10
    c, t = tabla(e["P"], e["i"], e["n"])
    cor, nocor = corriente(t, e["pagadas"])
    gasto = sum(r["interes"] for r in t[:12])
    pc = e["pc_sin"] + cor
    R["e10"] = dict(cuota=c, saldo12=saldo(t, 12), corriente=cor, no_corriente=nocor, razon_mal=e["ac"] / e["pc_sin"],
                    pc=pc, razon_ok=e["ac"] / pc, gasto=gasto, cobertura=e["r_operacional"] / gasto,
                    ac_necesario=e["minimo"] * pc, falta=e["minimo"] * pc - e["ac"],
                    porcion_max=e["ac"] / e["minimo"] - e["pc_sin"])
    return R


# ---------------------------------------------------------------- caso integrador

def cargar(nombre, ruta):
    spec = importlib.util.spec_from_file_location(nombre, ruta)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def caso():
    m = cargar("caso", AQUI / "caso_estampa_sur_contable.py")
    C = m.calcular()
    P, i, n = m.PRESTAMO["monto"], m.PRESTAMO["tasa_mensual"], m.PRESTAMO["cuotas"]
    c, t = tabla(P, i, n)
    propio = dict(cuota=c, interes_mes_1=t[0]["interes"], amortizacion_mes_1=t[0]["amort"],
                  saldo_tras_cuota_1=t[0]["fin"], interes_total=c * n - P)
    asiento("C-recibe (asiento 5 del caso)", "03-2026", {"Banco": P}, {"Préstamo bancario": P})
    asiento("C-6.1", "31-03", {"Gasto financiero": t[0]["interes"]}, {"Intereses por pagar": t[0]["interes"]})
    asiento("C-cuota1", "01-04", {"Préstamo bancario": t[0]["amort"], "Intereses por pagar": t[0]["interes"]}, {"Banco": c})
    cor, nocor = corriente(t, 0)
    doc = m.MAQUINA["neto"] - m.MAQUINA["pie"]
    # Finanzas, unidad 5: el valor presente del saldo de la máquina al 12 % anual
    fin = cargar("valentina", RAMO.parent / "finanzas" / "verificacion" / "caso_valentina.py")
    F = fin.calcular()
    vp = doc / (1 + fin.U5["tasa"]) ** fin.U5["anio_saldo"]
    assert abs(vp - F["u5"]["vp_saldo"]) < 1e-6 and doc == fin.U5["saldo_maquina"]
    extra = dict(P=P, i=i, n=n, t=t, corriente=cor, no_corriente=nocor, documento=doc,
                 pasivo_financiero=P + t[0]["interes"] + doc, corriente_total=cor + t[0]["interes"],
                 no_corriente_total=nocor + doc, gasto_u8=C["u8"]["gastos_financieros"], vp_documento=vp,
                 interes_implicito=doc - vp, total_pagado=c * n, tasa_pct=i)
    return C, propio, extra


def cadena(C, propio, extra):
    ok = all(abs(C["u6"][k] - v) < 1e-6 for k, v in propio.items())
    ok = ok and abs(extra["gasto_u8"] - propio["interes_mes_1"]) < 1e-6
    print("CADENA:", "OK, la pieza u6 calza con caso_estampa_sur_contable.py (y el gasto financiero de u8)"
          if ok else f"NO calza: {C['u6']} contra {propio}")
    return ok


# ---------------------------------------------------------------- planilla con fórmulas

def evaluar(wb):
    """Evalúa las fórmulas de la planilla (referencias, + − × ÷ ^ y SUMA) sin Excel."""
    from openpyxl.utils import range_boundaries, get_column_letter as L
    memo = {}

    def valor(hoja, ref):
        clave = (hoja, ref.replace("$", ""))
        if clave not in memo:
            v = wb[hoja][clave[1]].value
            if isinstance(v, str) and v.startswith("="):
                v = calcular(hoja, v[1:])
            memo[clave] = v if v is not None else 0
        return memo[clave]

    def rango(hoja, a, b):
        c1, f1, c2, f2 = range_boundaries(f"{a}:{b}".replace("$", ""))
        return [valor(hoja, f"{L(c)}{f}") for f in range(f1, f2 + 1) for c in range(c1, c2 + 1)]

    hoja_re = r"(?:'([^']+)'|([A-Za-z_]\w*))!"

    def calcular(hoja, f):
        f = re.sub(rf"SUM\((?:{hoja_re})?(\$?[A-Z]+\$?\d+):(\$?[A-Z]+\$?\d+)\)",
                   lambda mt: f"sum(_r({(mt.group(1) or mt.group(2) or hoja)!r},{mt.group(3)!r},{mt.group(4)!r}))", f)
        f = re.sub(rf"(?<![\w'])(?:{hoja_re})?(\$?[A-Z]{{1,2}}\$?\d+)(?![\w(])",
                   lambda mt: f"_v({(mt.group(1) or mt.group(2) or hoja)!r},{mt.group(3)!r})", f)
        return eval(f.replace("^", "**"), {"__builtins__": {}, "sum": sum, "_v": valor, "_r": rango})

    return valor


def hoja_prestamo(wb, titulo, P, i, n, pagadas, nota, primera=False):
    """Una hoja con los datos, la cuota, la tabla y la clasificación al cierre; todo con fórmulas."""
    from openpyxl.styles import Font
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo
    neg = Font(bold=True)
    ws["A1"], ws["B1"] = "Monto del préstamo", P
    ws["A2"], ws["B2"] = "Tasa mensual", i
    ws["A3"], ws["B3"] = "Número de cuotas", n
    ws["A4"], ws["B4"] = "Cuota (P·i / (1 − (1+i)^−n))", "=B1*B2/(1-(1+B2)^(0-B3))"
    ws["A5"], ws["B5"] = "Cuotas pagadas al cierre", pagadas
    ws["D1"] = nota
    enc = 7
    for col, txt in zip("ABCDEF", ["Cuota N.º", "Saldo inicial", "Interés", "Amortización", "Cuota", "Saldo final"]):
        ws[f"{col}{enc}"] = txt
        ws[f"{col}{enc}"].font = neg
    for k in range(1, n + 1):
        r = enc + k
        ws[f"A{r}"] = k
        ws[f"B{r}"] = "=B1" if k == 1 else f"=F{r - 1}"
        ws[f"C{r}"] = f"=B{r}*$B$2"
        ws[f"E{r}"] = "=$B$4"
        ws[f"D{r}"] = f"=E{r}-C{r}"
        ws[f"F{r}"] = f"=B{r}-D{r}"
    ult = enc + n
    tot = ult + 1
    ws[f"A{tot}"] = "Totales"
    ws[f"C{tot}"], ws[f"D{tot}"], ws[f"E{tot}"] = f"=SUM(C{enc + 1}:C{ult})", f"=SUM(D{enc + 1}:D{ult})", f"=SUM(E{enc + 1}:E{ult})"
    k0, k12 = pagadas, min(pagadas + 12, n)
    saldo_k = "B1" if k0 == 0 else f"F{enc + k0}"
    filas = [("Saldo de capital al cierre", f"={saldo_k}"),
             ("Interés devengado al cierre (de la cuota siguiente)", f"=C{enc + k0 + 1}"),
             ("Porción corriente (capital de las 12 cuotas siguientes)", f"=SUM(D{enc + k0 + 1}:D{enc + k12})"),
             ("Porción no corriente", f"=F{enc + k12}"),
             ("Comprobación: corriente + no corriente − saldo (debe dar 0)", f"=B{tot + 4}+B{tot + 5}-B{tot + 2}")]
    for j, (txt, fo) in enumerate(filas, start=2):
        ws[f"A{tot + j}"], ws[f"B{tot + j}"] = txt, fo
        ws[f"A{tot + j}"].font = neg if j in (4, 5) else Font()
    ws.column_dimensions["A"].width = 56
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16
        for r in list(range(enc + 1, tot + 1)) + list(range(tot + 2, tot + 7)):
            ws[f"{col}{r}"].number_format = "#,##0.00"
    ws["B1"].number_format = ws["B4"].number_format = "#,##0.00"
    ws["B2"].number_format = "0.00%"
    return dict(hoja=titulo, enc=enc, tot=tot)


def planilla(R, extra):
    from openpyxl import Workbook
    wb = Workbook()
    h1 = hoja_prestamo(wb, "Caso Estampa Sur", extra["P"], extra["i"], extra["n"], 0,
                       "Préstamo bancario de marzo. Cierre al 31-03: la cuota 1 se paga el 01-04.", primera=True)
    h2 = hoja_prestamo(wb, "Ejemplo Masa Madre", EJ["P"], EJ["i"], EJ["n"], EJ["pagadas_al_cierre"],
                       "Recibido el 01-10-2026; cuotas el día 1 desde el 01-11-2026; cierre al 31-12-2026.")
    destino = RAMO / "planillas" / "u06_prestamos.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)

    v = evaluar(wb)
    pares = []
    for h, t, cuo, cor, nocor, dev in [(h1, extra["t"], extra["t"][0]["cuota"], extra["corriente"], extra["no_corriente"], extra["t"][0]["interes"]),
                                       (h2, R["ej"]["t"], R["ej"]["cuota"], R["ej"]["corriente"], R["ej"]["no_corriente"], R["ej"]["i3"])]:
        s, e, tot = h["hoja"], h["enc"], h["tot"]
        pares += [(v(s, "B4"), cuo), (v(s, f"F{e + len(t)}"), 0), (v(s, f"C{tot}"), sum(r["interes"] for r in t)),
                  (v(s, f"D{tot}"), t[0]["ini"]), (v(s, f"B{tot + 3}"), dev), (v(s, f"B{tot + 4}"), cor),
                  (v(s, f"B{tot + 5}"), nocor), (v(s, f"B{tot + 6}"), 0)]
        pares += [(v(s, f"F{e + r['k']}"), r["fin"]) for r in t]
    ok = all(abs(a - b) < 1e-6 for a, b in pares)
    print(f"PLANILLA: {'OK' if ok else 'NO calza'}, {destino.relative_to(RAMO.parents[1])} ({len(pares)} celdas comprobadas)")
    return ok


# ---------------------------------------------------------------- cifras que el .md debe contener

def esperados(R, propio, extra):
    e = []
    j = R["ej"]
    e += tabla_md(j["t"], 0)
    e += [p(j[x]) for x in ("cuota", "interes_total", "total_pagado", "i3", "gasto_2026", "saldo_cierre", "corriente",
                            "no_corriente", "pasivo_total", "doce_cuotas", "error_corriente", "bruto", "diferidos",
                            "pagado_2026", "interes_pagado")]
    assert 0.70 < j["pct_2027"] < 0.75, "el texto dice «casi tres cuartos» de lo adeudado se paga en 2027"
    e += [fmt(j["amort_crece"], 4)]
    j = R["e1"]; e += [p(j[x]) for x in ("cuota", "int1", "amort1", "saldo1", "interes_total")]
    j = R["e2"]; e += [p(j[x]) for x in ("pagado", "interes", "amort", "saldo_ok")]
    j = R["e6"]; e += [p(j[x]) for x in ("bruto", "dif", "bruto_1", "dif_1", "neto_1")]
    j = R["e3"]; e += [p(j[x]) for x in ("cuota", "saldo5", "i6", "dev", "julio", "amort6", "gasto_sem", "corriente", "no_corriente")]
    j = R["e4"]; e += [p(j[x]) for x in ("iva", "banco", "bruto", "letras")]
    j = R["e5"]; e += [p(j[x]) for x in ("cuota", "saldo12", "corriente", "no_corriente", "doce", "int_futuro", "gasto_2026")]
    j = R["e7"]; e += [p(j[x]) for x in ("mensual", "i26", "i27", "total", "pago")]
    j = R["e8"]; e += [p(j[x]) for x in ("amort", "int1", "c1", "int2", "c2", "c8", "total", "saldo2")]
    j = R["e9"]; e += [p(j[x]) for x in ("cuota", "i1", "dev26", "amort1", "saldo1", "i2", "gasto27", "dev27", "interes_total")]
    j = R["e10"]; e += [p(j[x]) for x in ("cuota", "saldo12", "corriente", "no_corriente", "pc", "gasto", "ac_necesario", "falta", "porcion_max")]
    e += [fmt(j["razon_mal"], 2), fmt(j["razon_ok"], 2), fmt(j["cobertura"], 2)]
    e += [p2(v) for v in propio.values()]
    e += tabla_md(extra["t"], 2)
    e += [p2(extra[x]) for x in ("corriente", "no_corriente", "corriente_total", "no_corriente_total", "pasivo_financiero", "total_pagado")]
    e += [p(extra["documento"]), p(extra["vp_documento"]), p(extra["interes_implicito"]), pct(extra["tasa_pct"])]
    return e


if __name__ == "__main__":
    R = resultados()
    C, propio, extra = caso()
    for k, d in R.items():
        print(f"== {k}")
        for kk, v in d.items():
            if kk in ("t", "filas"):
                continue
            print(f"   {kk}: {round(v, 4) if isinstance(v, float) else v}")
    print("== caso (propio)", {k: round(v, 4) for k, v in propio.items()})
    print("== caso (extra)", {k: round(v, 4) for k, v in extra.items() if k != "t"})
    print(f"ASIENTOS: {len(ASIENTOS)} registrados, todos cuadran")
    if "--tablas" in sys.argv:
        print("\n-- tabla del ejemplo\n" + "\n".join(tabla_md(R["ej"]["t"], 0)))
        print("\n-- tabla del caso\n" + "\n".join(tabla_md(extra["t"], 2)))
    ok = cadena(C, propio, extra) and planilla(R, extra)
    if not ok:
        sys.exit(1)
    if "--revisar" in sys.argv:
        md = (RAMO / "unidades" / "06-prestamos-e-intereses.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %")
        faltan = [x for x in esperados(R, propio, extra) if x not in md]
        print("REVISIÓN:", "OK, todas las cifras están en el texto" if not faltan else f"faltan {len(faltan)}: {faltan}")
        sys.exit(1 if faltan else 0)
