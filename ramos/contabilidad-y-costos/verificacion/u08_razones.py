"""Unidad 8 · Estados financieros y razones financieras.

Recalcula todas las cifras de la unidad desde los datos. Para el caso, arma el balance de Estampa Sur
al 31 de marzo pasando los 12 asientos del caso, los ajustes de cierre, la producción y el costo de
ventas por las cuentas, y comprueba que cuadre. Genera la planilla con fórmulas y la comprueba contra
Python. Con --revisar busca cada resultado en el .md.

    py -3 -P u08_razones.py            # imprime los resultados y genera la planilla
    py -3 -P u08_razones.py --revisar  # además busca las cifras en el .md

Todos los datos son originales de esta base. No vienen de ningún libro ni curso.
"""
from pathlib import Path
import importlib.util, re, sys

sys.stdout.reconfigure(encoding="utf-8")
AQUI = Path(__file__).resolve().parent
DIAS_MARZO = 31


# ---------------------------------------------------------------- herramientas

def fmt(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, dec=1):
    return fmt(100 * x, dec) + " %"


def veces(x, dec=2):
    return fmt(x, dec)


def razones(d, dias=365):
    """Razones del curso desde un diccionario con los rubros. Flujos del período, saldos al cierre."""
    r = {}
    r["liquidez"] = d["ac"] / d["pc"]
    r["acida"] = (d["ac"] - d["inv"]) / d["pc"]
    r["ct"] = d["ac"] - d["pc"]
    r["d_p"] = d["pasivo"] / d["patrimonio"]
    r["d_a"] = d["pasivo"] / d["activo"]
    r["lp"] = d["pnc"] / d["pasivo"]
    r["rot_inv"] = d["cv"] / d["inv"]
    r["dias_inv"] = dias / r["rot_inv"]
    r["rot_cxc"] = d["ventas"] / d["cxc"]
    r["dias_cxc"] = dias / r["rot_cxc"]
    r["rot_cxc_iva"] = d["ventas"] * 1.19 / d["cxc"]
    r["dias_cxc_iva"] = dias / r["rot_cxc_iva"]
    r["mb"] = (d["ventas"] - d["cv"]) / d["ventas"]
    r["mn"] = d["un"] / d["ventas"]
    r["roa"] = d["ro"] / d["activo"]
    r["roe"] = d["un"] / d["patrimonio"]
    r["rot_act"] = d["ventas"] / d["activo"]
    r["cobertura"] = d["ro"] / d["gf"] if d.get("gf") else None
    return r


def cargar_caso():
    spec = importlib.util.spec_from_file_location("caso", AQUI / "caso_estampa_sur_contable.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m, m.calcular()


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · una ferretería al 31-12-2025
EJ = {"ventas": 96_400_000, "cv": 63_700_000, "gav": 22_950_000, "gf": 1_480_000, "impuesto": 2_105_000,
      "caja": 6_840_000, "clientes": 9_120_000, "existencias": 14_350_000, "seguro_ant": 1_260_000,
      "vehiculos": 18_900_000, "dep_acum": 5_670_000, "terreno": 11_500_000,
      "proveedores": 10_480_000, "iva_pagar": 1_150_000, "sueldos_pagar": 2_310_000,
      "prestamo_cp": 3_600_000, "prestamo_lp": 9_800_000, "capital": 20_000_000, "credito_nuevo": 5_800_000}
# 8.1 · estado de resultados por función
E1 = {"ventas": 41_850_000, "cv": 27_630_000, "adm": 6_920_000, "venta": 3_140_000, "gf": 865_000,
      "otros": 420_000, "impuesto": 1_010_000}
# 8.2 a 8.6 · una distribuidora (balance y resultados)
E2 = {"banco": 3_415_000, "clientes": 5_280_000, "mercaderias": 7_960_000, "arriendo_ant": 900_000,
      "maquinaria": 16_400_000, "dep_acum": 4_920_000, "terreno": 9_300_000, "proveedores": 6_140_000,
      "remuneraciones": 1_870_000, "iva_pagar": 640_000, "prestamo": 8_200_000, "prestamo_12m": 2_050_000,
      "capital": 18_000_000}
E5 = {"ventas": 58_300_000, "cv": 39_900_000, "inv_inicial": 7_240_000}
E6 = {"ro": 5_150_000}
# 8.7 · dos años de una misma empresa
E7 = {2024: {"ac": 12_600_000, "inv": 4_100_000, "pc": 6_300_000, "pasivo": 14_800_000, "patrimonio": 18_500_000,
             "ventas": 47_200_000, "un": 2_830_000},
      2025: {"ac": 15_900_000, "inv": 7_850_000, "pc": 9_940_000, "pasivo": 22_600_000, "patrimonio": 19_700_000,
             "ventas": 51_900_000, "un": 2_170_000}}
# 8.8 · cómo mueve cada operación la liquidez
E8 = {"ac": 17_640_000, "pc": 11_760_000, "inv": 5_880_000, "pago_prov": 2_940_000, "compra": 2_350_000,
      "prestamo_lp": 4_700_000, "maquina": 3_900_000}
# 8.9 · el método de inventario en las razones
E9 = {"ventas": 30_500_000, "disponible": 22_600_000, "if_fifo": 4_380_000, "if_ppp": 4_150_000,
      "otros_ac": 6_900_000, "pc": 7_200_000, "gastos": 7_400_000}
# 8.10 · el balance a partir de sus razones
E10 = {"activo": 40_000_000, "liquidez": 1.6, "acida": 0.9, "d_a": 0.45, "pc_de_pasivo": 0.60,
       "rot_inv": 5, "mb": 0.30, "roe": 0.12}


# ---------------------------------------------------------------- el caso: balance al 31 de marzo

def balance_caso(m, C):
    """Pasa por las cuentas el saldo inicial, los 12 asientos, los ajustes, la producción y el costo de
    ventas. Saldo deudor positivo, acreedor negativo."""
    ctas = {}

    def mov(cuenta, monto):
        ctas[cuenta] = ctas.get(cuenta, 0.0) + monto

    for k, v in m.BALANCE_INICIAL["activo"].items():
        mov(k, v)
    for k, v in m.BALANCE_INICIAL["pasivo"].items():
        mov(k, -v)
    mov("Capital", -C["u1"]["capital"])
    for a in C["u2"]["asientos"]:
        for k, v in a["debe"].items():
            mov(k, v)
        for k, v in a["haber"].items():
            mov(k, -v)
    # cierre del IVA: el crédito supera al débito y queda remanente
    deb, cred = -ctas.pop("IVA débito fiscal"), ctas.pop("IVA crédito fiscal")
    assert abs(cred - deb - C["u3"]["remanente"]) < 1e-6
    mov("Remanente de IVA crédito fiscal", cred - deb)
    # ajustes de cierre (unidades 4, 5 y 6)
    dep = C["u4"]["total_mes"]
    mov("Depreciación del mes", dep); mov("Depreciación acumulada", -dep)
    mov("Arriendo del mes", C["u5"]["arriendo_devengado"]); mov("Arriendo pagado por anticipado", -C["u5"]["arriendo_devengado"])
    mov("Seguro del mes", C["u5"]["seguro_devengado"]); mov("Seguro pagado por anticipado", -C["u5"]["seguro_devengado"])
    mov("Gastos financieros", C["u6"]["interes_mes_1"]); mov("Intereses por pagar", -C["u6"]["interes_mes_1"])
    # producción: materiales a estándar, variación de precio de las poleras (PPP, unidad 7)
    ppp = (m.BALANCE_INICIAL["activo"]["Poleras lisas básicas"] + C["u2"]["asientos"][0]["debe"]["Poleras lisas básicas"]) / \
          (900 + m.COMPRA_POLERAS["unidades"])
    consumo = C["u7"]["consumo"]
    mov("Poleras lisas básicas", -consumo * ppp)
    mov("Productos en proceso", consumo * 2_400)
    mov("Variación de precio de materias primas", consumo * (ppp - 2_400))
    otras = ctas["Otras prendas lisas"]
    mov("Otras prendas lisas", -otras); mov("Productos en proceso", otras)
    tinta_kg = sum(q * g for q, g in zip(m.PRODUCIDAS, m.TINTA_G)) / 1000
    mov("Tinta", -tinta_kg * m.TINTA_KG); mov("Productos en proceso", tinta_kg * m.TINTA_KG)
    otros_var = sum(q * o for q, o in zip(m.PRODUCIDAS, C["u9"]["otros_variables"]))
    mov("Productos en proceso", otros_var); mov("Insumos por facturar", -otros_var)
    for cuenta in ("Sueldos", "Depreciación del mes", "Arriendo del mes", "Seguro del mes"):   # CIF fijos
        mov("Productos en proceso", ctas[cuenta]); mov(cuenta, -ctas[cuenta])
    pp = ctas.pop("Productos en proceso")
    mov("Productos terminados", pp)
    costo_abs = C["u11"]["costo_absorcion_por_prenda"]
    cv = sum(q * c for q, c in zip(m.VENDIDAS, costo_abs))
    mov("Productos terminados", -cv); mov("Costo de ventas", cv)
    ctas = {k: v for k, v in ctas.items() if abs(v) > 1e-6}
    return ctas, dict(ppp=ppp, otros_var=otros_var, tinta_kg=tinta_kg, prod_total=pp, cv=cv)


def estados_caso(m, C):
    ctas, aux = balance_caso(m, C)
    u8 = C["u8"]
    var = ctas["Variación de precio de materias primas"]
    resultado_ctas = -(ctas["Ventas"] + ctas["Costo de ventas"] + ctas["Gastos de administración"] +
                       ctas["Gastos financieros"] + var)
    # el préstamo: lo que se amortiza en las cuotas 1 a 12 es pasivo corriente
    i, n, P = m.PRESTAMO["tasa_mensual"], m.PRESTAMO["cuotas"], m.PRESTAMO["monto"]
    cuota = P * i / (1 - (1 + i) ** -n)
    saldo12 = cuota * (1 - (1 + i) ** -(n - 12)) / i
    prest_cp = P - saldo12
    B = {
        "Banco": ctas["Banco"], "Clientes": ctas["Clientes"],
        "Remanente de IVA crédito fiscal": ctas["Remanente de IVA crédito fiscal"],
        "Poleras lisas básicas": ctas["Poleras lisas básicas"], "Tinta": ctas["Tinta"],
        "Productos terminados": ctas["Productos terminados"],
        "Arriendo pagado por anticipado": ctas["Arriendo pagado por anticipado"],
        "Seguro pagado por anticipado": ctas["Seguro pagado por anticipado"],
        "Maquinaria": ctas["Maquinaria"], "Depreciación acumulada": ctas["Depreciación acumulada"],
        "Proveedores": -ctas["Proveedores"], "Insumos por facturar": -ctas["Insumos por facturar"],
        "Intereses por pagar": -ctas["Intereses por pagar"], "Préstamo bancario (corto plazo)": prest_cp,
        "Préstamo bancario (largo plazo)": -ctas["Préstamo bancario"] - prest_cp,
        "Documentos por pagar a largo plazo": -ctas["Documentos por pagar a largo plazo"],
        "Capital": -ctas["Capital"], "Resultado del ejercicio": resultado_ctas,
    }
    inv = B["Poleras lisas básicas"] + B["Tinta"] + B["Productos terminados"]
    ac = sum(B[k] for k in list(B)[:8])
    anc = B["Maquinaria"] + B["Depreciación acumulada"]
    pc = B["Proveedores"] + B["Insumos por facturar"] + B["Intereses por pagar"] + B["Préstamo bancario (corto plazo)"]
    pnc = B["Préstamo bancario (largo plazo)"] + B["Documentos por pagar a largo plazo"]
    pat = B["Capital"] + B["Resultado del ejercicio"]
    assert abs(ac + anc - (pc + pnc + pat)) < 1e-4, "el balance del caso no cuadra"
    assert all(abs(v) < 1e-6 for k, v in ctas.items() if k in ("Otras prendas lisas",))
    ER = dict(ventas=u8["ventas"], cv=u8["costo_ventas"], mb=u8["margen_bruto"], adm=m.OTROS_FIJOS,
              ro=u8["resultado_operacional"], gf=u8["gastos_financieros"], uai=u8["utilidad_antes_impuesto"],
              var=var, cv_real=u8["costo_ventas"] + var, mb_real=u8["margen_bruto"] - var,
              ro_real=u8["resultado_operacional"] - var, uai_real=u8["utilidad_antes_impuesto"] - var)
    assert abs(ER["uai_real"] - resultado_ctas) < 1e-6, "el resultado de las cuentas no es el del estado"
    # recálculo propio del estado de resultados del caso
    assert abs(ctas["Ventas"] + u8["ventas"]) < 1e-6 and abs(ctas["Costo de ventas"] - u8["costo_ventas"]) < 1e-6
    # presentación a peso: cada partida se redondea y los totales se suman con las partidas redondeadas
    exacto_pt = B["Productos terminados"]
    B = {k: round(v) for k, v in B.items()}
    ER = {k: round(v) for k, v in ER.items()}
    # el peso de redondeo lo absorbe Productos terminados, que ya es un saldo calculado por diferencia
    otros = sum(v for k, v in list(B.items())[:10] if k != "Productos terminados")
    B["Productos terminados"] = sum(list(B.values())[10:16]) + B["Capital"] + B["Resultado del ejercicio"] - otros
    assert abs(B["Productos terminados"] - exacto_pt) <= 1.5
    for a, b, c in (("mb", "ventas", "cv"), ("mb_real", "ventas", "cv_real")):
        assert ER[a] == ER[b] - ER[c]
    assert ER["ro"] == ER["mb"] - ER["adm"] and ER["uai"] == ER["ro"] - ER["gf"]
    assert ER["uai_real"] == ER["uai"] - ER["var"] == B["Resultado del ejercicio"]
    inv = B["Poleras lisas básicas"] + B["Tinta"] + B["Productos terminados"]
    ac = sum(B[k] for k in list(B)[:8])
    anc = B["Maquinaria"] + B["Depreciación acumulada"]
    pc = B["Proveedores"] + B["Insumos por facturar"] + B["Intereses por pagar"] + B["Préstamo bancario (corto plazo)"]
    pnc = B["Préstamo bancario (largo plazo)"] + B["Documentos por pagar a largo plazo"]
    pat = B["Capital"] + B["Resultado del ejercicio"]
    assert ac + anc == pc + pnc + pat, "el balance redondeado no cuadra"
    tot = dict(ac=ac, anc=anc, activo=ac + anc, pc=pc, pnc=pnc, pasivo=pc + pnc, patrimonio=pat, inv=inv,
               maquinaria_neta=anc)
    d = dict(ac=ac, pc=pc, inv=inv, pasivo=pc + pnc, patrimonio=pat, activo=ac + anc, pnc=pnc, cv=ER["cv_real"],
             ventas=ER["ventas"], cxc=B["Clientes"], un=ER["uai_real"], ro=ER["ro_real"], gf=ER["gf"])
    r = razones(d, dias=DIAS_MARZO)
    r["mb_std"] = ER["mb"] / ER["ventas"]
    r["mn_std"] = ER["uai"] / ER["ventas"]
    r["acida_estricta"] = (ac - inv - B["Arriendo pagado por anticipado"] - B["Seguro pagado por anticipado"]) / pc
    r["roe_anual"] = 12 * r["roe"]
    pago = 10_000_000                     # decisión: pagar parte de Proveedores con el Banco
    r["pago"] = pago
    r["liq_pago"] = (ac - pago) / (pc - pago)
    r["dp_pago"] = (pc + pnc - pago) / pat
    r["banco_pago"] = B["Banco"] - pago
    r["pc_prov"] = B["Proveedores"] / pc
    r["prov_compras"] = -sum(a["haber"].get("Proveedores", 0) for a in C["u2"]["asientos"])
    return B, tot, ER, r, aux, dict(cuota=cuota, saldo12=saldo12, prest_cp=prest_cp)


# ---------------------------------------------------------------- resultados

def resultados(m, C):
    R = {}
    e = EJ
    mb = e["ventas"] - e["cv"]; ro = mb - e["gav"]; uai = ro - e["gf"]; un = uai - e["impuesto"]
    ac = e["caja"] + e["clientes"] + e["existencias"] + e["seguro_ant"]
    anc = e["vehiculos"] - e["dep_acum"] + e["terreno"]
    pc = e["proveedores"] + e["iva_pagar"] + e["sueldos_pagar"] + e["prestamo_cp"]
    pnc = e["prestamo_lp"]
    pat = ac + anc - pc - pnc
    retenidas = pat - e["capital"] - un
    d = dict(ac=ac, pc=pc, inv=e["existencias"], pasivo=pc + pnc, patrimonio=pat, activo=ac + anc, pnc=pnc,
             cv=e["cv"], ventas=e["ventas"], cxc=e["clientes"], un=un, ro=ro, gf=e["gf"])
    r = razones(d)
    R["ej"] = dict(mb=mb, ro=ro, uai=uai, un=un, ac=ac, anc=anc, activo=ac + anc, pc=pc, pnc=pnc, pasivo=pc + pnc,
                   pat=pat, retenidas=retenidas, r=r, veh_neto=e["vehiculos"] - e["dep_acum"],
                   dp_nuevo=(pc + pnc + e["credito_nuevo"]) / pat, holgura=pat - (pc + pnc),
                   cp_de_pasivo=pc / (pc + pnc))

    e = E1
    mb = e["ventas"] - e["cv"]; gastos = e["adm"] + e["venta"]; ro = mb - gastos
    rai = ro + e["otros"] - e["gf"]; un = rai - e["impuesto"]
    R["e1"] = dict(mb=mb, gastos=gastos, ro=ro, rai=rai, un=un, mb_pct=mb / e["ventas"], ro_pct=ro / e["ventas"],
                   un_pct=un / e["ventas"])

    e = E2
    ac = e["banco"] + e["clientes"] + e["mercaderias"] + e["arriendo_ant"]
    anc = e["maquinaria"] - e["dep_acum"] + e["terreno"]
    pc = e["proveedores"] + e["remuneraciones"] + e["iva_pagar"] + e["prestamo_12m"]
    pnc = e["prestamo"] - e["prestamo_12m"]
    pat = ac + anc - pc - pnc
    res = pat - e["capital"]
    R["e2"] = dict(ac=ac, anc=anc, activo=ac + anc, pc=pc, pnc=pnc, pasivo=pc + pnc, pat=pat, res=res,
                   maq_neta=e["maquinaria"] - e["dep_acum"])
    R["e3"] = dict(liquidez=ac / pc, acida=(ac - e["mercaderias"]) / pc, ct=ac - pc,
                   acida_estricta=(ac - e["mercaderias"] - e["arriendo_ant"]) / pc, sin_inv=ac - e["mercaderias"])
    R["e4"] = dict(d_p=(pc + pnc) / pat, d_a=(pc + pnc) / (ac + anc), lp=pnc / (pc + pnc), cp=pc / (pc + pnc),
                   pat_a=pat / (ac + anc))
    e5 = E5
    prom = (e5["inv_inicial"] + e["mercaderias"]) / 2
    rot_f = e5["cv"] / e["mercaderias"]; rot_p = e5["cv"] / prom
    rcxc = e5["ventas"] / e["clientes"]; rcxc_iva = e5["ventas"] * 1.19 / e["clientes"]
    R["e5"] = dict(rot_f=rot_f, dias_f=365 / rot_f, prom=prom, rot_p=rot_p, dias_p=365 / rot_p, rcxc=rcxc,
                   dias_cxc=365 / rcxc, rcxc_iva=rcxc_iva, dias_cxc_iva=365 / rcxc_iva, ventas_iva=e5["ventas"] * 1.19)
    un = res
    act = ac + anc
    mn = un / e5["ventas"]; ra = e5["ventas"] / act; mult = act / pat
    R["e6"] = dict(roa=E6["ro"] / act, roe=un / pat, mn=mn, rot=ra, mult=mult, dupont=mn * ra * mult,
                   roa_neto=un / act)
    assert abs(R["e6"]["dupont"] - R["e6"]["roe"]) < 1e-12

    out = {}
    for y, d in E7.items():
        out[y] = dict(liquidez=d["ac"] / d["pc"], acida=(d["ac"] - d["inv"]) / d["pc"], d_p=d["pasivo"] / d["patrimonio"],
                      mn=d["un"] / d["ventas"], roe=d["un"] / d["patrimonio"], ct=d["ac"] - d["pc"],
                      d_a=d["pasivo"] / (d["pasivo"] + d["patrimonio"]))
    out["crec"] = E7[2025]["ventas"] / E7[2024]["ventas"] - 1
    out["inv_crece"] = E7[2025]["inv"] / E7[2024]["inv"] - 1
    out["pc_crece"] = E7[2025]["pc"] - E7[2024]["pc"]
    R["e7"] = out

    e = E8
    base = dict(liq=e["ac"] / e["pc"], acida=(e["ac"] - e["inv"]) / e["pc"])
    a = (e["ac"] - e["pago_prov"]) / (e["pc"] - e["pago_prov"])
    a_acida = (e["ac"] - e["pago_prov"] - e["inv"]) / (e["pc"] - e["pago_prov"])
    bruto = e["compra"] * 1.19
    b = (e["ac"] + bruto) / (e["pc"] + bruto)
    b_acida = (e["ac"] + bruto - e["inv"] - e["compra"]) / (e["pc"] + bruto)
    c = (e["ac"] + e["prestamo_lp"]) / e["pc"]
    d_ = (e["ac"] - e["maquina"]) / e["pc"]
    R["e8"] = dict(base=base, a=a, a_acida=a_acida, b=b, b_acida=b_acida, c=c, d=d_, bruto=bruto, iva=e["compra"] * 0.19,
                   iva_maq=e["maquina"] * 0.19, ac_b=e["ac"] + bruto, pc_b=e["pc"] + bruto)

    e = E9
    out = {}
    for k in ("fifo", "ppp"):
        inv = e["if_" + k]
        cv = e["disponible"] - inv
        mb = e["ventas"] - cv
        ro = mb - e["gastos"]
        out[k] = dict(cv=cv, mb=mb, mb_pct=mb / e["ventas"], ro=ro, liq=(e["otros_ac"] + inv) / e["pc"],
                      acida=e["otros_ac"] / e["pc"], rot=cv / inv, dias=365 / (cv / inv))
    out["dif_ro"] = out["fifo"]["ro"] - out["ppp"]["ro"]
    out["dif_inv"] = e["if_fifo"] - e["if_ppp"]
    R["e9"] = out

    e = E10
    pasivo = e["d_a"] * e["activo"]; pat = e["activo"] - pasivo
    pc = e["pc_de_pasivo"] * pasivo; pnc = pasivo - pc
    ac = e["liquidez"] * pc; inv = ac - e["acida"] * pc; anc = e["activo"] - ac
    cv = e["rot_inv"] * inv; ventas = cv / (1 - e["mb"]); un = e["roe"] * pat
    R["e10"] = dict(pasivo=pasivo, pat=pat, pc=pc, pnc=pnc, ac=ac, inv=inv, anc=anc, cv=cv, ventas=ventas, un=un,
                    mn=un / ventas, sin_inv=ac - inv)
    # comprobación: las razones del balance reconstruido son las dadas
    assert abs(ac / pc - e["liquidez"]) < 1e-9 and abs((ac - inv) / pc - e["acida"]) < 1e-9

    B, tot, ER, r, aux, pr = estados_caso(m, C)
    R["caso"] = dict(B=B, tot=tot, ER=ER, r=r, aux=aux, pr=pr)
    return R


# ---------------------------------------------------------------- planilla con fórmulas

def evaluar(ws):
    cache = {}

    def valor(ref):
        if ref in cache:
            return cache[ref]
        v = ws[ref].value
        if isinstance(v, str) and v.startswith("="):
            expr = re.sub(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)",
                          lambda mm: "(" + "+".join(f"{mm.group(1)}{k}" for k in range(int(mm.group(2)), int(mm.group(4)) + 1)) + ")",
                          v[1:])
            expr = re.sub(r"\b([A-Z]{1,2}\d+)\b", lambda mm: repr(valor(mm.group(1))), expr)
            v = eval(expr)
        cache[ref] = v if v is not None else 0
        return cache[ref]
    return valor


def hoja_estados(wb, titulo, er, balance, dias):
    """er: lista de (rótulo, valor o None si es subtotal, fórmula de subtotal con nombres).
    balance: dict con listas por bloque. Devuelve la hoja y las celdas de cada razón."""
    from openpyxl.styles import Font
    ws = wb.create_sheet(titulo)
    ws["A1"] = titulo; ws["A1"].font = Font(bold=True)
    pos = {}
    r = 3
    ws.cell(r, 1, "Estado de resultados por función").font = Font(bold=True); r += 1
    for nombre, valor, formula in er:
        ws.cell(r, 1, nombre)
        if formula:
            ws.cell(r, 2, "=" + re.sub(r"\{(\w+)\}", lambda mm: pos[mm.group(1)], formula))
            ws.cell(r, 1).font = Font(bold=True)
        else:
            ws.cell(r, 2, valor)
        pos[nombre.split(" ·")[0].replace(" ", "_")] = f"B{r}"
        r += 1
    r += 1
    ws.cell(r, 1, "Balance (estado de situación financiera)").font = Font(bold=True); r += 1
    for bloque, cuentas in balance.items():
        ws.cell(r, 1, bloque).font = Font(italic=True); r += 1
        ini = r
        for nombre, valor in cuentas:
            ws.cell(r, 1, nombre); ws.cell(r, 2, valor); r += 1
        ws.cell(r, 1, "Total " + bloque.lower()); ws.cell(r, 2, f"=SUM(B{ini}:B{r - 1})")
        ws.cell(r, 1).font = Font(bold=True)
        pos["T_" + bloque] = f"B{r}"
        for nombre, _ in cuentas:
            pos[nombre.replace(" ", "_")] = f"B{ini + [c for c, _ in cuentas].index(nombre)}"
        r += 2
    return ws, pos, r


def escribir_razones(ws, pos, r, dias, inventario, cxc, ventas, cv, ro, un, gf):
    from openpyxl.styles import Font
    P = lambda k: pos[k]
    ac, pc = P("T_Activo corriente"), P("T_Pasivo corriente")
    anc, pnc, pat = P("T_Activo no corriente"), P("T_Pasivo no corriente"), P("T_Patrimonio")
    filas = [
        ("liquidez", "Liquidez corriente", f"={ac}/{pc}"),
        ("acida", "Prueba ácida", f"=({ac}-{inventario})/{pc}"),
        ("ct", "Capital de trabajo", f"={ac}-{pc}"),
        ("d_p", "Pasivo / patrimonio", f"=({pc}+{pnc})/{pat}"),
        ("d_a", "Pasivo / activo", f"=({pc}+{pnc})/({ac}+{anc})"),
        ("lp", "Pasivo de largo plazo / pasivo", f"={pnc}/({pc}+{pnc})"),
        ("rot_inv", "Rotación de inventario", f"={cv}/{inventario}"),
        ("dias_inv", f"Días de inventario ({dias} días)", f"={dias}/B{{rot_inv}}"),
        ("rot_cxc", "Rotación de cuentas por cobrar", f"={ventas}/{cxc}"),
        ("dias_cxc", f"Días de cobro ({dias} días)", f"={dias}/B{{rot_cxc}}"),
        ("mb", "Margen bruto", f"=({ventas}-{cv})/{ventas}"),
        ("mn", "Margen neto", f"={un}/{ventas}"),
        ("roa", "ROA (resultado operacional / activo)", f"={ro}/({ac}+{anc})"),
        ("roe", "ROE (resultado / patrimonio)", f"={un}/{pat}"),
        ("cobertura", "Cobertura de gastos financieros", f"={ro}/{gf}"),
    ]
    ws.cell(r, 1, "Razones").font = Font(bold=True); r += 1
    celdas = {}
    for clave, nombre, formula in filas:
        formula = re.sub(r"B\{(\w+)\}", lambda mm: celdas[mm.group(1)], formula)
        ws.cell(r, 1, nombre); ws.cell(r, 2, formula)
        ws.cell(r, 2).number_format = "0.0%" if clave in ("d_a", "lp", "mb", "mn", "roa", "roe") else (
            "#,##0" if clave == "ct" else "0.00")
        celdas[clave] = f"B{r}"; r += 1
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 16
    for row in ws.iter_rows(min_row=3, max_row=r):
        c = row[1]
        if c.number_format == "General":
            c.number_format = "#,##0"
    return celdas


def planilla(R):
    from openpyxl import Workbook
    wb = Workbook(); wb.remove(wb.active)
    c = R["caso"]; B, ER = c["B"], c["ER"]
    er = [("Ventas", ER["ventas"], None), ("Costo de ventas a estándar", -ER["cv"], None),
          ("Variación de precio de materias primas", -ER["var"], None),
          ("Margen bruto", None, "{Ventas}+{Costo_de_ventas_a_estándar}+{Variación_de_precio_de_materias_primas}"),
          ("Gastos de administración", -ER["adm"], None),
          ("Resultado operacional", None, "{Margen_bruto}+{Gastos_de_administración}"),
          ("Gastos financieros", -ER["gf"], None),
          ("Resultado antes de impuesto", None, "{Resultado_operacional}+{Gastos_financieros}")]
    ks = list(B)
    bal = {"Activo corriente": [(k, B[k]) for k in ks[:8]],
           "Activo no corriente": [(k, B[k]) for k in ks[8:10]],
           "Pasivo corriente": [(k, B[k]) for k in ks[10:14]],
           "Pasivo no corriente": [(k, B[k]) for k in ks[14:16]],
           "Patrimonio": [("Capital", B["Capital"])]}
    ws, pos, r = hoja_estados(wb, "Caso Estampa Sur marzo", er, bal, DIAS_MARZO)
    # el resultado del ejercicio en el patrimonio es la fórmula del estado de resultados
    ultima = int(pos["T_Patrimonio"][1:])
    ws.insert_rows(ultima)
    ws.cell(ultima, 1, "Resultado del ejercicio"); ws.cell(ultima, 2, f"={pos['Resultado_antes_de_impuesto']}")
    ws.cell(ultima + 1, 2, f"=SUM(B{int(pos['Capital'][1:])}:B{ultima})")
    pos["T_Patrimonio"] = f"B{ultima + 1}"
    ws.cell(ultima + 2, 1, "Activo − pasivo − patrimonio (debe dar 0)")
    ws.cell(ultima + 2, 2, f"={pos['T_Activo corriente']}+{pos['T_Activo no corriente']}-{pos['T_Pasivo corriente']}"
                           f"-{pos['T_Pasivo no corriente']}-{pos['T_Patrimonio']}")
    inv = f"({pos['Poleras_lisas_básicas']}+{pos['Tinta']}+{pos['Productos_terminados']})"
    cv = f"(-{pos['Costo_de_ventas_a_estándar']}-{pos['Variación_de_precio_de_materias_primas']})"
    cel = escribir_razones(ws, pos, ultima + 4, DIAS_MARZO, inv, pos["Clientes"], pos["Ventas"], cv,
                           pos["Resultado_operacional"], pos["Resultado_antes_de_impuesto"],
                           f"(-{pos['Gastos_financieros']})")
    hojas = [(ws, cel, c["r"], f"B{ultima + 2}")]

    e = EJ; j = R["ej"]
    er = [("Ventas", e["ventas"], None), ("Costo de ventas", -e["cv"], None),
          ("Margen bruto", None, "{Ventas}+{Costo_de_ventas}"), ("Gastos de administración y ventas", -e["gav"], None),
          ("Resultado operacional", None, "{Margen_bruto}+{Gastos_de_administración_y_ventas}"),
          ("Gastos financieros", -e["gf"], None),
          ("Resultado antes de impuesto", None, "{Resultado_operacional}+{Gastos_financieros}"),
          ("Impuesto a la renta", -e["impuesto"], None),
          ("Utilidad del ejercicio", None, "{Resultado_antes_de_impuesto}+{Impuesto_a_la_renta}")]
    bal = {"Activo corriente": [("Caja y banco", e["caja"]), ("Clientes", e["clientes"]), ("Existencias", e["existencias"]),
                                ("Seguro pagado por anticipado", e["seguro_ant"])],
           "Activo no corriente": [("Vehículos", e["vehiculos"]), ("Depreciación acumulada", -e["dep_acum"]),
                                   ("Terreno", e["terreno"])],
           "Pasivo corriente": [("Proveedores", e["proveedores"]), ("IVA por pagar", e["iva_pagar"]),
                                ("Sueldos por pagar", e["sueldos_pagar"]), ("Préstamo bancario (corto plazo)", e["prestamo_cp"])],
           "Pasivo no corriente": [("Préstamo bancario (largo plazo)", e["prestamo_lp"])],
           "Patrimonio": [("Capital", e["capital"]), ("Utilidades retenidas", j["retenidas"])]}
    ws, pos, r = hoja_estados(wb, "Ejemplo ferretería 2025", er, bal, 365)
    ultima = int(pos["T_Patrimonio"][1:])
    ws.insert_rows(ultima)
    ws.cell(ultima, 1, "Utilidad del ejercicio"); ws.cell(ultima, 2, f"={pos['Utilidad_del_ejercicio']}")
    ws.cell(ultima + 1, 2, f"=SUM(B{int(pos['Capital'][1:])}:B{ultima})")
    pos["T_Patrimonio"] = f"B{ultima + 1}"
    ws.cell(ultima + 2, 1, "Activo − pasivo − patrimonio (debe dar 0)")
    ws.cell(ultima + 2, 2, f"={pos['T_Activo corriente']}+{pos['T_Activo no corriente']}-{pos['T_Pasivo corriente']}"
                           f"-{pos['T_Pasivo no corriente']}-{pos['T_Patrimonio']}")
    cel = escribir_razones(ws, pos, ultima + 4, 365, pos["Existencias"], pos["Clientes"], pos["Ventas"],
                           f"(-{pos['Costo_de_ventas']})", pos["Resultado_operacional"], pos["Utilidad_del_ejercicio"],
                           f"(-{pos['Gastos_financieros']})")
    hojas.append((ws, cel, j["r"], f"B{ultima + 2}"))
    destino = AQUI.parents[0] / "planillas" / "u08_razones.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)
    ok = True
    for ws, cel, r, cuadre in hojas:
        v = evaluar(ws)
        ok &= abs(v(cuadre)) < 1e-4
        for k, ref in cel.items():
            ok &= abs(v(ref) - r[k]) < 1e-9
    print(f"PLANILLA: {'OK' if ok else 'NO calza'}, {destino.name}: estados, cuadre y razones con fórmulas")
    return ok


# ---------------------------------------------------------------- texto

def esperados(R, C):
    p = fmt
    e = []
    j = R["ej"]; r = j["r"]
    e += [p(x) for x in (j["mb"], j["ro"], j["uai"], j["un"], j["ac"], j["anc"], j["activo"], j["pc"], j["pnc"],
                         j["pasivo"], j["pat"], j["retenidas"], j["veh_neto"], j["holgura"], r["ct"])]
    e += [veces(r["liquidez"]), veces(r["acida"]), veces(r["d_p"]), pct(r["d_a"]), pct(r["lp"]), veces(r["rot_inv"]),
          fmt(r["dias_inv"], 1), veces(r["rot_cxc"]), fmt(r["dias_cxc"], 1), veces(r["rot_cxc_iva"]),
          fmt(r["dias_cxc_iva"], 1), pct(r["mb"]), pct(r["mn"]), pct(r["roa"]), pct(r["roe"]), fmt(r["cobertura"], 1),
          veces(j["dp_nuevo"]), pct(j["cp_de_pasivo"])]
    j = R["e1"]; e += [p(j["mb"]), p(j["gastos"]), p(j["ro"]), p(j["rai"]), p(j["un"]), pct(j["mb_pct"]),
                       pct(j["ro_pct"]), pct(j["un_pct"])]
    j = R["e2"]; e += [p(j[k]) for k in ("ac", "anc", "activo", "pc", "pnc", "pasivo", "pat", "res", "maq_neta")]
    j = R["e3"]; e += [veces(j["liquidez"]), veces(j["acida"]), p(j["ct"]), veces(j["acida_estricta"]), p(j["sin_inv"])]
    j = R["e4"]; e += [veces(j["d_p"]), pct(j["d_a"]), pct(j["lp"]), pct(j["cp"]), pct(j["pat_a"])]
    j = R["e5"]; e += [veces(j["rot_f"]), fmt(j["dias_f"], 1), p(j["prom"]), veces(j["rot_p"]), fmt(j["dias_p"], 1),
                       veces(j["rcxc"]), fmt(j["dias_cxc"], 1), veces(j["rcxc_iva"]), fmt(j["dias_cxc_iva"], 1),
                       p(j["ventas_iva"])]
    j = R["e6"]; e += [pct(j["roa"]), pct(j["roe"]), pct(j["mn"], 2), veces(j["rot"]), fmt(j["mult"], 3), pct(j["roa_neto"])]
    j = R["e7"]
    for y in (2024, 2025):
        e += [veces(j[y]["liquidez"]), veces(j[y]["acida"]), veces(j[y]["d_p"]), pct(j[y]["mn"]), pct(j[y]["roe"]),
              p(j[y]["ct"])]
    e += [pct(j["crec"]), pct(j["inv_crece"]), p(j["pc_crece"])]
    j = R["e8"]; e += [veces(j["base"]["liq"]), veces(j["base"]["acida"]), fmt(j["a"], 3), fmt(j["a_acida"], 3),
                       fmt(j["b"], 3), fmt(j["b_acida"], 3), fmt(j["c"], 3), fmt(j["d"], 3), p(j["bruto"]), p(j["iva"]),
                       p(j["ac_b"]), p(j["pc_b"]), p(j["iva_maq"])]
    j = R["e9"]
    for k in ("fifo", "ppp"):
        e += [p(j[k]["cv"]), p(j[k]["mb"]), pct(j[k]["mb_pct"]), p(j[k]["ro"]), fmt(j[k]["liq"], 3), veces(j[k]["rot"]),
              fmt(j[k]["dias"], 1)]
    e += [fmt(j["fifo"]["acida"], 3), p(j["dif_ro"]), p(j["dif_inv"])]
    j = R["e10"]; e += [p(j[k]) for k in ("pasivo", "pat", "pc", "pnc", "ac", "inv", "anc", "cv", "ventas", "un", "sin_inv")]
    e += [pct(j["mn"], 2)]
    c = R["caso"]; B, t, ER, r, aux, pr = c["B"], c["tot"], c["ER"], c["r"], c["aux"], c["pr"]
    e += [p(v) if v >= 0 else "−" + p(-v) for v in B.values()]
    e += [str(sum(C["u10"]["inventario_final"])) + " prendas"]
    e += [p(t[k]) for k in ("ac", "anc", "activo", "pc", "pnc", "pasivo", "patrimonio", "inv")]
    e += [p(ER[k]) for k in ("ventas", "cv", "mb", "adm", "ro", "gf", "uai", "var", "cv_real", "mb_real", "ro_real", "uai_real")]
    e += [veces(r["liquidez"]), veces(r["acida"]), p(r["ct"]), veces(r["d_p"]), pct(r["d_a"]), pct(r["lp"]),
          veces(r["rot_inv"]), fmt(r["dias_inv"], 1), veces(r["rot_cxc"]), fmt(r["dias_cxc"], 1), veces(r["rot_cxc_iva"]),
          fmt(r["dias_cxc_iva"], 1), pct(r["mb"]), pct(r["mn"]), pct(r["mb_std"]), pct(r["mn_std"]), pct(r["roa"]),
          pct(r["roe"]), fmt(r["cobertura"], 1), veces(r["acida_estricta"]), pct(r["roe_anual"], 0), pct(r["pc_prov"], 0),
          p(r["pago"]), veces(r["liq_pago"]), veces(r["dp_pago"]), p(r["banco_pago"])]
    e += [p(aux["otros_var"]), p(aux["prod_total"]), p(pr["cuota"]), p(pr["saldo12"]), fmt(aux["tinta_kg"], 3)]
    e += [p(C["u3"]["remanente"]), p(C["u4"]["total_mes"]), p(C["u5"]["arriendo_devengado"]), p(C["u5"]["seguro_devengado"]),
          p(C["u6"]["interes_mes_1"]), p(C["u10"]["cif_en_inventario"]), p(C["u9"]["cif_fijos_mes"])]
    return e


if __name__ == "__main__":
    m, C = cargar_caso()
    R = resultados(m, C)
    for k, d in R.items():
        if k == "caso":
            continue
        print(f"== {k}")
        for kk, v in d.items():
            if isinstance(v, dict):
                v = {a: (round(b, 4) if isinstance(b, float) else b) for a, b in v.items()}
            print(f"   {kk}: {round(v, 4) if isinstance(v, float) else v}")
    c = R["caso"]
    print("== caso · balance al 31 de marzo")
    for k, v in c["B"].items():
        print(f"   {k}: {fmt(v)}")
    for k in ("tot", "ER", "r", "aux", "pr"):
        print(f"== caso · {k}")
        for kk, v in c[k].items():
            print(f"   {kk}: {round(v, 4) if isinstance(v, float) else v}")
    ok_planilla = planilla(R)
    u8, ER = C["u8"], c["ER"]
    ok = all(abs(u8[a] - ER[b]) <= 0.5 for a, b in [("ventas", "ventas"), ("costo_ventas", "cv"), ("margen_bruto", "mb"),
                                                     ("resultado_operacional", "ro"), ("gastos_financieros", "gf"),
                                                     ("utilidad_antes_impuesto", "uai")])
    ok &= abs(u8["margen_neto"] - c["r"]["mn_std"]) < 1e-6 and abs(u8["margen_bruto_pct"] - c["r"]["mb_std"]) < 1e-6
    ok &= abs(C["u7"]["variacion_ppp"] - ER["var"]) < 1e-6 and abs(C["u7"]["saldo_ppp"] - c["B"]["Poleras lisas básicas"]) < 1e-6
    print("CADENA:", "OK, el estado de resultados y el balance calzan con caso_estampa_sur_contable.py" if ok else "NO calza")
    if not (ok and ok_planilla):
        sys.exit(1)
    if "--revisar" in sys.argv:
        md = (AQUI.parents[0] / "unidades" / "08-razones-financieras.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %")
        faltan = [x for x in esperados(R, C) if x not in md]
        print("REVISIÓN:", "OK, todas las cifras están en el texto" if not faltan else f"faltan {len(faltan)}: {faltan}")
        sys.exit(1 if faltan else 0)
