"""Unidad 3 · IVA.

Recalcula todas las cifras de la unidad desde los datos, genera la planilla
`planillas/u03_iva.xlsx` (con fórmulas) y la comprueba evaluando sus fórmulas en Python. Con
--revisar busca cada resultado y cada fila de las tablas generadas en el .md.

    py -3 -P u03_iva.py              # imprime resultados y tablas
    py -3 -P u03_iva.py --revisar    # además los busca en el .md

Todos los datos son originales de esta base (pasados por herramientas/colisiones.py). La única
cifra real es la tasa del IVA en Chile (19 %).
"""
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import importlib.util, re, sys

sys.stdout.reconfigure(encoding="utf-8")
AQUI = Path(__file__).resolve().parent
RAMO = AQUI.parent
IVA = 0.19


# ---------------------------------------------------------------- herramientas

def fmt(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def r0(x):
    """Redondeo a peso, mitad hacia arriba (como ROUND de Excel y como una factura)."""
    return int(Decimal(str(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def iva_de(neto):
    return r0(neto * IVA)


def neto_de_bruto(bruto):
    return r0(bruto / (1 + IVA))


def asiento(n, fecha, lineas, glosa):
    """lineas: [(cuenta, debe, haber)]. Comprueba que cuadra."""
    d = sum(x[1] for x in lineas)
    h = sum(x[2] for x in lineas)
    assert abs(d - h) < 1e-6, (n, d, h)
    return dict(n=n, fecha=fecha, lineas=lineas, glosa=glosa, total=d)


def md_asientos(asientos):
    filas = ["| N.º | Fecha | Cuenta | Debe | Haber |", "|---|---|---|---|---|"]
    for a in asientos:
        for k, (c, d, h) in enumerate(a["lineas"]):
            filas.append(f"| {a['n'] if k == 0 else ''} | {a['fecha'] if k == 0 else ''} | {c} | "
                         f"{fmt(d) if d else ''} | {fmt(h) if h else ''} |")
        filas.append(f"| | | *{a['glosa']}* | | |")
    return filas


def md_tabla(enc, filas):
    out = ["| " + " | ".join(enc) + " |", "|" + "---|" * len(enc)]
    out += ["| " + " | ".join(str(x) for x in f) + " |" for f in filas]
    return out


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · Ciclo Austral, taller de bicicletas, mayo de 2026
EJ = dict(repuestos=1_840_000, contado_repuestos=0.50, devolucion_bruta=107_100, soporte=620_000,
          ventas=3_260_000, contado_ventas=0.70, costo_vendido=1_130_000, luz_bruta=142_800)

# 3.1 · neto, IVA y bruto
E1 = dict(a_neto=46_000, b_bruto=83_300, c_bruto=1_249_500)
# 3.2 y 3.3 · librería: compra y venta
E2 = dict(compra=2_350_000, contado=0.30)
E3 = dict(fraccion=0.60, margen=0.40, contado=0.50)
# 3.4 · la cadena de un chaleco
E4 = dict(hilado=410_000, tejedora=1_160_000, tienda=1_870_000)
# 3.5 · devoluciones (sobre 3.2 y 3.3)
E5 = dict(dev_compra_bruta=71_400, dev_venta=0.05)
# 3.6 · liquidación de un almacén
E6 = dict(ventas=6_420_000, mercaderias=4_150_000, gas_luz=236_000, mantencion=310_000, sueldos=1_240_000)
# 3.7 · remanente que se usa al mes siguiente
E7 = dict(m1_ventas=2_940_000, m1_compras=1_310_000, elevador=5_640_000, m2_ventas=5_720_000, m2_compras=1_490_000)
# 3.8 · factura o boleta
E8 = dict(planchas=40, factura_neto=18_500, boleta_bruto=21_200)
# 3.9 · tres meses de una florería
E9 = dict(ventas=[1_820_000, 2_630_000, 4_960_000], compras=[1_160_000, 1_480_000, 2_140_000],
          camara=4_260_000)
# 3.10 · el precio del pan amasado
E10 = dict(precio_bruto=2_400, costo_neto=1_500, margen=0.40, kilos=3_150)


# ---------------------------------------------------------------- resultados

def resultados():
    R = {}

    # ----- ejemplo
    e = EJ
    iva1 = iva_de(e["repuestos"]); t1 = e["repuestos"] + iva1
    dev_neto = neto_de_bruto(e["devolucion_bruta"]); dev_iva = e["devolucion_bruta"] - dev_neto
    iva3 = iva_de(e["soporte"])
    iva4 = iva_de(e["ventas"]); t4 = e["ventas"] + iva4
    luz_neto = neto_de_bruto(e["luz_bruta"]); luz_iva = e["luz_bruta"] - luz_neto
    cf = iva1 - dev_iva + iva3 + luz_iva
    df = iva4
    pagar = df - cf
    A = [
        asiento(1, "04-05", [("Mercaderías (repuestos)", e["repuestos"], 0), ("IVA crédito fiscal", iva1, 0),
                             ("Banco", 0, r0(t1 * e["contado_repuestos"])),
                             ("Proveedores", 0, t1 - r0(t1 * e["contado_repuestos"]))],
                "Compra de repuestos, mitad por transferencia y mitad a 30 días"),
        asiento(2, "08-05", [("Proveedores", e["devolucion_bruta"], 0), ("Mercaderías (repuestos)", 0, dev_neto),
                             ("IVA crédito fiscal", 0, dev_iva)],
                "Devolución de repuestos fallados (nota de crédito del proveedor)"),
        asiento(3, "12-05", [("Herramientas y equipos", e["soporte"], 0), ("IVA crédito fiscal", iva3, 0),
                             ("Banco", 0, e["soporte"] + iva3)], "Compra de un soporte de taller por transferencia"),
        asiento(4, "20-05", [("Banco", r0(t4 * e["contado_ventas"]), 0),
                             ("Clientes", t4 - r0(t4 * e["contado_ventas"]), 0),
                             ("Ventas", 0, e["ventas"]), ("IVA débito fiscal", 0, iva4)],
                "Ventas del mes: 70 % al contado, 30 % a crédito a empresas en convenio"),
        asiento(5, "20-05", [("Costo de ventas", e["costo_vendido"], 0),
                             ("Mercaderías (repuestos)", 0, e["costo_vendido"])],
                "Costo de los repuestos vendidos (sin IVA)"),
        asiento(6, "28-05", [("Gasto en luz e internet", luz_neto, 0), ("IVA crédito fiscal", luz_iva, 0),
                             ("Banco", 0, e["luz_bruta"])], "Pago de luz e internet, IVA incluido"),
        asiento(7, "31-05", [("IVA débito fiscal", df, 0), ("IVA crédito fiscal", 0, cf),
                             ("IVA por pagar", 0, pagar)], "Liquidación del IVA de mayo"),
        asiento(8, "12-06", [("IVA por pagar", pagar, 0), ("Banco", 0, pagar)],
                "Pago del IVA de mayo en la declaración mensual"),
    ]
    R["ej"] = dict(iva1=iva1, t1=t1, banco1=r0(t1 * e["contado_repuestos"]), dev_neto=dev_neto, dev_iva=dev_iva,
                   iva3=iva3, t3=e["soporte"] + iva3, iva4=iva4, t4=t4, banco4=r0(t4 * e["contado_ventas"]),
                   clientes4=t4 - r0(t4 * e["contado_ventas"]), luz_neto=luz_neto, luz_iva=luz_iva,
                   cf_debe=iva1 + iva3 + luz_iva, cf=cf, df=df, pagar=pagar, asientos=A,
                   va=e["ventas"] - (e["repuestos"] - dev_neto + e["soporte"] + luz_neto),
                   margen_bruto=e["ventas"] - e["costo_vendido"])
    assert abs(R["ej"]["pagar"] - (df - cf)) < 1e-9

    # ----- 3.1
    e = E1
    a_iva = iva_de(e["a_neto"]); b_neto = neto_de_bruto(e["b_bruto"]); c_neto = neto_de_bruto(e["c_bruto"])
    R["e1"] = dict(a_iva=a_iva, a_bruto=e["a_neto"] + a_iva, b_neto=b_neto, b_iva=e["b_bruto"] - b_neto,
                   c_neto=c_neto, c_iva=e["c_bruto"] - c_neto, b_mal=r0(e["b_bruto"] * IVA),
                   b_dif=r0(e["b_bruto"] * IVA) - (e["b_bruto"] - b_neto))
    assert abs(b_neto * 1.19 - e["b_bruto"]) < 1 and abs(c_neto * 1.19 - e["c_bruto"]) < 1

    # ----- 3.2
    e = E2
    iva = iva_de(e["compra"]); t = e["compra"] + iva; banco = r0(t * e["contado"])
    A2 = [asiento(1, "02-04", [("Mercaderías", e["compra"], 0), ("IVA crédito fiscal", iva, 0),
                               ("Banco", 0, banco), ("Proveedores", 0, t - banco)],
                  "Compra de cuadernos, 30 % por transferencia y el resto a 60 días")]
    R["e2"] = dict(iva=iva, total=t, banco=banco, prov=t - banco, asientos=A2)

    # ----- 3.3
    e = E3
    costo = r0(E2["compra"] * e["fraccion"]); venta = r0(costo * (1 + e["margen"])); iva = iva_de(venta)
    t = venta + iva; banco = r0(t * e["contado"])
    A3 = [asiento(1, "15-04", [("Banco", banco, 0), ("Clientes", t - banco, 0), ("Ventas", 0, venta),
                               ("IVA débito fiscal", 0, iva)], "Venta del 60 % de los cuadernos, mitad al contado"),
          asiento(2, "15-04", [("Costo de ventas", costo, 0), ("Mercaderías", 0, costo)],
                  "Costo de los cuadernos vendidos")]
    R["e3"] = dict(costo=costo, venta=venta, iva=iva, total=t, banco=banco, clientes=t - banco, asientos=A3,
                   margen_bruto=venta - costo)

    # ----- 3.4
    e = E4
    etapas = [("Hilandería", 0, e["hilado"]), ("Tejedora", e["hilado"], e["tejedora"]),
              ("Tienda", e["tejedora"], e["tienda"])]
    filas = []
    for nombre, compra, venta in etapas:
        filas.append(dict(nombre=nombre, compra=compra, venta=venta, va=venta - compra, df=iva_de(venta),
                          cf=iva_de(compra), paga=iva_de(venta) - iva_de(compra)))
    total = sum(f["paga"] for f in filas)
    assert total == iva_de(e["tienda"])
    assert all(f["paga"] == iva_de(f["va"]) for f in filas)
    R["e4"] = dict(filas=filas, total=total, consumidor=e["tienda"] + iva_de(e["tienda"]))

    # ----- 3.5
    e = E5
    dn = neto_de_bruto(e["dev_compra_bruta"]); di = e["dev_compra_bruta"] - dn
    vn = r0(R["e3"]["venta"] * e["dev_venta"]); vi = iva_de(vn); vc = r0(R["e3"]["costo"] * e["dev_venta"])
    A5 = [asiento(1, "05-04", [("Proveedores", e["dev_compra_bruta"], 0), ("Mercaderías", 0, dn),
                               ("IVA crédito fiscal", 0, di)], "Devolución de cuadernos al proveedor"),
          asiento(2, "18-04", [("Ventas", vn, 0), ("IVA débito fiscal", vi, 0), ("Clientes", 0, vn + vi)],
                  "El cliente devuelve el 5 % de la venta del 15-04"),
          asiento(3, "18-04", [("Mercaderías", vc, 0), ("Costo de ventas", 0, vc)],
                  "Los cuadernos devueltos vuelven a bodega a su costo")]
    R["e5"] = dict(dn=dn, di=di, vn=vn, vi=vi, vt=vn + vi, vc=vc, asientos=A5,
                   cf_neto=R["e2"]["iva"] - di, df_neto=R["e3"]["iva"] - vi,
                   rem=(R["e2"]["iva"] - di) - (R["e3"]["iva"] - vi))
    assert R["e5"]["rem"] > 0

    # ----- 3.6
    e = E6
    df = iva_de(e["ventas"]); cfs = [iva_de(e[k]) for k in ("mercaderias", "gas_luz", "mantencion")]
    cf = sum(cfs); pagar = df - cf
    A6 = [asiento(1, "31-08", [("IVA débito fiscal", df, 0), ("IVA crédito fiscal", 0, cf),
                               ("IVA por pagar", 0, pagar)], "Liquidación del IVA de agosto"),
          asiento(2, "12-09", [("IVA por pagar", pagar, 0), ("Banco", 0, pagar)], "Pago del IVA de agosto")]
    R["e6"] = dict(df=df, cfs=cfs, cf=cf, pagar=pagar, asientos=A6,
                   mal_sueldos=pagar - iva_de(e["sueldos"]))

    # ----- 3.7
    e = E7
    df1 = iva_de(e["m1_ventas"]); cf1 = iva_de(e["m1_compras"]) + iva_de(e["elevador"]); rem1 = cf1 - df1
    df2 = iva_de(e["m2_ventas"]); cf2 = iva_de(e["m2_compras"]); antes = df2 - cf2; pagar2 = antes - rem1
    assert rem1 > 0 and pagar2 > 0
    A7 = [asiento(1, "31-07", [("IVA débito fiscal", df1, 0), ("IVA remanente", rem1, 0),
                               ("IVA crédito fiscal", 0, cf1)], "Liquidación de julio: queda remanente"),
          asiento(2, "31-08", [("IVA débito fiscal", df2, 0), ("IVA crédito fiscal", 0, cf2),
                               ("IVA remanente", 0, rem1), ("IVA por pagar", 0, pagar2)],
                  "Liquidación de agosto: se usa el remanente de julio")]
    R["e7"] = dict(df1=df1, cf1=cf1, iva_elevador=iva_de(e["elevador"]), rem1=rem1, df2=df2, cf2=cf2,
                   antes=antes, pagar2=pagar2, asientos=A7,
                   sin_elevador_1=df1 - iva_de(e["m1_compras"]))

    # ----- 3.8
    e = E8
    fac_bruto = e["factura_neto"] + iva_de(e["factura_neto"])
    R["e8"] = dict(fac_bruto=fac_bruto, costo_fac=e["planchas"] * e["factura_neto"],
                   costo_bol=e["planchas"] * e["boleta_bruto"],
                   ahorro=e["planchas"] * (e["boleta_bruto"] - e["factura_neto"]),
                   persona_fac=e["planchas"] * fac_bruto, persona_bol=e["planchas"] * e["boleta_bruto"],
                   persona_ahorro=e["planchas"] * (fac_bruto - e["boleta_bruto"]),
                   iva_fac=e["planchas"] * iva_de(e["factura_neto"]))

    # ----- 3.9
    e = E9
    meses, rem = [], 0
    for k in range(3):
        df = iva_de(e["ventas"][k]); cf = iva_de(e["compras"][k]) + (iva_de(e["camara"]) if k == 0 else 0)
        neto_mes = df - cf - rem
        pagar = max(neto_mes, 0); nuevo_rem = max(-neto_mes, 0)
        meses.append(dict(df=df, cf=cf, rem_ant=rem, pagar=pagar, rem=nuevo_rem))
        rem = nuevo_rem
    tot_df = sum(m["df"] for m in meses); tot_cf = sum(m["cf"] for m in meses)
    assert sum(m["pagar"] for m in meses) == tot_df - tot_cf
    sin_camara = sum(max(iva_de(v) - iva_de(c), 0) for v, c in zip(e["ventas"], e["compras"]))
    m3 = meses[2]
    A9 = [asiento(1, "31-05", [("IVA débito fiscal", m3["df"], 0), ("IVA crédito fiscal", 0, m3["cf"]),
                               ("IVA remanente", 0, m3["rem_ant"]), ("IVA por pagar", 0, m3["pagar"])],
                  "Liquidación de mayo: se agota el remanente")]
    R["e9"] = dict(meses=meses, tot_df=tot_df, tot_cf=tot_cf, iva_camara=iva_de(e["camara"]),
                   sin_camara=sin_camara, asientos=A9)

    # ----- 3.10
    e = E10
    neto = e["precio_bruto"] / (1 + IVA)
    precio_obj = r0(e["costo_neto"] * (1 + e["margen"]) * (1 + IVA))
    neto_obj = precio_obj / (1 + IVA)
    df = r0(e["kilos"] * precio_obj * IVA / (1 + IVA)); cf = r0(e["kilos"] * e["costo_neto"] * IVA)
    R["e10"] = dict(neto=neto, iva_kilo=e["precio_bruto"] - neto, margen_real=neto - e["costo_neto"],
                    margen_creido=e["precio_bruto"] - e["costo_neto"],
                    margen_pct=(neto - e["costo_neto"]) / e["costo_neto"],
                    precio_obj=precio_obj, neto_obj=neto_obj, df=df, cf=cf, pagar=df - cf,
                    va=e["kilos"] * (neto_obj - e["costo_neto"]))
    assert abs(R["e10"]["pagar"] - R["e10"]["va"] * IVA) < 1

    R["caso"] = caso()
    return R


# ---------------------------------------------------------------- caso integrador

def cargar_caso():
    spec = importlib.util.spec_from_file_location("caso", AQUI / "caso_estampa_sur_contable.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def caso():
    m = cargar_caso()
    C = m.calcular()
    # recálculo propio desde los datos del caso
    ventas = sum(q * p for q, p in zip(m.VENDIDAS, m.PRECIO))
    tinta_kg = sum(q * g for q, g in zip(m.PRODUCIDAS, m.TINTA_G)) / 1000 - 12 + m.TINTA_SOBRANTE_KG
    lineas = [(1, "Poleras lisas básicas", m.COMPRA_POLERAS["unidades"] * m.COMPRA_POLERAS["precio"]),
              (2, "Otras prendas lisas", sum(q * c for q, c in zip(m.PRODUCIDAS[1:], m.PRENDA_LISA[1:]))),
              (3, "Tinta", r0(tinta_kg * m.TINTA_KG)),   # a peso, la mitad hacia arriba, como el script
              (4, "Segunda máquina", m.MAQUINA["neto"]),
              (9, "Seguro anual", m.SEGURO_ANUAL),
              (12, "Otros gastos de administración", m.OTROS_FIJOS)]
    # cada documento se redondea a peso (neto e IVA), igual que el script
    debito = r0(ventas * IVA)
    credito = sum(r0(n * IVA) for _, _, n in lineas)
    remanente = credito - debito
    sin_maquina = debito - (credito - r0(m.MAQUINA["neto"] * IVA))
    abril_hip = sin_maquina - remanente
    return dict(C=C["u3"], asientos=C["u2"]["asientos"], ventas=ventas, debito=debito, credito=credito, remanente=remanente,
                lineas=lineas, tinta_kg=tinta_kg,
                iva_maquina=m.MAQUINA["neto"] * IVA, sin_maquina=sin_maquina, abril_hip=abril_hip,
                recuperado_marzo=m.MAQUINA["neto"] * IVA - remanente)


def chequear_caso(R):
    c = R["caso"]
    ok = (abs(c["C"]["debito"] - c["debito"]) < 1e-6 and abs(c["C"]["credito"] - c["credito"]) < 1e-6
          and abs(c["C"]["remanente"] - c["remanente"]) < 1e-6 and c["C"]["a_pagar"] == 0
          and all(abs(a["debe"].get("IVA crédito fiscal", 0) - r0(n * IVA)) < 1e-6
                  for a in c["asientos"] for k, _, n in c["lineas"] if a["n"] == k))
    return ok


# ---------------------------------------------------------------- tablas para el .md

def tablas(R):
    T = {}
    j = R["ej"]
    T["ej_diario"] = md_asientos(j["asientos"])
    T["ej_t"] = md_tabla(["IVA crédito fiscal · Debe", "IVA crédito fiscal · Haber", "IVA débito fiscal · Debe",
                          "IVA débito fiscal · Haber"],
                         [[f"(1) {fmt(j['iva1'])}", f"(2) {fmt(j['dev_iva'])}", "", f"(4) {fmt(j['iva4'])}"],
                          [f"(3) {fmt(j['iva3'])}", "", "", ""],
                          [f"(6) {fmt(j['luz_iva'])}", "", "", ""],
                          [f"Suma {fmt(j['cf_debe'])}", f"Suma {fmt(j['dev_iva'])}", "", f"Suma {fmt(j['df'])}"],
                          [f"**Saldo deudor {fmt(j['cf'])}**", "", "", f"**Saldo acreedor {fmt(j['df'])}**"]])
    T["e2"] = md_asientos(R["e2"]["asientos"])
    T["e3"] = md_asientos(R["e3"]["asientos"])
    T["e4"] = md_tabla(["Etapa", "Compra neta", "Venta neta", "Valor agregado", "Débito", "Crédito", "Paga al fisco"],
                       [[f["nombre"], fmt(f["compra"]), fmt(f["venta"]), fmt(f["va"]), fmt(f["df"]), fmt(f["cf"]),
                         fmt(f["paga"])] for f in R["e4"]["filas"]]
                       + [["**Total**", "", "", fmt(sum(f["va"] for f in R["e4"]["filas"])), "", "",
                           f"**{fmt(R['e4']['total'])}**"]])
    T["e5"] = md_asientos(R["e5"]["asientos"])
    T["e6"] = md_asientos(R["e6"]["asientos"])
    T["e7"] = md_asientos(R["e7"]["asientos"])
    T["e9"] = md_tabla(["Mes", "Débito", "Crédito del mes", "Remanente que llega", "IVA a pagar", "Remanente que sigue"],
                       [[mes, fmt(m["df"]), fmt(m["cf"]), fmt(m["rem_ant"]), fmt(m["pagar"]), fmt(m["rem"])]
                        for mes, m in zip(["Marzo", "Abril", "Mayo"], R["e9"]["meses"])])
    T["e9_a"] = md_asientos(R["e9"]["asientos"])
    c = R["caso"]
    T["caso"] = md_tabla(["Asiento", "Compra", "Neto", "IVA crédito fiscal"],
                         [[str(n), nombre, fmt(neto), fmt(r0(neto * IVA))] for n, nombre, neto in c["lineas"]]
                         + [["", "**Crédito fiscal de marzo**", "", f"**{fmt(c['credito'])}**"],
                            ["6", "Ventas del mes (débito fiscal)", fmt(c["ventas"]), f"**{fmt(c['debito'])}**"]])
    T["caso_a"] = md_asientos([asiento("Liq.", "31-03", [("IVA débito fiscal", r0(c["debito"]), 0),
                                                     ("IVA remanente", r0(c["remanente"]), 0),
                                                     ("IVA crédito fiscal", 0, r0(c["credito"]))],
                                       "Liquidación del IVA de marzo: queda remanente")])
    return T


# ---------------------------------------------------------------- planilla (con fórmulas)

def evaluar(ws):
    """Evalúa las fórmulas de una hoja de openpyxl (solo + − × ÷, SUM, MAX, MIN, ROUND, ABS)."""
    cache = {}

    def val(ref):
        ref = ref.replace("$", "")
        if ref in cache:
            return cache[ref]
        v = ws[ref].value
        if isinstance(v, str) and v.startswith("="):
            v = calc(v[1:])
        cache[ref] = 0 if v is None else v
        return cache[ref]

    def rng(c1, f1, c2, f2):
        from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
        return [val(f"{gl(c)}{f}") for c in range(ci(c1), ci(c2) + 1) for f in range(f1, f2 + 1)]

    def _round(x, d=0):
        q = Decimal(1).scaleb(-int(d))
        return float(Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP))

    def _sum(*a):
        return sum(sum(x) if isinstance(x, list) else x for x in a)

    def calc(f):
        f = re.sub(r"\$?([A-Z]{1,2})\$?(\d+):\$?([A-Z]{1,2})\$?(\d+)", r'rng("\1",\2,"\3",\4)', f)
        f = re.sub(r"(?<![A-Za-z\"])\$?([A-Z]{1,2})\$?(\d+)(?![\d\"])", r'val("\1\2")', f)
        f = f.replace("SUM(", "_sum(").replace("ROUND(", "_round(").replace("MAX(", "max(") \
             .replace("MIN(", "min(").replace("ABS(", "abs(")
        return eval(f, {"val": val, "rng": rng, "_sum": _sum, "_round": _round, "max": max, "min": min, "abs": abs})

    return val


def planilla(R):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    negrita, gris = Font(bold=True), PatternFill("solid", fgColor="F2F1EF")
    wb = Workbook()

    # Hoja 1 · Datos y libro diario del ejemplo
    ws = wb.active; ws.title = "Ejemplo"
    e = EJ
    datos = [("Tasa de IVA", IVA), ("Repuestos, neto", e["repuestos"]), ("Parte al contado de la compra", e["contado_repuestos"]),
             ("Devolución, IVA incluido", e["devolucion_bruta"]), ("Soporte de taller, neto", e["soporte"]),
             ("Ventas del mes, neto", e["ventas"]), ("Parte al contado de las ventas", e["contado_ventas"]),
             ("Costo de lo vendido", e["costo_vendido"]), ("Luz e internet, IVA incluido", e["luz_bruta"])]
    ws["A1"] = "Ciclo Austral · IVA de mayo (datos en azul; todo lo demás es fórmula)"; ws["A1"].font = negrita
    for k, (n, v) in enumerate(datos, start=3):
        ws.cell(k, 1, n); c = ws.cell(k, 2, v); c.font = Font(color="1F4E9E")
    # B3 tasa, B4 repuestos, B5 contado, B6 dev bruta, B7 soporte, B8 ventas, B9 contado v, B10 costo, B11 luz
    ws["A13"] = "Libro diario"; ws["A13"].font = negrita
    enc = ["N.º", "Fecha", "Cuenta", "Debe", "Haber"]
    for k, h in enumerate(enc, start=1):
        c = ws.cell(14, k, h); c.font = negrita; c.fill = gris
    filas = [
        (1, "04-05", "Mercaderías (repuestos)", "=B4", None),
        (None, None, "IVA crédito fiscal", "=ROUND(B4*B3,0)", None),
        (None, None, "Banco", None, "=ROUND((B4+ROUND(B4*B3,0))*B5,0)"),
        (None, None, "Proveedores", None, "=B4+ROUND(B4*B3,0)-E17"),
        (2, "08-05", "Proveedores", "=B6", None),
        (None, None, "Mercaderías (repuestos)", None, "=ROUND(B6/(1+B3),0)"),
        (None, None, "IVA crédito fiscal", None, "=B6-E20"),
        (3, "12-05", "Herramientas y equipos", "=B7", None),
        (None, None, "IVA crédito fiscal", "=ROUND(B7*B3,0)", None),
        (None, None, "Banco", None, "=D22+D23"),
        (4, "20-05", "Banco", "=ROUND((B8+ROUND(B8*B3,0))*B9,0)", None),
        (None, None, "Clientes", "=B8+ROUND(B8*B3,0)-D25", None),
        (None, None, "Ventas", None, "=B8"),
        (None, None, "IVA débito fiscal", None, "=ROUND(B8*B3,0)"),
        (5, "20-05", "Costo de ventas", "=B10", None),
        (None, None, "Mercaderías (repuestos)", None, "=B10"),
        (6, "28-05", "Gasto en luz e internet", "=ROUND(B11/(1+B3),0)", None),
        (None, None, "IVA crédito fiscal", "=B11-D31", None),
        (None, None, "Banco", None, "=B11"),
        (7, "31-05", "IVA débito fiscal", "=E28", None),
        (None, None, "IVA crédito fiscal", None, "=D16-E21+D23+D32"),
        (None, None, "IVA por pagar", None, "=D34-E35"),
        (8, "12-06", "IVA por pagar", "=E36", None),
        (None, None, "Banco", None, "=D37"),
    ]
    for k, f in enumerate(filas, start=15):
        for col, v in enumerate(f, start=1):
            if v is not None:
                ws.cell(k, col, v)
    fin = 15 + len(filas)
    ws.cell(fin, 3, "Sumas").font = negrita
    ws.cell(fin, 4, f"=SUM(D15:D{fin - 1})").font = negrita
    ws.cell(fin, 5, f"=SUM(E15:E{fin - 1})").font = negrita
    ws.cell(fin + 1, 3, "Diferencia (debe dar 0)")
    ws.cell(fin + 1, 4, f"=D{fin}-E{fin}")
    ws["G14"] = "Liquidación"; ws["G14"].font = negrita
    ws["G15"], ws["H15"] = "Débito fiscal", "=E28"
    ws["G16"], ws["H16"] = "Crédito fiscal", "=D16-E21+D23+D32"
    ws["G17"], ws["H17"] = "IVA por pagar", "=MAX(H15-H16,0)"
    ws["G18"], ws["H18"] = "Remanente", "=MAX(H16-H15,0)"
    for col, w in zip("ABCDEFGH", [34, 9, 28, 14, 14, 3, 16, 14]):
        ws.column_dimensions[col].width = w

    # Hoja 2 · Tres meses de remanente (ejercicio 3.9)
    w2 = wb.create_sheet("Remanente 3.9")
    w2["A1"] = "Florería · tres meses de IVA (datos en azul)"; w2["A1"].font = negrita
    enc = ["Mes", "Ventas netas", "Compras netas", "Activo fijo neto", "Débito", "Crédito", "Remanente que llega",
           "IVA a pagar", "Remanente que sigue"]
    for k, h in enumerate(enc, start=1):
        c = w2.cell(3, k, h); c.font = negrita; c.fill = gris
    w2["K3"], w2["L3"] = "Tasa", IVA
    for k, mes in enumerate(["Marzo", "Abril", "Mayo"]):
        f = 4 + k
        w2.cell(f, 1, mes)
        for col, v in zip((2, 3, 4), (E9["ventas"][k], E9["compras"][k], E9["camara"] if k == 0 else 0)):
            w2.cell(f, col, v).font = Font(color="1F4E9E")
        w2.cell(f, 5, f"=ROUND(B{f}*$L$3,0)")
        w2.cell(f, 6, f"=ROUND(C{f}*$L$3,0)+ROUND(D{f}*$L$3,0)")
        w2.cell(f, 7, 0 if k == 0 else f"=I{f - 1}")
        w2.cell(f, 8, f"=MAX(E{f}-F{f}-G{f},0)")
        w2.cell(f, 9, f"=MAX(F{f}+G{f}-E{f},0)")
    w2["A8"], w2["H8"] = "Total pagado en los tres meses", "=SUM(H4:H6)"
    w2["A9"], w2["H9"] = "Débito total menos crédito total", "=SUM(E4:E6)-SUM(F4:F6)"
    for col in "ABCDEFGHI":
        w2.column_dimensions[col].width = 16

    # Hoja 3 · Caso integrador
    w3 = wb.create_sheet("Caso marzo")
    c = R["caso"]
    w3["A1"] = "Estampa Sur · IVA de marzo (netos del caso integrador)"; w3["A1"].font = negrita
    for k, h in enumerate(["Asiento", "Operación", "Neto", "IVA"], start=1):
        x = w3.cell(3, k, h); x.font = negrita; x.fill = gris
    w3["F3"], w3["G3"] = "Tasa", IVA
    for k, (n, nombre, neto) in enumerate(c["lineas"]):
        f = 4 + k
        w3.cell(f, 1, n); w3.cell(f, 2, nombre); w3.cell(f, 3, neto).font = Font(color="1F4E9E")
        w3.cell(f, 4, f"=ROUND(C{f}*$G$3,0)")
    f = 4 + len(c["lineas"])
    w3.cell(f, 2, "Crédito fiscal").font = negrita; w3.cell(f, 4, f"=SUM(D4:D{f - 1})")
    w3.cell(f + 1, 1, 6); w3.cell(f + 1, 2, "Ventas (débito fiscal)")
    w3.cell(f + 1, 3, c["ventas"]).font = Font(color="1F4E9E"); w3.cell(f + 1, 4, f"=ROUND(C{f + 1}*$G$3,0)")
    w3.cell(f + 2, 2, "IVA por pagar"); w3.cell(f + 2, 4, f"=MAX(D{f + 1}-D{f},0)")
    w3.cell(f + 3, 2, "Remanente para abril"); w3.cell(f + 3, 4, f"=MAX(D{f}-D{f + 1},0)")
    w3.column_dimensions["B"].width = 32
    for col in "CD":
        w3.column_dimensions[col].width = 16

    destino = RAMO / "planillas" / "u03_iva.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)

    # comprobación: se evalúan las fórmulas y se comparan con Python
    v1, v2, v3 = evaluar(ws), evaluar(w2), evaluar(w3)
    j = R["ej"]
    chk = {
        "planilla: el diario cuadra": abs(v1(f"D{fin + 1}")) < 1e-6 and abs(v1(f"D{fin}") - sum(a["total"] for a in j["asientos"])) < 1e-6,
        "planilla: débito y crédito del ejemplo": v1("H15") == j["df"] and v1("H16") == j["cf"],
        "planilla: IVA por pagar del ejemplo": v1("H17") == j["pagar"] and v1("H18") == 0,
        "planilla: los tres meses de 3.9": all(v2(f"H{4 + k}") == m["pagar"] and v2(f"I{4 + k}") == m["rem"]
                                               for k, m in enumerate(R["e9"]["meses"])),
        "planilla: lo pagado en 3.9 = débito − crédito": v2("H8") == v2("H9"),
        "planilla: remanente del caso": abs(v3(f"D{f + 3}") - c["remanente"]) < 1e-6 and v3(f"D{f + 2}") == 0,
    }
    return destino, chk


# ---------------------------------------------------------------- esperados

def esperados(R, T):
    p = fmt
    e = []
    j = R["ej"]
    e += [p(j[k]) for k in ("iva1", "t1", "banco1", "dev_neto", "dev_iva", "iva3", "t3", "iva4", "t4", "banco4",
                            "clientes4", "luz_neto", "luz_iva", "cf_debe", "cf", "df", "pagar", "va", "margen_bruto")]
    j = R["e1"]; e += [p(j[k]) for k in ("a_iva", "a_bruto", "b_neto", "b_iva", "c_neto", "c_iva", "b_mal", "b_dif")]
    j = R["e2"]; e += [p(j[k]) for k in ("iva", "total", "banco", "prov")]
    j = R["e3"]; e += [p(j[k]) for k in ("costo", "venta", "iva", "total", "banco", "clientes", "margen_bruto")]
    j = R["e4"]; e += [p(j["total"]), p(j["consumidor"])]
    j = R["e5"]; e += [p(j[k]) for k in ("dn", "di", "vn", "vi", "vt", "vc", "cf_neto", "df_neto", "rem")]
    j = R["e6"]; e += [p(j["df"]), p(j["cf"]), p(j["pagar"]), p(j["mal_sueldos"])] + [p(x) for x in j["cfs"]]
    j = R["e7"]; e += [p(j[k]) for k in ("df1", "cf1", "iva_elevador", "rem1", "df2", "cf2", "antes", "pagar2", "sin_elevador_1")]
    j = R["e8"]; e += [p(j[k]) for k in ("fac_bruto", "costo_fac", "costo_bol", "ahorro", "persona_fac",
                                         "persona_bol", "persona_ahorro", "iva_fac")]
    j = R["e9"]; e += [p(j["tot_df"]), p(j["tot_cf"]), p(j["iva_camara"]), p(j["sin_camara"])]
    j = R["e10"]; e += [p(j["neto"], 2), p(j["iva_kilo"], 2), p(j["margen_real"], 2), p(j["margen_creido"]),
                        fmt(100 * j["margen_pct"], 1) + " %", p(j["precio_obj"]), p(j["neto_obj"]), p(j["df"]),
                        p(j["cf"]), p(j["pagar"]), p(j["va"])]
    c = R["caso"]
    e += [p(c["debito"]), p(c["credito"]), p(c["remanente"]), p(c["ventas"]), p(c["iva_maquina"]),
          p(c["sin_maquina"]), p(c["abril_hip"]), p(c["recuperado_marzo"])]
    for filas in T.values():
        e += filas
    return e


# ---------------------------------------------------------------- principal

if __name__ == "__main__":
    R = resultados()
    T = tablas(R)
    for k, v in R.items():
        if k == "caso":
            continue
        print(f"== {k}")
        for kk, vv in v.items():
            if kk == "asientos":
                continue
            print(f"   {kk}: {vv}")
    c = R["caso"]
    print("== caso", {k: v for k, v in c.items() if k not in ("C",)})
    for k, filas in T.items():
        print(f"\n--- tabla {k}")
        print("\n".join(filas))
    destino, chk = planilla(R)
    chk["caso: débito, crédito y remanente calzan con el script"] = chequear_caso(R)
    print(f"\nPlanilla: {destino}")
    print("CHEQUEOS:")
    for k, v in chk.items():
        print(f"   {'OK ' if v else 'MAL'} {k}")
    if not all(chk.values()):
        sys.exit(1)
    print("CADENA: OK, la pieza del caso calza con caso_estampa_sur_contable.py")
    if "--revisar" in sys.argv:
        md = (RAMO / "unidades" / "03-iva.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %")
        faltan = [x for x in esperados(R, T) if x not in md]
        print("\nREVISIÓN:", "OK, todas las cifras están en el texto" if not faltan else f"faltan {len(faltan)}: {faltan}")
        sys.exit(1 if faltan else 0)
