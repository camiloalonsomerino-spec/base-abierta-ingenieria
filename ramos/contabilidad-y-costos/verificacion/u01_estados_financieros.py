"""Unidad 1 · La contabilidad y los estados financieros.

Recalcula todas las cifras de la unidad desde los datos, genera la planilla
`planillas/u01_estados_financieros.xlsx` (con fórmulas) y comprueba sus resultados evaluando esas
fórmulas en Python. Con --revisar busca cada resultado escrito tal cual en el .md.

    py -3 -P u01_estados_financieros.py            # imprime los resultados y genera la planilla
    py -3 -P u01_estados_financieros.py --revisar  # además los busca en el .md

Todos los datos son originales de esta base (pasados por herramientas/colisiones.py). No vienen de
ningún libro ni curso.
"""
from pathlib import Path
import importlib.util, re, sys

sys.stdout.reconfigure(encoding="utf-8")
AQUI = Path(__file__).resolve().parent
RAMO = AQUI.parent


# ---------------------------------------------------------------- herramientas

def fmt(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, dec=1):
    return fmt(100 * x, dec) + " %"


# Tipos de cuenta: A activo, CA complementaria de activo (resta), P pasivo, PAT patrimonio,
# RET retiros o dividendos (restan patrimonio), I ingreso, G gasto.
def efecto(mov, tipo):
    """Efecto de un movimiento {cuenta: delta} en activo, pasivo y patrimonio."""
    dA = sum(v for k, v in mov.items() if tipo[k] == "A") - sum(v for k, v in mov.items() if tipo[k] == "CA")
    dP = sum(v for k, v in mov.items() if tipo[k] == "P")
    dPat = (sum(v for k, v in mov.items() if tipo[k] in ("PAT", "I"))
            - sum(v for k, v in mov.items() if tipo[k] in ("G", "RET")))
    return dA, dP, dPat


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · la bicicletería Rueda Libre, abril de 2026 (sin IVA ni depreciación)
TIPO_EJ = {"Banco": "A", "Clientes": "A", "Mercaderías": "A", "Arriendo pagado por anticipado": "A",
           "Equipos y herramientas": "A", "Proveedores": "P", "Cuentas por pagar": "P",
           "Ingreso percibido por adelantado": "P", "Capital": "PAT", "Retiros del dueño": "RET",
           "Ventas de repuestos": "I", "Ingresos por reparaciones": "I", "Costo de ventas": "G",
           "Sueldos": "G", "Arriendo": "G", "Electricidad": "G"}
EJ = [  # (letra, fecha, glosa, movimientos, flujo de caja: O operación, I inversión, F financiamiento)
    ("a", "01-04-2026", "Matías aporta $6.300.000 al banco", {"Banco": 6_300_000, "Capital": 6_300_000}, "F"),
    ("b", "02-04-2026", "Compra herramientas y un banco de trabajo", {"Equipos y herramientas": 1_450_000, "Banco": -1_450_000}, "I"),
    ("c", "03-04-2026", "Compra repuestos: 40 % al contado y el resto a 30 días",
     {"Mercaderías": 2_150_000, "Banco": -860_000, "Proveedores": 1_290_000}, "O"),
    ("d", "05-04-2026", "Paga el arriendo de abril y mayo",
     {"Banco": -670_000, "Arriendo": 335_000, "Arriendo pagado por anticipado": 335_000}, "O"),
    ("e", "15-04-2026", "Vende repuestos que costaron $1.240.000 con 60 % de recargo; mitad al contado",
     {"Banco": 992_000, "Clientes": 992_000, "Ventas de repuestos": 1_984_000,
      "Costo de ventas": 1_240_000, "Mercaderías": -1_240_000}, "O"),
    ("f", "15-04-2026", "Cobra reparaciones al contado", {"Banco": 730_000, "Ingresos por reparaciones": 730_000}, "O"),
    ("g", "20-04-2026", "Un cliente paga por adelantado una mantención de mayo",
     {"Banco": 265_000, "Ingreso percibido por adelantado": 265_000}, "O"),
    ("h", "30-04-2026", "Paga el sueldo del ayudante", {"Banco": -560_000, "Sueldos": 560_000}, "O"),
    ("i", "30-04-2026", "Llega la cuenta de luz de abril; se paga en mayo", {"Electricidad": 48_000, "Cuentas por pagar": 48_000}, None),
    ("j", "30-04-2026", "Matías retira dinero para uso personal", {"Banco": -185_000, "Retiros del dueño": 185_000}, "F"),
]
EJ_RECARGO = 0.60

# 1.2 · ecuación contable con incógnitas
E2 = {"A": {"activo": 37_400_000, "pasivo": 15_900_000},
      "B": {"pasivo": 8_150_000, "patrimonio": 22_600_000},
      "C": {"activo": 52_000_000, "pasivo": 31_500_000, "d_activo": 6_300_000, "d_pasivo": -2_750_000}}

# 1.3 · peluquería: efecto de seis transacciones
TIPO_E3 = {"Banco": "A", "Muebles y equipos": "A", "Proveedores": "P", "Ingreso percibido por adelantado": "P",
           "Capital": "PAT", "Ingresos por servicios": "I", "Arriendo": "G"}
E3 = [("1", {"Banco": 4_500_000, "Capital": 4_500_000}),
      ("2", {"Muebles y equipos": 1_680_000, "Proveedores": 1_680_000}),
      ("3", {"Banco": 395_000, "Ingresos por servicios": 395_000}),
      ("4", {"Banco": -600_000, "Proveedores": -600_000}),
      ("5", {"Banco": -280_000, "Arriendo": 280_000}),
      ("6", {"Banco": 90_000, "Ingreso percibido por adelantado": 90_000})]

# 1.4 · estudio de diseño en mayo: devengado contra caja
E4 = {"facturado_mayo": 3_150_000, "cobrado_de_mayo": 1_900_000, "cobro_de_abril": 700_000,
      "sueldos_mayo": 1_240_000, "arriendo_mensual": 410_000, "internet_mayo": 36_000}

# 1.5 · ferretería: balance clasificado con el capital como incógnita
E5 = {"corriente": {"Caja": 185_000, "Banco": 3_920_000, "Clientes": 2_640_000, "Mercaderías": 7_350_000,
                    "Seguro pagado por anticipado": 270_000},
      "no_corriente": {"Terreno": 18_500_000, "Vehículo": 11_800_000, "Depreciación acumulada del vehículo": -2_950_000},
      "pasivo_c": {"Proveedores": 4_110_000, "Sueldos por pagar": 960_000, "IVA por pagar": 515_000,
                   "Préstamo bancario de corto plazo": 1_800_000},
      "pasivo_nc": {"Préstamo bancario de largo plazo": 9_600_000},
      "patrimonio_conocido": {"Utilidades retenidas": 4_380_000, "Resultado del ejercicio": 1_735_000}}

# 1.6 · distribuidora: estado de resultados clasificado (sin impuesto a la renta)
E6 = {"ventas": 84_600_000, "costo_ventas": 52_450_000,
      "operacionales": {"Sueldos de administración": 9_720_000, "Publicidad": 1_380_000, "Arriendo de oficina": 3_600_000,
                        "Depreciación del ejercicio": 2_150_000, "Distribución": 4_270_000},
      "intereses": 1_940_000, "utilidad_venta_vehiculo": 860_000,
      "distractores": {"Clientes": 9_300_000, "Préstamo bancario": 18_000_000, "Capital": 25_000_000,
                       "Arriendo pagado por anticipado": 600_000}}

# 1.7 · del balance inicial al final con aportes y dividendos
E7 = {"A0": 63_400_000, "P0": 27_500_000, "A1": 71_900_000, "P1": 26_300_000, "aportes": 4_700_000, "dividendos": 3_250_000}

# 1.8 · estado de flujo de efectivo de un mes
E8 = {"saldo_inicial": 2_380_000,
      "mov": [("Cobros a clientes", 12_640_000, "O"), ("Pagos a proveedores", -7_910_000, "O"),
              ("Sueldos pagados", -2_450_000, "O"), ("Intereses pagados", -115_000, "O"),
              ("Compra de una camioneta", -9_200_000, "I"), ("Venta de una máquina usada", 1_260_000, "I"),
              ("Préstamo recibido", 5_800_000, "F"), ("Amortización del préstamo antiguo", -750_000, "F"),
              ("Dividendos pagados", -870_000, "F")]}

# 1.9 · utilidad con la caja en rojo
E9 = {"ventas": 18_500_000, "contado": 0.25, "costo_ventas": 11_100_000, "compras_contado": 14_800_000,
      "gastos_pagados": 4_250_000}

# 1.10 · balance con errores de principio
E10 = {"utilidad": 31_500_000, "activo": 148_000_000, "pasivo": 52_000_000,
       "terreno_costo": 43_500_000, "terreno_tasacion": 61_000_000,
       "seguro_anual": 1_440_000, "meses_consumidos": 1,
       "anticipo_cliente": 2_300_000, "retiro_dueno": 760_000, "dep_acumulada": 5_600_000}


# ---------------------------------------------------------------- resultados

def resultados():
    R = {}

    # --- ejemplo
    saldo, filas, acum = {}, [], [0, 0, 0]
    for letra, fecha, glosa, mov, _ in EJ:
        for k, v in mov.items():
            saldo[k] = saldo.get(k, 0) + v
        dA, dP, dPat = efecto(mov, TIPO_EJ)
        assert dA == dP + dPat, letra
        acum = [acum[0] + dA, acum[1] + dP, acum[2] + dPat]
        filas.append(dict(letra=letra, dA=dA, dP=dP, dPat=dPat, A=acum[0], P=acum[1], Pat=acum[2]))
    ing = saldo["Ventas de repuestos"] + saldo["Ingresos por reparaciones"]
    margen_bruto = ing - saldo["Costo de ventas"]
    gastos_op = saldo["Sueldos"] + saldo["Arriendo"] + saldo["Electricidad"]
    utilidad = margen_bruto - gastos_op
    flujo = {t: sum(m.get("Banco", 0) for *_, m, f in EJ if f == t) for t in "OIF"}
    cobros = sum(m["Banco"] for *_, m, f in EJ if f == "O" and m.get("Banco", 0) > 0)
    pagos = -sum(m["Banco"] for *_, m, f in EJ if f == "O" and m.get("Banco", 0) < 0)
    activos = {k: v for k, v in saldo.items() if TIPO_EJ[k] == "A"}
    pasivos = {k: v for k, v in saldo.items() if TIPO_EJ[k] == "P"}
    pat_final = saldo["Capital"] + utilidad - saldo["Retiros del dueño"]
    assert abs(saldo["Ventas de repuestos"] - saldo["Costo de ventas"] * (1 + EJ_RECARGO)) < 1e-6
    caja_resultado = cobros - pagos  # ingresos y gastos mirados con la caja (solo operación)
    R["ej"] = dict(filas=filas, saldo=saldo, ingresos=ing, margen_bruto=margen_bruto, gastos_op=gastos_op,
                   utilidad=utilidad, flujo=flujo, cobros=cobros, pagos=pagos, variacion=sum(flujo.values()),
                   activo=sum(activos.values()), pasivo=sum(pasivos.values()), pat_final=pat_final,
                   activo_corriente=sum(v for k, v in activos.items() if k != "Equipos y herramientas"),
                   caja_resultado=caja_resultado, sin_cobrar=saldo["Clientes"], en_bodega=saldo["Mercaderías"])
    assert R["ej"]["activo"] == R["ej"]["pasivo"] + pat_final
    assert R["ej"]["variacion"] == saldo["Banco"]
    # puente entre utilidad y flujo operacional
    puente = (utilidad - saldo["Clientes"] - saldo["Mercaderías"] - saldo["Arriendo pagado por anticipado"]
              + saldo["Proveedores"] + saldo["Cuentas por pagar"] + saldo["Ingreso percibido por adelantado"])
    assert puente == flujo["O"]
    R["ej"]["puente"] = puente

    # --- 1.2
    e = E2
    pat_A = e["A"]["activo"] - e["A"]["pasivo"]
    act_B = e["B"]["pasivo"] + e["B"]["patrimonio"]
    pat_C0 = e["C"]["activo"] - e["C"]["pasivo"]
    pat_C1 = (e["C"]["activo"] + e["C"]["d_activo"]) - (e["C"]["pasivo"] + e["C"]["d_pasivo"])
    R["e2"] = dict(pat_A=pat_A, endeuda_A=e["A"]["pasivo"] / e["A"]["activo"], act_B=act_B,
                   pat_C0=pat_C0, pat_C1=pat_C1, d_pat_C=pat_C1 - pat_C0,
                   A_C1=e["C"]["activo"] + e["C"]["d_activo"], P_C1=e["C"]["pasivo"] + e["C"]["d_pasivo"])

    # --- 1.3
    acum, filas3 = [0, 0, 0], []
    for n, mov in E3:
        dA, dP, dPat = efecto(mov, TIPO_E3)
        assert dA == dP + dPat
        acum = [acum[0] + dA, acum[1] + dP, acum[2] + dPat]
        filas3.append(dict(n=n, dA=dA, dP=dP, dPat=dPat, A=acum[0], P=acum[1], Pat=acum[2]))
    R["e3"] = dict(filas=filas3)

    # --- 1.4
    e = E4
    gastos_dev = e["sueldos_mayo"] + e["arriendo_mensual"] + e["internet_mayo"]
    ent = e["cobrado_de_mayo"] + e["cobro_de_abril"]
    sal = e["sueldos_mayo"] + e["arriendo_mensual"]
    R["e4"] = dict(gastos_dev=gastos_dev, util_dev=e["facturado_mayo"] - gastos_dev, entradas=ent, salidas=sal,
                   neto_caja=ent - sal, por_cobrar=e["facturado_mayo"] - e["cobrado_de_mayo"])
    R["e4"]["diferencia"] = R["e4"]["util_dev"] - R["e4"]["neto_caja"]
    # la diferencia se explica partida por partida
    # (el arriendo pagado en mayo, de junio, y el devengado, de mayo, son iguales y se cancelan)
    assert R["e4"]["diferencia"] == R["e4"]["por_cobrar"] - e["cobro_de_abril"] - e["internet_mayo"]

    # --- 1.5
    e = E5
    ac = sum(e["corriente"].values()); anc = sum(e["no_corriente"].values())
    pc = sum(e["pasivo_c"].values()); pnc = sum(e["pasivo_nc"].values())
    pat = ac + anc - pc - pnc
    R["e5"] = dict(ac=ac, anc=anc, activo=ac + anc, pc=pc, pnc=pnc, pasivo=pc + pnc, pat=pat,
                   capital=pat - sum(e["patrimonio_conocido"].values()), capital_trabajo=ac - pc,
                   vehiculo_neto=e["no_corriente"]["Vehículo"] + e["no_corriente"]["Depreciación acumulada del vehículo"])

    # --- 1.6
    e = E6
    mb = e["ventas"] - e["costo_ventas"]
    gop = sum(e["operacionales"].values())
    ro = mb - gop
    rno = e["utilidad_venta_vehiculo"] - e["intereses"]
    R["e6"] = dict(mb=mb, gop=gop, ro=ro, rno=rno, rai=ro + rno, mb_pct=mb / e["ventas"], ro_pct=ro / e["ventas"],
                   rai_pct=(ro + rno) / e["ventas"])

    # --- 1.7
    e = E7
    p0, p1 = e["A0"] - e["P0"], e["A1"] - e["P1"]
    R["e7"] = dict(pat0=p0, pat1=p1, d_pat=p1 - p0, utilidad=p1 - p0 - e["aportes"] + e["dividendos"],
                   d_activo=e["A1"] - e["A0"])

    # --- 1.8
    e = E8
    f = {t: sum(v for _, v, c in e["mov"] if c == t) for t in "OIF"}
    R["e8"] = dict(O=f["O"], I=f["I"], F=f["F"], var=sum(f.values()), final=e["saldo_inicial"] + sum(f.values()))

    # --- 1.9
    e = E9
    util = e["ventas"] - e["costo_ventas"] - e["gastos_pagados"]
    caja_op = e["ventas"] * e["contado"] - e["compras_contado"] - e["gastos_pagados"]
    d_clientes = e["ventas"] * (1 - e["contado"])
    d_inv = e["compras_contado"] - e["costo_ventas"]
    assert abs(util - d_clientes - d_inv - caja_op) < 1e-6
    contado_min = (e["costo_ventas"] + e["gastos_pagados"]) / e["ventas"]
    R["e9"] = dict(util=util, caja_op=caja_op, d_clientes=d_clientes, d_inv=d_inv, contado_min=contado_min,
                   cobrado=e["ventas"] * e["contado"])

    # --- 1.10
    e = E10
    plusvalia = e["terreno_tasacion"] - e["terreno_costo"]
    seguro_anticipado = e["seguro_anual"] * (12 - e["meses_consumidos"]) / 12
    seguro_gasto = e["seguro_anual"] - seguro_anticipado
    util = e["utilidad"] - plusvalia + seguro_anticipado - e["anticipo_cliente"] + e["retiro_dueno"]
    activo = e["activo"] - plusvalia + seguro_anticipado - e["dep_acumulada"]
    pasivo = e["pasivo"] + e["anticipo_cliente"] - e["dep_acumulada"]
    pat0 = e["activo"] - e["pasivo"]
    pat1 = activo - pasivo
    assert abs((pat1 - pat0) - ((util - e["utilidad"]) - e["retiro_dueno"])) < 1e-6
    R["e10"] = dict(plusvalia=plusvalia, seguro_anticipado=seguro_anticipado, seguro_gasto=seguro_gasto, util=util,
                    activo=activo, pasivo=pasivo, pat0=pat0, pat1=pat1, d_pat=pat1 - pat0,
                    d_util=util - e["utilidad"])
    return R


# ---------------------------------------------------------------- caso integrador

TIPO_CASO = {"Banco": "A", "Clientes": "A", "Poleras lisas básicas": "A", "Otras prendas lisas": "A", "Tinta": "A",
             "IVA crédito fiscal": "A", "Arriendo pagado por anticipado": "A", "Seguro pagado por anticipado": "A",
             "Maquinaria": "A", "Depreciación acumulada": "CA", "Proveedores": "P", "IVA débito fiscal": "P",
             "Documentos por pagar a largo plazo": "P", "Préstamo bancario": "P", "Capital": "PAT",
             "Ventas": "I", "Sueldos": "G", "Gastos de administración": "G"}
CORRIENTE = ["Banco", "Clientes", "Poleras lisas básicas", "Tinta", "Arriendo pagado por anticipado"]


def cargar_caso():
    spec = importlib.util.spec_from_file_location("caso", AQUI / "caso_estampa_sur_contable.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def caso():
    m = cargar_caso()
    C = m.calcular()
    a, p = m.BALANCE_INICIAL["activo"], m.BALANCE_INICIAL["pasivo"]
    ac = sum(a[k] for k in CORRIENTE)
    anc = sum(v for k, v in a.items() if k not in CORRIENTE)
    total = ac + anc
    pasivo = sum(p.values())
    # tres operaciones de marzo en la ecuación (asientos 5, 4 y 7 del script)
    ops = {}
    for f in C["u2"]["asientos"]:
        if f["n"] in (4, 5, 7):
            mov = {}
            for k, v in f["debe"].items():
                mov[k] = mov.get(k, 0) + (v if TIPO_CASO[k] in ("A", "G", "RET") else -v)
            for k, v in f["haber"].items():
                mov[k] = mov.get(k, 0) + (-v if TIPO_CASO[k] in ("A", "G", "RET") else v)
            ops[f["n"]] = efecto(mov, TIPO_CASO)
    propio = dict(ac=ac, anc=anc, total_activo=total, total_pasivo=pasivo, capital=total - pasivo,
                  capital_trabajo=ac - pasivo, maquinaria_neta=a["Maquinaria"] + a["Depreciación acumulada"],
                  deuda_pct=pasivo / total, ops=ops, poleras=a["Poleras lisas básicas"], tinta=a["Tinta"],
                  activo_bruto=total - a["Depreciación acumulada"])
    return C["u1"], propio, a, p


# ---------------------------------------------------------------- planilla (con fórmulas)

def generar_planilla(R, a, p):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance 01-03"
    ws["A1"] = "Estampa Sur · balance al 1 de marzo (pesos)"; ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["Activo corriente", ""])
    fila = 4
    ini = fila
    for k in CORRIENTE:
        ws.cell(fila, 1, k); ws.cell(fila, 2, a[k]); fila += 1
    ws.cell(fila, 1, "Total activo corriente"); ws.cell(fila, 2, f"=SUM(B{ini}:B{fila-1})"); f_ac = fila; fila += 2
    ws.cell(fila, 1, "Activo no corriente"); fila += 1
    ini = fila
    for k, v in a.items():
        if k in CORRIENTE:
            continue
        ws.cell(fila, 1, k); ws.cell(fila, 2, v); fila += 1
    ws.cell(fila, 1, "Total activo no corriente"); ws.cell(fila, 2, f"=SUM(B{ini}:B{fila-1})"); f_anc = fila; fila += 2
    ws.cell(fila, 1, "TOTAL ACTIVO"); ws.cell(fila, 2, f"=B{f_ac}+B{f_anc}"); f_ta = fila; fila += 2
    ws.cell(fila, 1, "Pasivo corriente"); fila += 1
    ini = fila
    for k, v in p.items():
        ws.cell(fila, 1, k); ws.cell(fila, 2, v); fila += 1
    ws.cell(fila, 1, "TOTAL PASIVO"); ws.cell(fila, 2, f"=SUM(B{ini}:B{fila-1})"); f_tp = fila; fila += 2
    ws.cell(fila, 1, "Capital (patrimonio = activo − pasivo)"); ws.cell(fila, 2, f"=B{f_ta}-B{f_tp}"); f_cap = fila; fila += 1
    ws.cell(fila, 1, "TOTAL PASIVO + PATRIMONIO"); ws.cell(fila, 2, f"=B{f_tp}+B{f_cap}"); f_tpp = fila; fila += 1
    ws.cell(fila, 1, "Diferencia (debe dar 0)"); ws.cell(fila, 2, f"=B{f_ta}-B{f_tpp}"); f_dif = fila; fila += 1
    ws.cell(fila, 1, "Capital de trabajo"); ws.cell(fila, 2, f"=B{f_ac}-B{f_tp}"); f_ct = fila
    celdas_caso = dict(ac=f"B{f_ac}", anc=f"B{f_anc}", total_activo=f"B{f_ta}", total_pasivo=f"B{f_tp}",
                       capital=f"B{f_cap}", diferencia=f"B{f_dif}", capital_trabajo=f"B{f_ct}")
    ws.column_dimensions["A"].width = 42; ws.column_dimensions["B"].width = 16

    # hoja 2 · ejemplo: análisis de transacciones y estados
    w2 = wb.create_sheet("Rueda Libre abril")
    cuentas = list(TIPO_EJ)
    w2.append(["Transacción", "Fecha"] + cuentas)
    for letra, fecha, glosa, mov, _ in EJ:
        w2.append([letra, fecha] + [mov.get(k, None) for k in cuentas])
    n = len(EJ) + 1
    fila_saldo = n + 1
    w2.cell(fila_saldo, 1, "Saldo al 30-04")
    from openpyxl.utils import get_column_letter as L
    col = {k: L(3 + i) for i, k in enumerate(cuentas)}
    for k in cuentas:
        w2[f"{col[k]}{fila_saldo}"] = f"=SUM({col[k]}2:{col[k]}{n})"
    s = lambda k: f"{col[k]}{fila_saldo}"
    r = fila_saldo + 2
    filas = [
        ("Estado de resultados de abril", None),
        ("Ingresos", f"={s('Ventas de repuestos')}+{s('Ingresos por reparaciones')}"),
        ("Costo de ventas", f"={s('Costo de ventas')}"),
        ("Margen bruto", "=B{r1}-B{r2}"),
        ("Gastos de operación", f"={s('Sueldos')}+{s('Arriendo')}+{s('Electricidad')}"),
        ("Utilidad del mes", "=B{r3}-B{r4}"),
        ("Balance al 30-04", None),
        ("Total activo", "=" + "+".join(s(k) for k in cuentas if TIPO_EJ[k] == "A")),
        ("Total pasivo", "=" + "+".join(s(k) for k in cuentas if TIPO_EJ[k] == "P")),
        ("Patrimonio final (capital + utilidad − retiros)", f"={s('Capital')}+B{{r5}}-{s('Retiros del dueño')}"),
        ("Diferencia (debe dar 0)", "=B{r6}-B{r7}-B{r8}"),
    ]
    pos = {}
    for i, (rot, _) in enumerate(filas):
        pos[i] = r + i
    ref = dict(r1=pos[1], r2=pos[2], r3=pos[3], r4=pos[4], r5=pos[5], r6=pos[7], r7=pos[8], r8=pos[9])
    for i, (rot, form) in enumerate(filas):
        w2.cell(pos[i], 1, rot)
        if form:
            w2.cell(pos[i], 2, form.format(**ref))
    celdas_ej = dict(ingresos=f"B{pos[1]}", margen_bruto=f"B{pos[3]}", gastos_op=f"B{pos[4]}", utilidad=f"B{pos[5]}",
                     activo=f"B{pos[7]}", pasivo=f"B{pos[8]}", pat_final=f"B{pos[9]}", diferencia=f"B{pos[10]}",
                     banco=s("Banco"))
    w2.column_dimensions["A"].width = 44
    destino = RAMO / "planillas" / "u01_estados_financieros.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)
    return destino, celdas_caso, celdas_ej


class Evaluador:
    """Evalúa las fórmulas de la planilla (SUM, SUMIF, MAX, MIN, ABS, ROUND y aritmética)."""
    REF = re.compile(r"\"[^\"]*\"|(?:(?:'([^']+)'|([A-Za-z_][\w]*))!)?(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?")

    def __init__(self, ruta):
        from openpyxl import load_workbook
        self.wb = load_workbook(ruta)
        self.cache = {}

    def valor(self, hoja, ref):
        ref = ref.replace("$", "")
        if (hoja, ref) in self.cache:
            return self.cache[(hoja, ref)]
        v = self.wb[hoja][ref].value
        if isinstance(v, str) and v.startswith("="):
            v = self._eval(hoja, v[1:])
        self.cache[(hoja, ref)] = v
        return v

    def rango(self, hoja, a, b):
        from openpyxl.utils import range_boundaries, get_column_letter
        c1, r1, c2, r2 = range_boundaries(f"{a}:{b}".replace("$", ""))
        return [self.valor(hoja, f"{get_column_letter(c)}{r}") for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]

    def _eval(self, hoja, f):
        def sub(m):
            if m.group(0).startswith('"'):
                return m.group(0)
            h = m.group(1) or m.group(2) or hoja
            if m.group(4):
                return f"_R({h!r},{m.group(3)!r},{m.group(4)!r})"
            return f"_V({h!r},{m.group(3)!r})"
        expr = self.REF.sub(sub, f)
        expr = re.sub(r"\b(SUMIF|SUM|MAX|MIN|ABS|ROUND)\(", lambda m: "_" + m.group(1).lower() + "(", expr)
        num = lambda x: x if isinstance(x, (int, float)) else 0
        plano = lambda args: [y for x in args for y in (x if isinstance(x, list) else [x])]
        texto_o_num = lambda v: v if isinstance(v, str) else num(v)
        env = dict(_R=self.rango, _V=lambda h, r: texto_o_num(self.valor(h, r)),
                   _sum=lambda *a: sum(num(x) for x in plano(a)),
                   _sumif=lambda rg, c, sr: sum(num(s) for x, s in zip(rg, sr) if x == c),
                   _max=lambda *a: max(num(x) for x in plano(a)), _min=lambda *a: min(num(x) for x in plano(a)),
                   _abs=abs, _round=round)
        return eval(expr, {"__builtins__": {}}, env)


# ---------------------------------------------------------------- cifras esperadas en el texto

def esperados(R, u1, propio):
    p = fmt
    e = []
    j = R["ej"]
    for f in j["filas"]:
        e += [p(f["A"]), p(f["P"]), p(f["Pat"])]
    s = j["saldo"]
    e += [p(s["Banco"]), p(s["Clientes"]), p(s["Mercaderías"]), p(s["Arriendo pagado por anticipado"]),
          p(s["Proveedores"]), p(j["ingresos"]), p(j["margen_bruto"]), p(j["gastos_op"]), p(j["utilidad"]),
          p(j["flujo"]["O"]), p(j["flujo"]["I"]), p(j["flujo"]["F"]), p(j["cobros"]), p(j["pagos"]),
          p(j["activo"]), p(j["pasivo"]), p(j["pat_final"]), p(j["activo_corriente"])]
    j = R["e2"]; e += [p(j["pat_A"]), pct(j["endeuda_A"]), p(j["act_B"]), p(j["pat_C0"]), p(j["pat_C1"]), p(j["d_pat_C"]),
                       p(j["A_C1"]), p(j["P_C1"])]
    for f in R["e3"]["filas"]:
        e += [p(f["A"]), p(f["P"]), p(f["Pat"])]
    j = R["e4"]; e += [p(j["gastos_dev"]), p(j["util_dev"]), p(j["entradas"]), p(j["salidas"]), p(j["neto_caja"]),
                       p(j["diferencia"]), p(j["por_cobrar"])]
    j = R["e5"]; e += [p(j["ac"]), p(j["anc"]), p(j["activo"]), p(j["pc"]), p(j["pnc"]), p(j["pasivo"]), p(j["pat"]),
                       p(j["capital"]), p(j["capital_trabajo"]), p(j["vehiculo_neto"])]
    j = R["e6"]; e += [p(j["mb"]), p(j["gop"]), p(j["ro"]), p(j["rno"]), p(j["rai"]), pct(j["mb_pct"]), pct(j["ro_pct"]),
                       pct(j["rai_pct"])]
    j = R["e7"]; e += [p(j["pat0"]), p(j["pat1"]), p(j["d_pat"]), p(j["utilidad"]), p(j["d_activo"])]
    j = R["e8"]; e += [p(j["O"]), p(j["I"]), p(j["F"]), p(j["var"]), p(j["final"])]
    j = R["e9"]; e += [p(j["util"]), p(j["caja_op"]), p(j["d_clientes"]), p(j["d_inv"]), pct(j["contado_min"]),
                       p(j["cobrado"])]
    j = R["e10"]; e += [p(j["plusvalia"]), p(j["seguro_anticipado"]), p(j["seguro_gasto"]), p(j["util"]), p(j["activo"]),
                        p(j["pasivo"]), p(j["pat0"]), p(j["pat1"]), p(j["d_pat"]), p(j["d_util"])]
    e += [p(u1["total_activo"]), p(u1["total_pasivo"]), p(u1["capital"]), p(propio["ac"]), p(propio["anc"]),
          p(propio["capital_trabajo"]), p(propio["maquinaria_neta"]), pct(propio["deuda_pct"]), p(propio["poleras"]),
          p(propio["tinta"]), p(propio["activo_bruto"])]
    for n, (dA, dP, dPat) in propio["ops"].items():
        e += [p(round(dA)), p(round(dP)), p(round(dPat))]
    return e


if __name__ == "__main__":
    R = resultados()
    u1, propio, a, p = caso()
    for k, d in R.items():
        print(f"== {k}")
        for kk, v in d.items():
            if kk == "filas":
                for f in v:
                    print("     ", {x: (fmt(y) if isinstance(y, (int, float)) else y) for x, y in f.items()})
            else:
                print(f"   {kk}: {v}")
    print("== caso (script)", u1)
    print("== caso (propio)", {k: v for k, v in propio.items()})

    ruta, cc, ce = generar_planilla(R, a, p)
    ev = Evaluador(ruta)
    hc, he = "Balance 01-03", "Rueda Libre abril"
    j = R["ej"]
    chequeos = {
        "caso: total activo = script": propio["total_activo"] == u1["total_activo"],
        "caso: total pasivo = script": propio["total_pasivo"] == u1["total_pasivo"],
        "caso: capital = script": propio["capital"] == u1["capital"],
        "caso: las tres operaciones mantienen la ecuación": all(abs(x - y - z) < 1e-6 for x, y, z in propio["ops"].values()),
        "planilla: balance del caso": all(abs(ev.valor(hc, cc[k]) - propio[k]) < 1e-6
                                         for k in ("ac", "anc", "total_activo", "total_pasivo", "capital", "capital_trabajo")),
        "planilla: la diferencia del balance es 0": abs(ev.valor(hc, cc["diferencia"])) < 1e-6,
        "planilla: estados del ejemplo": all(abs(ev.valor(he, ce[k]) - j[k]) < 1e-6
                                            for k in ("ingresos", "margen_bruto", "gastos_op", "utilidad", "activo", "pasivo", "pat_final")),
        "planilla: saldo del banco del ejemplo": ev.valor(he, ce["banco"]) == j["saldo"]["Banco"],
        "planilla: la diferencia del ejemplo es 0": abs(ev.valor(he, ce["diferencia"])) < 1e-6,
    }
    print(f"\nPLANILLA: {ruta}")
    print("CHEQUEOS:")
    for k, v in chequeos.items():
        print(f"   {'OK ' if v else 'MAL'} {k}")
    if not all(chequeos.values()):
        sys.exit(1)
    print("CADENA: OK, la pieza del caso calza con caso_estampa_sur_contable.py")
    if "--revisar" in sys.argv:
        md = (RAMO / "unidades" / "01-contabilidad-y-estados-financieros.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %").replace("−", "-")
        faltan = [x for x in esperados(R, u1, propio) if x not in md]
        print("\nREVISIÓN:", "OK, todas las cifras están en el texto" if not faltan else f"faltan {len(faltan)}: {faltan}")
        sys.exit(1 if faltan else 0)
