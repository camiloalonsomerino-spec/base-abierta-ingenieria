"""Unidad 7 · Inventarios: FIFO y precio promedio ponderado.

Recalcula todas las cifras de la unidad desde los datos, genera la planilla de las tarjetas de
existencias (con fórmulas) y comprueba que la planilla dé lo mismo que Python. Con --revisar busca
cada resultado, y cada fila de las tarjetas, escrito tal cual en el .md.

    py -3 -P u07_inventarios.py            # imprime los resultados y genera la planilla
    py -3 -P u07_inventarios.py --tablas   # además imprime las tarjetas en Markdown
    py -3 -P u07_inventarios.py --revisar  # además busca las cifras en el .md

Todos los datos son originales de esta base. No vienen de ningún libro ni curso.
"""
from pathlib import Path
import importlib.util, re, sys

sys.stdout.reconfigure(encoding="utf-8")
AQUI = Path(__file__).resolve().parent
IVA = 0.19


# ---------------------------------------------------------------- herramientas

def fmt(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, dec=1):
    return fmt(100 * x, dec) + " %"


def pesos(x):
    return "$" + fmt(x)


def tarjeta_fifo(eventos):
    """Tarjeta FIFO. Cada evento: (fecha, tipo, unidades, precio o None, referencia de venta o None)."""
    lotes, filas, ventas = [], [], {}
    for fecha, tipo, q, p, ref in eventos:
        ent, sal = [], []
        if tipo in ("inicial", "compra"):
            lotes.append([q, p]); ent = [(q, p)]
        elif tipo == "dev_compra":            # sale al precio de la compra que se devuelve
            lote = [l for l in lotes if l[1] == p][-1]
            lote[0] -= q; sal = [(q, p)]
        elif tipo in ("venta", "merma"):      # salen primero los lotes más antiguos
            resta = q
            while resta:
                t = min(resta, lotes[0][0])
                sal.append((t, lotes[0][1])); lotes[0][0] -= t; resta -= t
                if lotes[0][0] == 0:
                    lotes.pop(0)
            ventas[fecha] = list(sal)
        elif tipo == "dev_venta":             # vuelve al costo con que salió (las últimas unidades de esa venta)
            resta = q
            for t, pv in reversed(ventas[ref]):
                u = min(resta, t)
                destino = [l for l in lotes if l[1] == pv]
                if destino:
                    destino[0][0] += u
                else:
                    lotes.insert(0, [u, pv])
                ent.append((u, pv)); resta -= u
                if not resta:
                    break
        lotes = [l for l in lotes if l[0]]
        filas.append(dict(fecha=fecha, tipo=tipo, ent=ent, sal=sal,
                          su=sum(l[0] for l in lotes), sv=sum(l[0] * l[1] for l in lotes),
                          detalle=[tuple(l) for l in lotes]))
    return filas


def tarjeta_ppp(eventos):
    """Tarjeta con precio promedio ponderado móvil: el promedio se recalcula con cada entrada."""
    su, sv, filas, ventas = 0, 0.0, [], {}
    for fecha, tipo, q, p, ref in eventos:
        ent, sal = [], []
        prom = sv / su if su else 0
        if tipo in ("inicial", "compra"):
            su += q; sv += q * p; ent = [(q, p)]
        elif tipo == "dev_compra":
            su -= q; sv -= q * p; sal = [(q, p)]
        elif tipo in ("venta", "merma"):
            su -= q; sv -= q * prom; sal = [(q, prom)]; ventas[fecha] = prom
        elif tipo == "dev_venta":
            su += q; sv += q * ventas[ref]; ent = [(q, ventas[ref])]
        filas.append(dict(fecha=fecha, tipo=tipo, ent=ent, sal=sal, su=su, sv=sv, detalle=sv / su))
    return filas


NOMBRE = {"inicial": "Saldo inicial", "compra": "Compra", "dev_compra": "Devolución de compra",
          "venta": "Venta", "merma": "Merma", "dev_venta": "Devolución de venta", "consumo": "Consumo"}


def total(partes):
    return sum(q * p for q, p in partes)


def celda(partes):
    if not partes:
        return ""
    s = " + ".join(f"{fmt(q)} × {fmt(p)}" for q, p in partes)
    return s + (f" = {fmt(total(partes))}" if len(partes) > 1 or partes[0][0] != 1 else "")


def tabla_md(filas, metodo, nombres=None):
    nombres = nombres or {}
    cab = "| Fecha | Movimiento | Entra | Sale | Saldo (u) | Saldo (pesos) | " + \
          ("Lotes en bodega |" if metodo == "fifo" else "Promedio |")
    out = [cab, "|---|---|---|---|---|---|---|"]
    for f in filas:
        det = " + ".join(f"{fmt(q)} × {fmt(p)}" for q, p in f["detalle"]) if metodo == "fifo" else fmt(f["detalle"])
        out.append(f"| {f['fecha']} | {nombres.get(f['tipo'], NOMBRE[f['tipo']])} | {celda(f['ent'])} | "
                   f"{celda(f['sal'])} | {fmt(f['su'])} | {fmt(f['sv'])} | {det} |")
    return out


def resumen(filas):
    cv = sum(total(f["sal"]) for f in filas if f["tipo"] == "venta") - \
         sum(total(f["ent"]) for f in filas if f["tipo"] == "dev_venta")
    merma = sum(total(f["sal"]) for f in filas if f["tipo"] == "merma")
    return dict(costo_ventas=cv, merma=merma, saldo_u=filas[-1]["su"], saldo=filas[-1]["sv"])


def asiento(n, fecha, glosa, debe, haber):
    assert abs(sum(debe.values()) - sum(haber.values())) < 1e-6, f"asiento {n} no cuadra"
    return dict(n=n, fecha=fecha, glosa=glosa, debe=debe, haber=haber)


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · la bicicletería y sus cascos (mayo de 2026)
EJ = [("01-05-2026", "inicial", 40, 18_500, None),
      ("06-05-2026", "compra", 60, 19_700, None),
      ("12-05-2026", "venta", 70, None, None),
      ("20-05-2026", "compra", 50, 20_300, None),
      ("27-05-2026", "venta", 55, None, None),
      ("31-05-2026", "merma", 1, None, None)]
EJ_PRECIO = {"12-05-2026": 32_900, "27-05-2026": 33_500}
EJ_FISICO = 24
# 7.1 · sistema periódico
E1 = {"ii": 3_250_000, "compras": 8_940_000, "if": 2_715_000, "ventas": 14_600_000}
# 7.2 a 7.4 · un solo movimiento de compra y una venta
E2 = [("02-06-2026", "inicial", 120, 4_300, None), ("09-06-2026", "compra", 200, 4_580, None),
      ("16-06-2026", "venta", 250, None, None)]
E4 = {"precio_venta": 7_200}
# 7.5 · PPP con devoluciones y merma (pintura en galones)
E5 = [("01-07-2026", "inicial", 314, 8_230, None), ("04-07-2026", "compra", 506, 8_640, None),
      ("07-07-2026", "dev_compra", 35, 8_640, None), ("15-07-2026", "venta", 437, None, None),
      ("18-07-2026", "dev_venta", 23, None, "15-07-2026"), ("22-07-2026", "compra", 406, 8_920, None),
      ("28-07-2026", "venta", 479, None, None), ("31-07-2026", "merma", 3, None, None)]
E5_FISICO = 295
# 7.6 · FIFO contra PPP con precios que suben (rollos de film)
E6 = [("01-08-2026", "inicial", 470, 1_240, None), ("05-08-2026", "compra", 705, 1_295, None),
      ("12-08-2026", "venta", 1_013, None, None), ("19-08-2026", "compra", 801, 1_380, None),
      ("26-08-2026", "venta", 836, None, None)]
E6_PRECIO = {"12-08-2026": 2_090, "26-08-2026": 2_140}
# 7.7 · costo estándar y variación de precio (harina en sacos)
E7 = {"estandar": 14_200, "ii": 60, "compra": 140, "precio": 14_650, "consumo": 150}
# 7.8 · merma en el sistema permanente y error de conteo en el periódico
E8 = {"tarjeta": 1_240, "fisico": 1_212, "ppp": 3_150, "ii": 4_100_000, "compras": 11_350_000, "error_u": 40}
# 7.9 · precios que bajan
E9 = [("01-09-2026", "inicial", 380, 6_900, None), ("08-09-2026", "compra", 608, 6_510, None),
      ("14-09-2026", "venta", 713, None, None), ("21-09-2026", "compra", 475, 6_120, None),
      ("28-09-2026", "venta", 597, None, None)]
# 7.10 · tarjeta FIFO completa con devoluciones, asientos e IVA (mochilas)
E10 = [("01-10-2026", "inicial", 83, 16_100, None), ("05-10-2026", "compra", 127, 16_900, None),
       ("09-10-2026", "venta", 154, None, None), ("13-10-2026", "dev_venta", 9, None, "09-10-2026"),
       ("20-10-2026", "compra", 96, 17_450, None), ("24-10-2026", "dev_compra", 6, 17_450, None),
       ("27-10-2026", "venta", 117, None, None), ("31-10-2026", "merma", 1, None, None)]
E10_BRUTO = 35_581
E10_FISICO = 37


def caso_eventos(m):
    ini = m.BALANCE_INICIAL["activo"]["Poleras lisas básicas"] // 2_400
    c = m.COMPRA_POLERAS
    ev = [("1 de marzo", "inicial", ini, 2_400, None), ("semana 1", "compra", c["unidades"], c["precio"], None)]
    ev += [(sem, "venta", q, None, None) for q, sem in m.CONSUMO_POLERAS]
    return ev


def cargar_caso():
    spec = importlib.util.spec_from_file_location("caso", AQUI / "caso_estampa_sur_contable.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m, m.calcular()


# ---------------------------------------------------------------- resultados

def resultados(m):
    R = {}
    # ejemplo
    ff, fp = tarjeta_fifo(EJ), tarjeta_ppp(EJ)
    ventas = sum(q * EJ_PRECIO[f] for f, t, q, _, _ in EJ if t == "venta")
    rf, rp = resumen(ff), resumen(fp)
    disponible = sum(q * p for _, t, q, p, _ in EJ if t in ("inicial", "compra"))
    u_disp = sum(q for _, t, q, _, _ in EJ if t in ("inicial", "compra"))
    u_vend = sum(q for _, t, q, _, _ in EJ if t == "venta")
    compra = EJ[1][2] * EJ[1][3]
    v1 = EJ[2][2] * EJ_PRECIO["12-05-2026"]
    asientos = [
        asiento(1, "06-05-2026", "Compra de 60 cascos a crédito",
                {"Existencias": compra, "IVA crédito fiscal": compra * IVA}, {"Proveedores": compra * (1 + IVA)}),
        asiento(2, "12-05-2026", "Venta de 70 cascos a crédito",
                {"Clientes": v1 * (1 + IVA)}, {"Ventas": v1, "IVA débito fiscal": v1 * IVA}),
        asiento(3, "12-05-2026", "Costo de la venta del 12-05 (FIFO)",
                {"Costo de ventas": total(ff[2]["sal"])}, {"Existencias": total(ff[2]["sal"])}),
        asiento(4, "31-05-2026", "Merma detectada en el conteo (FIFO)",
                {"Pérdida por merma": total(ff[5]["sal"])}, {"Existencias": total(ff[5]["sal"])}),
    ]
    R["ej"] = dict(fifo=ff, ppp=fp, rf=rf, rp=rp, ventas=ventas, v1=v1, v2=ventas - v1,
                   mb_fifo=ventas - rf["costo_ventas"], mb_ppp=ventas - rp["costo_ventas"],
                   dif=rp["costo_ventas"] - rf["costo_ventas"], disponible=disponible,
                   ppp_periodico=disponible / u_disp, cv_periodico=u_vend * disponible / u_disp,
                   if_periodico=EJ_FISICO * disponible / u_disp,
                   merma_periodico=(u_disp - u_vend - EJ_FISICO) * disponible / u_disp,
                   cv_periodico_total=disponible - EJ_FISICO * disponible / u_disp,
                   asientos=asientos, compra=compra, saldo_teorico=u_disp - u_vend)
    assert ff[-1]["su"] == EJ_FISICO and fp[-1]["su"] == EJ_FISICO
    for r in (rf, rp):
        assert abs(r["costo_ventas"] + r["merma"] + r["saldo"] - disponible) < 1e-6

    # 7.1
    cv = E1["ii"] + E1["compras"] - E1["if"]
    R["e1"] = dict(cv=cv, mb=E1["ventas"] - cv, mb_pct=(E1["ventas"] - cv) / E1["ventas"],
                   disponible=E1["ii"] + E1["compras"])
    # 7.2 a 7.4
    f2, p2 = tarjeta_fifo(E2), tarjeta_ppp(E2)
    r2f, r2p = resumen(f2), resumen(p2)
    compra2 = E2[1][2] * E2[1][3]
    venta4 = E2[2][2] * E4["precio_venta"]
    a4 = [asiento(1, "09-06-2026", "Compra a crédito", {"Existencias": compra2, "IVA crédito fiscal": compra2 * IVA},
                  {"Proveedores": compra2 * (1 + IVA)}),
          asiento(2, "16-06-2026", "Venta al contado", {"Banco": venta4 * (1 + IVA)},
                  {"Ventas": venta4, "IVA débito fiscal": venta4 * IVA}),
          asiento(3, "16-06-2026", "Costo de la venta (PPP)", {"Costo de ventas": r2p["costo_ventas"]},
                  {"Existencias": r2p["costo_ventas"]})]
    R["e2"] = dict(fifo=f2, rf=r2f, prom=p2[1]["detalle"], rp=r2p, disponible=E2[0][2] * E2[0][3] + compra2,
                   dif=r2p["costo_ventas"] - r2f["costo_ventas"])
    R["e4"] = dict(asientos=a4, compra=compra2, venta=venta4, mb=venta4 - r2p["costo_ventas"],
                   mb_pct=(venta4 - r2p["costo_ventas"]) / venta4)
    # 7.5
    p5 = tarjeta_ppp(E5)
    R["e5"] = dict(ppp=p5, r=resumen(p5),
                   disponible=sum(q * p for _, t, q, p, _ in E5 if t in ("inicial", "compra")) -
                   sum(q * p for _, t, q, p, _ in E5 if t == "dev_compra"))
    assert p5[-1]["su"] == E5_FISICO
    assert abs(R["e5"]["r"]["costo_ventas"] + R["e5"]["r"]["merma"] + R["e5"]["r"]["saldo"] - R["e5"]["disponible"]) < 1e-6
    # 7.6
    f6, p6 = tarjeta_fifo(E6), tarjeta_ppp(E6)
    v6 = sum(q * E6_PRECIO[f] for f, t, q, _, _ in E6 if t == "venta")
    disp6 = sum(q * p for _, t, q, p, _ in E6 if t in ("inicial", "compra"))
    u6 = sum(q for _, t, q, _, _ in E6 if t in ("inicial", "compra"))
    uv6 = sum(q for _, t, q, _, _ in E6 if t == "venta")
    r6f, r6p = resumen(f6), resumen(p6)
    R["e6"] = dict(fifo=f6, ppp=p6, rf=r6f, rp=r6p, ventas=v6, disponible=disp6,
                   mb_fifo=v6 - r6f["costo_ventas"], mb_ppp=v6 - r6p["costo_ventas"],
                   dif=r6p["costo_ventas"] - r6f["costo_ventas"],
                   prom_periodico=disp6 / u6, cv_periodico=uv6 * disp6 / u6, if_periodico=(u6 - uv6) * disp6 / u6,
                   mb_periodico=v6 - uv6 * disp6 / u6)
    # 7.7
    e = E7
    ev7 = [("", "inicial", e["ii"], e["estandar"], None), ("", "compra", e["compra"], e["precio"], None),
           ("", "venta", e["consumo"], None, None)]
    c7f, c7p = resumen(tarjeta_fifo(ev7)), resumen(tarjeta_ppp(ev7))
    std = e["consumo"] * e["estandar"]
    saldo_u = e["ii"] + e["compra"] - e["consumo"]
    R["e7"] = dict(cf=c7f["costo_ventas"], cp=c7p["costo_ventas"], std=std, var_f=c7f["costo_ventas"] - std,
                   var_p=c7p["costo_ventas"] - std, prom=(e["ii"] * e["estandar"] + e["compra"] * e["precio"]) /
                   (e["ii"] + e["compra"]), var_compra=e["compra"] * (e["precio"] - e["estandar"]),
                   bod_f=c7f["saldo"] - saldo_u * e["estandar"], bod_p=c7p["saldo"] - saldo_u * e["estandar"],
                   por_saco=(e["ii"] * e["estandar"] + e["compra"] * e["precio"]) / (e["ii"] + e["compra"]) - e["estandar"],
                   mes_siguiente=e["consumo"] * (e["precio"] - e["estandar"]), dif_unit=e["precio"] - e["estandar"])
    assert abs(R["e7"]["var_f"] + R["e7"]["bod_f"] - R["e7"]["var_compra"]) < 1e-6
    assert abs(R["e7"]["var_p"] + R["e7"]["bod_p"] - R["e7"]["var_compra"]) < 1e-6
    # 7.8
    e = E8
    merma = (e["tarjeta"] - e["fisico"]) * e["ppp"]
    if_ok = e["fisico"] * e["ppp"]
    if_mal = (e["fisico"] + e["error_u"]) * e["ppp"]
    R["e8"] = dict(merma_u=e["tarjeta"] - e["fisico"], merma=merma, if_ok=if_ok, cv_ok=e["ii"] + e["compras"] - if_ok,
                   if_mal=if_mal, cv_mal=e["ii"] + e["compras"] - if_mal, inflada=if_mal - if_ok,
                   merma_pct=(e["tarjeta"] - e["fisico"]) / e["tarjeta"])
    # 7.9
    f9, p9 = tarjeta_fifo(E9), tarjeta_ppp(E9)
    r9f, r9p = resumen(f9), resumen(p9)
    R["e9"] = dict(fifo=f9, ppp=p9, rf=r9f, rp=r9p, dif=r9f["costo_ventas"] - r9p["costo_ventas"],
                   disponible=sum(q * p for _, t, q, p, _ in E9 if t in ("inicial", "compra")),
                   reposicion=r9f["saldo_u"] * E9[3][3])
    # 7.10
    f10 = tarjeta_fifo(E10)
    r10 = resumen(f10)
    neto = round(E10_BRUTO / (1 + IVA))
    uv = sum(q for _, t, q, _, _ in E10 if t == "venta") - sum(q for _, t, q, _, _ in E10 if t == "dev_venta")
    ventas10 = uv * neto
    c1 = E10[1][2] * E10[1][3]; c2 = E10[4][2] * E10[4][3]; dc = E10[5][2] * E10[5][3]
    ivd = ventas10 * IVA
    ivc = (c1 + c2 - dc) * IVA
    v_a, d_a, v_b = 154 * neto, 9 * neto, 117 * neto
    a10 = [
        asiento(1, "05-10-2026", "Compra de 127 mochilas por transferencia", {"Existencias": c1, "IVA crédito fiscal": c1 * IVA},
                {"Banco": c1 * (1 + IVA)}),
        asiento(2, "09-10-2026", "Venta de 154 mochilas", {"Banco": v_a * (1 + IVA)}, {"Ventas": v_a, "IVA débito fiscal": v_a * IVA}),
        asiento(3, "09-10-2026", "Costo de la venta del 09-10", {"Costo de ventas": total(f10[2]["sal"])},
                {"Existencias": total(f10[2]["sal"])}),
        asiento(4, "13-10-2026", "Devolución de 9 mochilas de la venta del 09-10", {"Ventas": d_a, "IVA débito fiscal": d_a * IVA},
                {"Banco": d_a * (1 + IVA)}),
        asiento(5, "13-10-2026", "Las 9 mochilas vuelven a bodega a su costo", {"Existencias": total(f10[3]["ent"])},
                {"Costo de ventas": total(f10[3]["ent"])}),
        asiento(6, "20-10-2026", "Compra de 96 mochilas por transferencia", {"Existencias": c2, "IVA crédito fiscal": c2 * IVA},
                {"Banco": c2 * (1 + IVA)}),
        asiento(7, "24-10-2026", "Devolución de 6 mochilas al proveedor", {"Banco": dc * (1 + IVA)},
                {"Existencias": dc, "IVA crédito fiscal": dc * IVA}),
        asiento(8, "27-10-2026", "Venta de 117 mochilas", {"Banco": v_b * (1 + IVA)}, {"Ventas": v_b, "IVA débito fiscal": v_b * IVA}),
        asiento(9, "27-10-2026", "Costo de la venta del 27-10", {"Costo de ventas": total(f10[6]["sal"])},
                {"Existencias": total(f10[6]["sal"])}),
        asiento(10, "31-10-2026", "Merma del conteo", {"Pérdida por merma": total(f10[7]["sal"])},
                {"Existencias": total(f10[7]["sal"])}),
        asiento(11, "12-11-2026", "Pago del IVA de octubre", {"IVA débito fiscal": ivd},
                {"IVA crédito fiscal": ivc, "Banco": ivd - ivc}),
    ]
    assert f10[-1]["su"] == E10_FISICO
    disp10 = E10[0][2] * E10[0][3] + c1 + c2 - dc
    assert abs(r10["costo_ventas"] + r10["merma"] + r10["saldo"] - disp10) < 1e-6
    R["e10"] = dict(fifo=f10, r=r10, neto=neto, unidades=uv, ventas=ventas10, mb=ventas10 - r10["costo_ventas"],
                    resultado=ventas10 - r10["costo_ventas"] - r10["merma"], ivd=ivd, ivc=ivc, pagar=ivd - ivc,
                    asientos=a10, disponible=disp10, mb_pct=(ventas10 - r10["costo_ventas"]) / ventas10)

    # caso
    ev = caso_eventos(m)
    cf, cp = tarjeta_fifo(ev), tarjeta_ppp(ev)
    rcf, rcp = resumen(cf), resumen(cp)
    consumo = sum(q for q, _ in m.CONSUMO_POLERAS)
    std = consumo * 2_400
    compra = m.COMPRA_POLERAS["unidades"] * m.COMPRA_POLERAS["precio"]
    prom = cp[1]["detalle"]
    var_semana_p = [q * (prom - 2_400) for q, _ in m.CONSUMO_POLERAS]
    var_semana_f = [total(cf[2]["sal"]) - m.CONSUMO_POLERAS[0][0] * 2_400,
                    total(cf[3]["sal"]) - m.CONSUMO_POLERAS[1][0] * 2_400]
    mc = m.MARGEN[0]
    R["caso"] = dict(fifo=cf, ppp=cp, consumo=consumo, costo_fifo=rcf["costo_ventas"], saldo_fifo=rcf["saldo"],
                     ppp_=prom, costo_ppp=rcp["costo_ventas"], saldo_ppp=rcp["saldo"], costo_estandar=std,
                     variacion_fifo=rcf["costo_ventas"] - std, variacion_ppp=rcp["costo_ventas"] - std,
                     compra=compra, iva_compra=compra * IVA, total_compra=compra * (1 + IVA),
                     var_compra=m.COMPRA_POLERAS["unidades"] * (m.COMPRA_POLERAS["precio"] - 2_400),
                     bod_fifo=rcf["saldo"] - rcf["saldo_u"] * 2_400, bod_ppp=rcp["saldo"] - rcp["saldo_u"] * 2_400,
                     var_semana_p=var_semana_p, var_semana_f=var_semana_f,
                     std_semana=[q * 2_400 for q, _ in m.CONSUMO_POLERAS],
                     real_semana_p=[q * prom for q, _ in m.CONSUMO_POLERAS],
                     real_semana_f=[total(cf[2]["sal"]), total(cf[3]["sal"])],
                     unit_fifo=rcf["costo_ventas"] / consumo, sobre_std=prom - 2_400,
                     mc=mc, mc_real=mc - (prom - 2_400), mc_nuevo=mc - (m.COMPRA_POLERAS["precio"] - 2_400),
                     extra_mes=consumo * (m.COMPRA_POLERAS["precio"] - 2_400),
                     saldo_inicial=m.BALANCE_INICIAL["activo"]["Poleras lisas básicas"], saldo_u=rcp["saldo_u"],
                     caida_mc=(prom - 2_400) / mc, caida_mc_nuevo=(m.COMPRA_POLERAS["precio"] - 2_400) / mc)
    c = R["caso"]
    assert abs(c["variacion_fifo"] + c["bod_fifo"] - c["var_compra"]) < 1e-6
    assert abs(c["variacion_ppp"] + c["bod_ppp"] - c["var_compra"]) < 1e-6
    return R


# ---------------------------------------------------------------- planilla con fórmulas

def evaluar(ws):
    """Evalúa las fórmulas simples de la hoja (+ − × ÷, referencias y SUM) para comprobarlas en Python."""
    cache = {}

    def valor(ref):
        if ref in cache:
            return cache[ref]
        v = ws[ref].value
        if isinstance(v, str) and v.startswith("="):
            expr = v[1:]
            expr = re.sub(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)",
                          lambda mm: "(" + "+".join(f"{mm.group(1)}{k}" for k in range(int(mm.group(2)), int(mm.group(4)) + 1)) + ")",
                          expr)
            expr = re.sub(r"\b([A-Z]{1,2}\d+)\b", lambda mm: repr(valor(mm.group(1))), expr)
            v = eval(expr)
        cache[ref] = v if v is not None else 0
        return cache[ref]
    return valor


def hoja_tarjeta(wb, titulo, eventos, metodo, nombres=None):
    """Una tarjeta por hoja. FIFO: una fila por lote que sale. PPP: el promedio es fórmula."""
    from openpyxl.styles import Font
    nombres = nombres or {}
    ws = wb.create_sheet(titulo)
    ws.append([titulo + (" · FIFO" if metodo == "fifo" else " · precio promedio ponderado")])
    ws["A1"].font = Font(bold=True)
    cab = ["Fecha", "Movimiento", "Lote", "Entra u", "Entra precio", "Entra $", "Sale u", "Sale precio", "Sale $",
           "Saldo u", "Saldo $"] + (["Promedio"] if metodo == "ppp" else [])
    ws.append(cab)
    for c in ws[2]:
        c.font = Font(bold=True)
    fila = 3
    lote_fila = {}          # precio -> fila donde entró ese lote (FIFO)
    venta_fila = {}         # fecha de venta -> filas de salida
    marcas = {"venta": [], "dev_venta": [], "merma": []}
    filas = tarjeta_fifo(eventos) if metodo == "fifo" else tarjeta_ppp(eventos)
    for (fecha, tipo, q, p, ref), f in zip(eventos, filas):
        partes = []
        if f["ent"]:
            partes += [("ent", u, pr) for u, pr in f["ent"]]
        if f["sal"]:
            partes += [("sal", u, pr) for u, pr in f["sal"]]
        for k, (lado, u, pr) in enumerate(partes):
            r = fila
            prev = r - 1
            ws.cell(r, 1, fecha); ws.cell(r, 2, nombres.get(tipo, NOMBRE[tipo]))
            if lado == "ent":
                ws.cell(r, 4, u)
                if tipo in ("inicial", "compra"):
                    ws.cell(r, 5, pr); lote_fila.setdefault(pr, r)
                elif metodo == "fifo":
                    ws.cell(r, 5, f"=E{lote_fila[pr]}")
                else:                                   # devolución de venta al promedio con que salió
                    ws.cell(r, 5, f"=H{venta_fila[ref][0]}")
                ws.cell(r, 6, f"=D{r}*E{r}")
                ws.cell(r, 3, fmt(pr) if metodo == "fifo" else "")
            else:
                ws.cell(r, 7, u)
                if metodo == "fifo" or tipo == "dev_compra":
                    ws.cell(r, 8, f"=E{lote_fila[pr]}")
                else:
                    ws.cell(r, 8, f"=L{prev}")
                ws.cell(r, 9, f"=G{r}*H{r}")
                ws.cell(r, 3, fmt(pr) if metodo == "fifo" else "")
                venta_fila.setdefault(fecha, []).append(r)
            if r == 3:
                ws.cell(r, 10, f"=D{r}-G{r}"); ws.cell(r, 11, f"=F{r}-I{r}")
            else:
                ws.cell(r, 10, f"=J{prev}+D{r}-G{r}"); ws.cell(r, 11, f"=K{prev}+F{r}-I{r}")
            if metodo == "ppp":
                ws.cell(r, 12, f"=K{r}/J{r}")
            if tipo in marcas:
                marcas[tipo].append((r, "F" if lado == "ent" else "I"))
            fila += 1
    fila += 1
    sumar = lambda lst: "+".join(f"{col}{r}" for r, col in lst) or "0"
    ws.cell(fila, 2, "Costo de ventas"); ws.cell(fila, 11, f"={sumar(marcas['venta'])}-({sumar(marcas['dev_venta'])})")
    ws.cell(fila + 1, 2, "Merma"); ws.cell(fila + 1, 11, f"={sumar(marcas['merma'])}")
    ws.cell(fila + 2, 2, "Saldo final"); ws.cell(fila + 2, 11, f"=K{fila - 2}")
    for col, ancho in zip("ABCDEFGHIJKL", [13, 22, 10, 9, 12, 13, 9, 12, 13, 9, 13, 11]):
        ws.column_dimensions[col].width = ancho
    for row in ws.iter_rows(min_row=3, max_row=fila + 2):
        for c in row[3:]:
            c.number_format = "#,##0"
    return ws, dict(costo=f"K{fila}", merma=f"K{fila + 1}", saldo=f"K{fila + 2}", ultima=fila + 2)


def planilla(R, m):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    nom_caso = {"venta": "Consumo", "inicial": "Saldo inicial", "compra": "Compra"}
    ev = caso_eventos(m)
    hojas = {}
    hojas["caso_fifo"] = hoja_tarjeta(wb, "Caso poleras FIFO", ev, "fifo", nom_caso)
    hojas["caso_ppp"] = hoja_tarjeta(wb, "Caso poleras PPP", ev, "ppp", nom_caso)
    for clave in ("caso_fifo", "caso_ppp"):
        ws, ref = hojas[clave]
        r = ref["ultima"] + 2
        ws.cell(r, 2, "Consumo (u)"); ws.cell(r, 11, R["caso"]["consumo"])
        ws.cell(r + 1, 2, "Costo estándar ($2.400)"); ws.cell(r + 1, 11, f"=K{r}*2400")
        ws.cell(r + 2, 2, "Variación de precio"); ws.cell(r + 2, 11, f"={ref['costo']}-K{r + 1}")
        ws.cell(r + 2, 2).font = Font(bold=True)
        for k in range(3):
            ws.cell(r + k, 11).number_format = "#,##0"
        ref["variacion"] = f"K{r + 2}"
    hojas["ej_fifo"] = hoja_tarjeta(wb, "Ejemplo cascos FIFO", EJ, "fifo")
    hojas["ej_ppp"] = hoja_tarjeta(wb, "Ejemplo cascos PPP", EJ, "ppp")
    hojas["e5"] = hoja_tarjeta(wb, "Ejercicio 7.5 PPP", E5, "ppp")
    hojas["e10"] = hoja_tarjeta(wb, "Ejercicio 7.10 FIFO", E10, "fifo")
    destino = AQUI.parents[0] / "planillas" / "u07_inventarios.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)
    # comprobar la planilla contra Python
    esperado = {"caso_fifo": (R["caso"]["costo_fifo"], 0, R["caso"]["saldo_fifo"]),
                "caso_ppp": (R["caso"]["costo_ppp"], 0, R["caso"]["saldo_ppp"]),
                "ej_fifo": (R["ej"]["rf"]["costo_ventas"], R["ej"]["rf"]["merma"], R["ej"]["rf"]["saldo"]),
                "ej_ppp": (R["ej"]["rp"]["costo_ventas"], R["ej"]["rp"]["merma"], R["ej"]["rp"]["saldo"]),
                "e5": (R["e5"]["r"]["costo_ventas"], R["e5"]["r"]["merma"], R["e5"]["r"]["saldo"]),
                "e10": (R["e10"]["r"]["costo_ventas"], R["e10"]["r"]["merma"], R["e10"]["r"]["saldo"])}
    ok = True
    for clave, (ws, ref) in hojas.items():
        v = evaluar(ws)
        got = (v(ref["costo"]), v(ref["merma"]), v(ref["saldo"]))
        ok &= all(abs(a - b) < 1e-6 for a, b in zip(got, esperado[clave]))
        if "variacion" in ref:
            want = R["caso"]["variacion_fifo"] if clave == "caso_fifo" else R["caso"]["variacion_ppp"]
            ok &= abs(v(ref["variacion"]) - want) < 1e-6
    print(f"PLANILLA: {'OK' if ok else 'NO calza'}, {destino.name} con {len(hojas)} hojas de fórmulas")
    return ok


# ---------------------------------------------------------------- texto

def tablas(R):
    t = {}
    nom_caso = {"venta": "Consumo"}
    t["caso_fifo"] = tabla_md(R["caso"]["fifo"], "fifo", nom_caso)
    t["caso_ppp"] = tabla_md(R["caso"]["ppp"], "ppp", nom_caso)
    t["ej_fifo"] = tabla_md(R["ej"]["fifo"], "fifo")
    t["ej_ppp"] = tabla_md(R["ej"]["ppp"], "ppp")
    t["e5"] = tabla_md(R["e5"]["ppp"], "ppp")
    t["e6_fifo"] = tabla_md(R["e6"]["fifo"], "fifo")
    t["e6_ppp"] = tabla_md(R["e6"]["ppp"], "ppp")
    t["e9_fifo"] = tabla_md(R["e9"]["fifo"], "fifo")
    t["e9_ppp"] = tabla_md(R["e9"]["ppp"], "ppp")
    t["e10"] = tabla_md(R["e10"]["fifo"], "fifo")
    return t


def asientos_md(lista):
    fila = lambda *c: "| " + " | ".join(c) + " |"
    out = ["| N.º | Fecha | Cuenta | Debe | Haber |", "|---|---|---|---|---|"]
    for a in lista:
        primero = True
        for k, v in a["debe"].items():
            out.append(fila(str(a["n"]) if primero else "", a["fecha"] if primero else "", k, fmt(v), ""))
            primero = False
        for k, v in a["haber"].items():
            out.append(fila("", "", k, "", fmt(v)))
        out.append(fila("", "", f"*{a['glosa']}*", "", ""))
    return out


def esperados(R, C):
    p = fmt
    e = []
    for filas in tablas(R).values():
        e += filas[2:]
    j = R["ej"]
    e += [p(j["ventas"]), p(j["v1"]), p(j["v2"]), p(j["mb_fifo"]), p(j["mb_ppp"]), p(j["dif"]), p(j["disponible"]),
          p(j["ppp_periodico"]), p(j["cv_periodico"]), p(j["if_periodico"]), p(j["cv_periodico_total"]), p(j["rf"]["costo_ventas"]),
          p(j["rp"]["costo_ventas"]), p(j["rf"]["merma"]), p(j["rp"]["merma"]), p(j["rf"]["saldo"]), p(j["rp"]["saldo"])]
    e += asientos_md(j["asientos"])[2:]
    j = R["e1"]; e += [p(j["cv"]), p(j["mb"]), pct(j["mb_pct"]), p(j["disponible"])]
    j = R["e2"]; e += [p(j["rf"]["costo_ventas"]), p(j["rf"]["saldo"]), p(j["prom"]), p(j["rp"]["costo_ventas"]),
                       p(j["rp"]["saldo"]), p(j["disponible"]), p(j["dif"])]
    j = R["e4"]; e += asientos_md(j["asientos"])[2:] + [p(j["mb"]), pct(j["mb_pct"])]
    j = R["e5"]; e += [p(j["r"]["costo_ventas"]), p(j["r"]["merma"]), p(j["r"]["saldo"]), p(j["disponible"])]
    j = R["e6"]; e += [p(j["ventas"]), p(j["rf"]["costo_ventas"]), p(j["rp"]["costo_ventas"]), p(j["mb_fifo"]),
                       p(j["mb_ppp"]), p(j["dif"]), p(j["disponible"]), fmt(j["prom_periodico"], 2), p(j["cv_periodico"]),
                       p(j["if_periodico"]), p(j["mb_periodico"]), p(j["rf"]["saldo"]), p(j["rp"]["saldo"])]
    j = R["e7"]; e += [p(j["cf"]), p(j["cp"]), p(j["std"]), p(j["var_f"]), p(j["var_p"]), p(j["prom"]),
                       p(j["var_compra"]), p(j["bod_f"]), p(j["bod_p"]), p(j["por_saco"]), p(j["mes_siguiente"])]
    j = R["e8"]; e += [p(j["merma"]), p(j["if_ok"]), p(j["cv_ok"]), p(j["if_mal"]), p(j["cv_mal"]), p(j["inflada"]),
                       pct(j["merma_pct"], 2)]
    j = R["e9"]; e += [p(j["rf"]["costo_ventas"]), p(j["rp"]["costo_ventas"]), p(j["dif"]), p(j["rf"]["saldo"]),
                       p(j["rp"]["saldo"]), p(j["disponible"]), p(j["reposicion"])]
    j = R["e10"]; e += [p(j["neto"]), p(j["ventas"]), p(j["r"]["costo_ventas"]), p(j["mb"]), p(j["resultado"]), p(j["ivd"]),
                        p(j["ivc"]), p(j["pagar"]), p(j["disponible"]), pct(j["mb_pct"])]
    e += asientos_md(j["asientos"])[2:]
    c = R["caso"]
    e += [p(c[k]) for k in ("consumo", "costo_fifo", "saldo_fifo", "ppp_", "costo_ppp", "saldo_ppp", "costo_estandar",
                            "variacion_fifo", "variacion_ppp", "compra", "iva_compra", "total_compra", "var_compra",
                            "bod_fifo", "bod_ppp", "mc", "mc_real", "mc_nuevo", "extra_mes", "saldo_inicial", "sobre_std")]
    e += [p(x) for x in c["var_semana_p"] + c["var_semana_f"] + c["std_semana"] + c["real_semana_p"] + c["real_semana_f"]]
    e += [fmt(c["unit_fifo"], 2), pct(c["caida_mc"]), pct(c["caida_mc_nuevo"])]
    return e


if __name__ == "__main__":
    m, C = cargar_caso()
    R = resultados(m)
    for k, d in R.items():
        print(f"== {k}")
        for kk, v in d.items():
            if kk in ("fifo", "ppp", "asientos"):
                continue
            print(f"   {kk}: {round(v, 4) if isinstance(v, float) else v}")
    if "--tablas" in sys.argv:
        for k, filas in tablas(R).items():
            print(f"\n### {k}"); print("\n".join(filas))
        for k in ("ej", "e4", "e10"):
            print(f"\n### asientos {k}"); print("\n".join(asientos_md(R[k]["asientos"])))
    ok_planilla = planilla(R, m)
    u7, c = C["u7"], R["caso"]
    pares = [("consumo", "consumo"), ("costo_fifo", "costo_fifo"), ("saldo_fifo", "saldo_fifo"), ("ppp", "ppp_"),
             ("costo_ppp", "costo_ppp"), ("saldo_ppp", "saldo_ppp"), ("costo_estandar", "costo_estandar"),
             ("variacion_fifo", "variacion_fifo"), ("variacion_ppp", "variacion_ppp")]
    ok = all(abs(u7[a] - c[b]) < 1e-6 for a, b in pares)
    print("CADENA:", "OK, la tarjeta de las poleras calza con caso_estampa_sur_contable.py" if ok else "NO calza")
    if not (ok and ok_planilla):
        sys.exit(1)
    if "--revisar" in sys.argv:
        md = (AQUI.parents[0] / "unidades" / "07-inventarios.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %")
        faltan = [x for x in esperados(R, C) if x not in md]
        print("REVISIÓN:", "OK, todas las cifras están en el texto" if not faltan else f"faltan {len(faltan)}: {faltan}")
        sys.exit(1 if faltan else 0)
