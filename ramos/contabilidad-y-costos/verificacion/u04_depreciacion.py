"""Unidad 4 · Activo fijo y depreciación.

Recalcula todas las cifras de la unidad desde los datos, genera la planilla
`planillas/u04_depreciacion.xlsx` (los cuatro métodos del ejemplo y la tabla mensual de las dos
máquinas del caso, con fórmulas) y la comprueba evaluando sus fórmulas en Python. Con --revisar
busca cada resultado y cada fila de las tablas generadas en el .md.

    py -3 -P u04_depreciacion.py             # imprime resultados y tablas
    py -3 -P u04_depreciacion.py --revisar   # además los busca en el .md

Todos los datos son originales de esta base (pasados por herramientas/colisiones.py).
Convención declarada: el mes de compra se deprecia completo (como en el curso).
"""
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import importlib.util, re, sys

sys.stdout.reconfigure(encoding="utf-8")
AQUI = Path(__file__).resolve().parent
RAMO = AQUI.parent
IVA = 0.19


# ---------------------------------------------------------------- herramientas

def fmt(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def r0(x):
    return int(Decimal(str(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def asiento(n, fecha, lineas, glosa):
    d = sum(x[1] for x in lineas)
    h = sum(x[2] for x in lineas)
    assert abs(d - h) < 1e-6, (n, d, h)
    return dict(n=n, fecha=fecha, lineas=lineas, glosa=glosa, total=d)


def md_asientos(asientos, dec=0):
    filas = ["| N.º | Fecha | Cuenta | Debe | Haber |", "|---|---|---|---|---|"]
    for a in asientos:
        for k, (c, d, h) in enumerate(a["lineas"]):
            filas.append(f"| {a['n'] if k == 0 else ''} | {a['fecha'] if k == 0 else ''} | {c} | "
                         f"{fmt(d, dec) if d else ''} | {fmt(h, dec) if h else ''} |")
        filas.append(f"| | | *{a['glosa']}* | | |")
    return filas


def md_tabla(enc, filas):
    out = ["| " + " | ".join(enc) + " |", "|" + "---|" * len(enc)]
    out += ["| " + " | ".join(str(x) for x in f) + " |" for f in filas]
    return out


def lineal(C, VR, n):
    return [(C - VR) / n] * n


def por_unidades(C, VR, unidades):
    t = (C - VR) / sum(unidades)
    return [t * u for u in unidades], t


def tasa_saldo(C, VR, n):
    """Tasa que, aplicada n veces sobre el valor libro, deja exactamente el valor residual."""
    return 1 - (VR / C) ** (1 / n)


def saldo_decreciente(C, VR, n):
    g, vl, d = tasa_saldo(C, VR, n), C, []
    for _ in range(n):
        d.append(vl * g); vl -= vl * g
    return d, g


def suma_digitos(C, VR, n):
    s = n * (n + 1) / 2
    return [(C - VR) * (n - k) / s for k in range(n)], s


def valores_libro(C, deps):
    vl, out = C, []
    for d in deps:
        vl -= d; out.append(vl)
    return out


# ---------------------------------------------------------------- datos

# Ejemplo resuelto · Panadería Los Aromos, horno rotatorio comprado el 10-03-2026
EJ = dict(precio=8_460_000, flete=248_000, instalacion=332_000, capacitacion=164_000, residual=940_000,
          vida=6, bandejas=[29_400, 31_800, 28_600, 26_200, 24_300, 21_700], meses_2026=10,
          meses_venta=48, precio_venta=3_980_000)
# 4.1 · costo de una guillotina de imprenta
E1 = dict(lista=6_780_000, descuento=0.05, flete=187_000, seguro=43_500, instalacion=238_500, reparacion=97_000)
# 4.2 · camioneta, lineal
E2 = dict(C=18_960_000, VR=4_560_000, n=6, anios=3)
# 4.3 · servidor comprado el 18-08-2026
E3 = dict(C=3_312_000, VR=432_000, n=4, meses=5)
# 4.4 · terreno y galpón comprados juntos
E4 = dict(pagado=94_300_000, tas_terreno=41_000_000, tas_galpon=61_500_000, vida=40, VR=6_580_000)
# 4.5 · furgón por kilómetros
E5 = dict(C=21_780_000, VR=3_780_000, km=[72_400, 80_600, 65_900, 81_100])
# 4.6 · torno, saldo decreciente
E6 = dict(C=28_800_000, VR=900_000, n=5)
# 4.7 · grúa horquilla, suma de dígitos
E7 = dict(C=15_460_000, VR=1_460_000, n=4)
# 4.8 · venta con pérdida (compra 01-07-2023, venta 30-09-2026)
E8 = dict(C=11_460_000, VR=660_000, meses_vida=108, meses_hasta_2025=30, meses_2026=9, precio=6_840_000)
# 4.9 · ¿cuándo se compró? (saldos al 31-12-2025)
E9 = dict(C=26_400_000, DA=9_900_000, VR=2_640_000, vida_anios=8, meses_2026=5)
# 4.10 · cambio de estimación y decisión de vender
E10 = dict(C=12_640_000, VR=640_000, n=8, revision=3, vida_nueva=6, VR_nuevo=1_240_000, anio_oferta=5,
           oferta=2_900_000, margen_un_anio=1_950_000)


# ---------------------------------------------------------------- resultados

def resultados():
    R = {}
    e = EJ
    C = e["precio"] + e["flete"] + e["instalacion"]
    base = C - e["residual"]
    iva = r0((C + e["capacitacion"]) * IVA)
    L = lineal(C, e["residual"], e["vida"])
    U, tu = por_unidades(C, e["residual"], e["bandejas"])
    S, g = saldo_decreciente(C, e["residual"], e["vida"])
    D, sd = suma_digitos(C, e["residual"], e["vida"])
    for m in (L, U, S, D):
        assert abs(sum(m) - base) < 1e-6
    mensual = L[0] / 12
    dep26 = mensual * e["meses_2026"]
    da_venta = mensual * e["meses_venta"]
    vl_venta = C - da_venta
    util = e["precio_venta"] - vl_venta
    A = [asiento(1, "10-03-2026", [("Maquinaria (horno)", C, 0), ("Gasto en capacitación", e["capacitacion"], 0),
                                   ("IVA crédito fiscal", iva, 0), ("Banco", 0, C + e["capacitacion"] + iva)],
                 "Compra e instalación del horno rotatorio, por transferencia"),
         asiento(2, "31-12-2026", [("Gasto por depreciación", dep26, 0), ("Depreciación acumulada", 0, dep26)],
                 "Depreciación de marzo a diciembre (10 meses)"),
         asiento(3, "fin del año 4", [("Banco", e["precio_venta"], 0), ("Depreciación acumulada", da_venta, 0),
                                      ("Maquinaria (horno)", 0, C), ("Utilidad en venta de activo fijo", 0, util)],
                 "Venta del horno tras 48 meses de uso")]
    R["ej"] = dict(C=C, base=base, iva=iva, banco=C + e["capacitacion"] + iva, L=L, U=U, tu=tu, S=S, g=g, D=D, sd=sd,
                   vl_L=valores_libro(C, L), vl_U=valores_libro(C, U), vl_S=valores_libro(C, S),
                   vl_D=valores_libro(C, D), mensual=mensual, dep26=dep26, vl26=C - dep26,
                   da_venta=da_venta, vl_venta=vl_venta, util=util, asientos=A,
                   acum2_L=sum(L[:2]), acum2_S=sum(S[:2]), acum2_D=sum(D[:2]))
    assert abs(R["ej"]["vl_S"][-1] - e["residual"]) < 1e-6

    e = E1
    precio = e["lista"] * (1 - e["descuento"])
    C = precio + e["flete"] + e["seguro"] + e["instalacion"]
    iva = r0((C + e["reparacion"]) * IVA)
    A1 = [asiento(1, "06-04", [("Maquinaria (guillotina)", C, 0), ("Gasto en reparaciones", e["reparacion"], 0),
                               ("IVA crédito fiscal", iva, 0), ("Banco", 0, C + e["reparacion"] + iva)],
                  "Compra, traslado e instalación de la guillotina")]
    R["e1"] = dict(precio=precio, desc=e["lista"] - precio, C=C, iva=iva, total=C + e["reparacion"] + iva, asientos=A1)

    e = E2
    anual = (e["C"] - e["VR"]) / e["n"]
    R["e2"] = dict(anual=anual, mensual=anual / 12, da=anual * e["anios"], vl=e["C"] - anual * e["anios"],
                   vl_fin=e["VR"])

    e = E3
    anual = (e["C"] - e["VR"]) / e["n"]
    dep = anual / 12 * e["meses"]
    A3 = [asiento(1, "31-12-2026", [("Gasto por depreciación", dep, 0), ("Depreciación acumulada", 0, dep)],
                  "Depreciación del servidor, agosto a diciembre")]
    R["e3"] = dict(anual=anual, mensual=anual / 12, dep=dep, vl=e["C"] - dep, asientos=A3)

    e = E4
    tot = e["tas_terreno"] + e["tas_galpon"]
    terreno = e["pagado"] * e["tas_terreno"] / tot
    galpon = e["pagado"] * e["tas_galpon"] / tot
    dep = (galpon - e["VR"]) / e["vida"]
    mal = (e["pagado"] - e["VR"]) / e["vida"]
    R["e4"] = dict(tot=tot, pct_terreno=e["tas_terreno"] / tot, terreno=terreno, galpon=galpon, dep=dep, mal=mal,
                   exceso=mal - dep)

    e = E5
    deps, t = por_unidades(e["C"], e["VR"], e["km"])
    R["e5"] = dict(tasa=t, deps=deps, vl=valores_libro(e["C"], deps), total_km=sum(e["km"]),
                   lineal=(e["C"] - e["VR"]) / len(e["km"]))

    e = E6
    deps, g = saldo_decreciente(e["C"], e["VR"], e["n"])
    R["e6"] = dict(g=g, deps=deps, vl=valores_libro(e["C"], deps), lineal=(e["C"] - e["VR"]) / e["n"],
                   dif1=deps[0] - (e["C"] - e["VR"]) / e["n"], doble=2 / e["n"])
    assert abs(g - 0.5) < 1e-12

    e = E7
    deps, s = suma_digitos(e["C"], e["VR"], e["n"])
    lin = (e["C"] - e["VR"]) / e["n"]
    R["e7"] = dict(s=s, deps=deps, vl=valores_libro(e["C"], deps), lineal=lin, dif1=deps[0] - lin,
                   acum2=sum(deps[:2]), acum2_lin=2 * lin, vl2=e["C"] - sum(deps[:2]), vl2_lin=e["C"] - 2 * lin)

    e = E8
    m = (e["C"] - e["VR"]) / e["meses_vida"]
    da25 = m * e["meses_hasta_2025"]; d26 = m * e["meses_2026"]; da = da25 + d26
    vl = e["C"] - da; res = e["precio"] - vl
    A8 = [asiento(1, "30-09-2026", [("Gasto por depreciación", d26, 0), ("Depreciación acumulada", 0, d26)],
                  "Depreciación de enero a septiembre, antes de vender"),
          asiento(2, "30-09-2026", [("Banco", e["precio"], 0), ("Depreciación acumulada", da, 0),
                                    ("Pérdida en venta de activo fijo", -res, 0), ("Maquinaria", 0, e["C"])],
                  "Venta de la máquina por transferencia")]
    assert res < 0
    R["e8"] = dict(mensual=m, da25=da25, d26=d26, da=da, vl=vl, perdida=-res, asientos=A8)

    e = E9
    m = (e["C"] - e["VR"]) / (12 * e["vida_anios"])
    meses = e["DA"] / m
    assert abs(meses - round(meses)) < 1e-9
    meses = round(meses)
    d26 = m * e["meses_2026"]
    A9 = [asiento(1, "31-05-2026", [("Gasto por depreciación", d26, 0), ("Depreciación acumulada", 0, d26)],
                  "Depreciación de enero a mayo de 2026")]
    R["e9"] = dict(mensual=m, anual=12 * m, meses=meses, d26=d26, da=e["DA"] + d26, vl=e["C"] - e["DA"] - d26,
                   quedan=12 * e["vida_anios"] - meses - e["meses_2026"], asientos=A9)
    # 40 meses hasta el 31-12-2025 contando el mes de compra: septiembre de 2022
    assert meses == 4 + 12 * 3

    e = E10
    a1 = (e["C"] - e["VR"]) / e["n"]
    da3 = a1 * e["revision"]; vl3 = e["C"] - da3
    resto = e["vida_nueva"] - e["revision"]
    a2 = (vl3 - e["VR_nuevo"]) / resto
    anios_nuevos = e["anio_oferta"] - e["revision"]
    da5 = da3 + a2 * anios_nuevos; vl5 = e["C"] - da5
    perdida = vl5 - e["oferta"]
    mantener = e["margen_un_anio"] + e["VR_nuevo"]
    R["e10"] = dict(a1=a1, da3=da3, vl3=vl3, a2=a2, da5=da5, vl5=vl5, perdida=perdida, mantener=mantener,
                    ventaja=mantener - e["oferta"], vl6=vl5 - a2)
    assert abs(R["e10"]["vl6"] - e["VR_nuevo"]) < 1e-6

    R["caso"] = caso()
    return R


# ---------------------------------------------------------------- caso integrador

def caso():
    spec = importlib.util.spec_from_file_location("caso", AQUI / "caso_estampa_sur_contable.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    C = m.calcular()["u4"]
    p = m.MAQUINAS["pulpo 6 (inicial)"]; s = m.MAQUINAS["segunda máquina"]
    dp = (p["costo"] - p["residual"]) / p["vida_meses"]
    ds = (s["costo"] - s["residual"]) / s["vida_meses"]
    da0 = -m.BALANCE_INICIAL["activo"]["Depreciación acumulada"]
    usados = da0 / dp
    assert abs(usados - round(usados)) < 1e-9
    # tabla mensual desde marzo de 2026
    filas, dap, das = [], da0, 0.0
    restan_p = p["vida_meses"] - round(usados)
    for k in range(1, restan_p + 1):
        d1 = max(0, min(dp, p["costo"] - dap - p["residual"])); dap += d1
        d2 = max(0, min(ds, s["costo"] - das - s["residual"])); das += d2
        filas.append(dict(k=k, dp=d1, dap=dap, vlp=p["costo"] - dap, ds=d2, das=das, vls=s["costo"] - das))
    anual = []
    for a in range(restan_p // 12):
        bloque = filas[12 * a:12 * a + 12]
        anual.append(dict(a=a + 1, dp=sum(f["dp"] for f in bloque), vlp=bloque[-1]["vlp"],
                          ds=sum(f["ds"] for f in bloque), vls=bloque[-1]["vls"]))
    f1 = filas[0]
    return dict(C=C, dp=dp, ds=ds, total=dp + ds, da0=da0, usados=round(usados), restan_p=restan_p,
                vl0_p=p["costo"] - da0, f1=f1, filas=filas, anual=anual,
                maquinaria=p["costo"] + s["costo"], da_31=da0 + dp + ds,
                vl_31=p["costo"] + s["costo"] - da0 - dp - ds, cif_fijos=m.calcular()["u9"]["cif_fijos_mes"],
                p=p, s=s)


def chequear_caso(R):
    c = R["caso"]
    C = c["C"]
    return (abs(C["depreciacion_mensual"]["pulpo 6 (inicial)"] - c["dp"]) < 1e-6
            and abs(C["depreciacion_mensual"]["segunda máquina"] - c["ds"]) < 1e-6
            and abs(C["total_mes"] - c["total"]) < 1e-6
            and abs(c["filas"][59]["vls"] - c["s"]["residual"]) < 1e-6
            and abs(c["filas"][-1]["vlp"] - c["p"]["residual"]) < 1e-6)


# ---------------------------------------------------------------- tablas para el .md

def tablas(R):
    T = {}
    j = R["ej"]
    T["ej_metodos"] = md_tabla(["Año", "Bandejas", "Lineal", "Unidades producidas", "Saldo decreciente", "Suma de dígitos"],
                               [[k + 1, fmt(EJ["bandejas"][k]), fmt(j["L"][k]), fmt(j["U"][k]), fmt(j["S"][k]),
                                 fmt(j["D"][k])] for k in range(EJ["vida"])]
                               + [["**Total**", fmt(sum(EJ["bandejas"])), fmt(sum(j["L"])), fmt(sum(j["U"])),
                                   fmt(sum(j["S"])), fmt(sum(j["D"]))]])
    T["ej_vl"] = md_tabla(["Al final del año", "Lineal", "Unidades producidas", "Saldo decreciente", "Suma de dígitos"],
                          [[k + 1, fmt(j["vl_L"][k]), fmt(j["vl_U"][k]), fmt(j["vl_S"][k]), fmt(j["vl_D"][k])]
                           for k in range(EJ["vida"])])
    T["ej_diario"] = md_asientos(j["asientos"])
    T["e1"] = md_asientos(R["e1"]["asientos"])
    T["e3"] = md_asientos(R["e3"]["asientos"])
    j = R["e5"]
    T["e5"] = md_tabla(["Año", "Kilómetros", "Depreciación", "Valor libro"],
                       [[k + 1, fmt(E5["km"][k]), fmt(j["deps"][k]), fmt(j["vl"][k])] for k in range(len(E5["km"]))])
    j = R["e6"]
    vlini = [E6["C"]] + j["vl"][:-1]
    T["e6"] = md_tabla(["Año", "Valor libro al inicio", "Depreciación (50 %)", "Valor libro al final", "Lineal"],
                       [[k + 1, fmt(vlini[k]), fmt(j["deps"][k]), fmt(j["vl"][k]), fmt(j["lineal"])]
                        for k in range(E6["n"])])
    j = R["e7"]
    T["e7"] = md_tabla(["Año", "Fracción", "Suma de dígitos", "Valor libro", "Lineal", "Valor libro lineal"],
                       [[k + 1, f"{E7['n'] - k}/{int(j['s'])}", fmt(j["deps"][k]), fmt(j["vl"][k]), fmt(j["lineal"]),
                         fmt(E7["C"] - j["lineal"] * (k + 1))] for k in range(E7["n"])])
    T["e8"] = md_asientos(R["e8"]["asientos"])
    T["e9"] = md_asientos(R["e9"]["asientos"])
    c = R["caso"]
    T["caso_anual"] = md_tabla(["Año (de marzo a febrero)", "Depreciación pulpo 6", "Valor libro pulpo 6",
                                "Depreciación segunda máquina", "Valor libro segunda máquina"],
                               [[f"{a['a']} (03-{2025 + a['a']} a 02-{2026 + a['a']})", fmt(a["dp"]), fmt(a["vlp"]),
                                 fmt(a["ds"]), fmt(a["vls"])] for a in c["anual"]])
    T["caso_a"] = md_asientos([asiento("Ajuste", "31-03", [("Depreciación de maquinaria (costo indirecto)", c["total"], 0),
                                                           ("Depreciación acumulada", 0, c["total"])],
                                       "Depreciación de marzo de las dos máquinas")], dec=2)
    return T


# ---------------------------------------------------------------- planilla (con fórmulas)

def evaluar(ws):
    """Evalúa las fórmulas de una hoja (+ − × ÷ ^, SUM, MAX, MIN, ROUND, ABS)."""
    cache = {}

    def val(ref):
        ref = ref.replace("$", "")
        if ref in cache:
            return cache[ref]
        v = ws[ref].value
        if isinstance(v, str) and v.startswith("="):
            v = calc(v[1:])
        cache[ref] = 0 if v is None else v
        return cache[ref]

    def rng(c1, f1, c2, f2):
        from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
        return [val(f"{gl(c)}{f}") for c in range(ci(c1), ci(c2) + 1) for f in range(f1, f2 + 1)]

    def _round(x, d=0):
        return float(Decimal(str(x)).quantize(Decimal(1).scaleb(-int(d)), rounding=ROUND_HALF_UP))

    def _sum(*a):
        return sum(sum(x) if isinstance(x, list) else x for x in a)

    def calc(f):
        f = re.sub(r"\$?([A-Z]{1,2})\$?(\d+):\$?([A-Z]{1,2})\$?(\d+)", r'rng("\1",\2,"\3",\4)', f)
        f = re.sub(r"(?<![A-Za-z\"])\$?([A-Z]{1,2})\$?(\d+)(?![\d\"])", r'val("\1\2")', f)
        f = f.replace("SUM(", "_sum(").replace("ROUND(", "_round(").replace("MAX(", "max(") \
             .replace("MIN(", "min(").replace("ABS(", "abs(").replace("^", "**")
        return eval(f, {"val": val, "rng": rng, "_sum": _sum, "_round": _round, "max": max, "min": min, "abs": abs})

    return val


def planilla(R):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    negrita, gris, azul = Font(bold=True), PatternFill("solid", fgColor="F2F1EF"), Font(color="1F4E9E")
    wb = Workbook()

    # Hoja 1 · los cuatro métodos del ejemplo
    ws = wb.active; ws.title = "Métodos"
    ws["A1"] = "Panadería Los Aromos · horno rotatorio (datos en azul; lo demás es fórmula)"; ws["A1"].font = negrita
    datos = [("Precio neto", EJ["precio"]), ("Flete", EJ["flete"]), ("Instalación", EJ["instalacion"]),
             ("Costo del activo", "=B3+B4+B5"), ("Valor residual", EJ["residual"]), ("Vida útil (años)", EJ["vida"]),
             ("Base depreciable", "=B6-B7"), ("Tasa de saldo decreciente", "=1-(B7/B6)^(1/B8)"),
             ("Suma de los dígitos", "=B8*(B8+1)/2"), ("Bandejas en toda la vida", "=SUM(B16:B21)"),
             ("Depreciación por bandeja", "=B9/B12")]
    for k, (n, v) in enumerate(datos, start=3):
        ws.cell(k, 1, n); c = ws.cell(k, 2, v)
        if not (isinstance(v, str) and v.startswith("=")):
            c.font = azul
    enc = ["Año", "Bandejas", "Lineal", "Unidades", "Saldo decreciente", "Suma de dígitos",
           "VL lineal", "VL unidades", "VL saldo decr.", "VL suma díg."]
    for k, h in enumerate(enc, start=1):
        c = ws.cell(15, k, h); c.font = negrita; c.fill = gris
    for k in range(EJ["vida"]):
        f = 16 + k
        ws.cell(f, 1, k + 1); ws.cell(f, 2, EJ["bandejas"][k]).font = azul
        ws.cell(f, 3, "=$B$9/$B$8")
        ws.cell(f, 4, f"=B{f}*$B$13")
        ws.cell(f, 5, f"=$B$6*$B$10" if k == 0 else f"=I{f - 1}*$B$10")
        ws.cell(f, 6, f"=($B$8-A{f}+1)/$B$11*$B$9")
        for col, dcol in zip("GHIJ", "CDEF"):
            prev = "$B$6" if k == 0 else f"{col}{f - 1}"
            ws[f"{col}{f}"] = f"={prev}-{dcol}{f}"
    ws["A22"] = "Total"; ws["A22"].font = negrita
    for col in "BCDEF":
        ws[f"{col}22"] = f"=SUM({col}16:{col}21)"
    ws["A24"], ws["B24"] = "Depreciación mensual (lineal)", "=C16/12"
    ws["A25"], ws["B25"] = "Depreciación de 2026 (10 meses)", "=B24*10"
    ws["A26"], ws["B26"] = "Valor libro al 31-12-2026", "=B6-B25"
    ws["A27"], ws["B27"] = "Precio de venta tras 48 meses", EJ["precio_venta"]; ws["B27"].font = azul
    ws["A28"], ws["B28"] = "Utilidad (pérdida) en la venta", "=B27-(B6-B24*48)"
    ws.column_dimensions["A"].width = 32
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 15

    # Hoja 2 · la tabla mensual de las dos máquinas del caso
    w2 = wb.create_sheet("Caso marzo")
    c = R["caso"]
    w2["A1"] = "Estampa Sur · depreciación mensual desde marzo de 2026"; w2["A1"].font = negrita
    d2 = [("Pulpo 6 · costo", c["p"]["costo"]), ("Pulpo 6 · valor residual", c["p"]["residual"]),
          ("Pulpo 6 · vida útil (meses)", c["p"]["vida_meses"]), ("Pulpo 6 · dep. acumulada al 01-03", c["da0"]),
          ("Segunda máquina · costo neto", c["s"]["costo"]), ("Segunda máquina · valor residual", c["s"]["residual"]),
          ("Segunda máquina · vida útil (meses)", c["s"]["vida_meses"]),
          ("Pulpo 6 · depreciación mensual", "=(B3-B4)/B5"), ("Segunda máquina · depreciación mensual", "=(B7-B8)/B9")]
    for k, (n, v) in enumerate(d2, start=3):
        w2.cell(k, 1, n); x = w2.cell(k, 2, v)
        if not (isinstance(v, str) and v.startswith("=")):
            x.font = azul
    enc = ["Mes", "Período", "Dep. pulpo 6", "DA pulpo 6", "VL pulpo 6", "Dep. segunda", "DA segunda",
           "VL segunda", "Total del mes"]
    for k, h in enumerate(enc, start=1):
        x = w2.cell(14, k, h); x.font = negrita; x.fill = gris
    for k, fila in enumerate(c["filas"]):
        f = 15 + k
        mes = (2 + k) % 12 + 1; anio = 2026 + (2 + k) // 12
        w2.cell(f, 1, k + 1); w2.cell(f, 2, f"{mes:02d}-{anio}")
        pda = "$B$6" if k == 0 else f"D{f - 1}"
        sda = "0" if k == 0 else f"G{f - 1}"
        w2.cell(f, 3, f"=MAX(0,MIN($B$10,$B$3-{pda}-$B$4))")
        w2.cell(f, 4, f"={pda}+C{f}")
        w2.cell(f, 5, f"=$B$3-D{f}")
        w2.cell(f, 6, f"=MAX(0,MIN($B$11,$B$7-{sda}-$B$8))")
        w2.cell(f, 7, f"={sda}+F{f}")
        w2.cell(f, 8, f"=$B$7-G{f}")
        w2.cell(f, 9, f"=C{f}+F{f}")
    ultima = 14 + len(c["filas"])
    w2.column_dimensions["A"].width = 38
    for col in "BCDEFGHI":
        w2.column_dimensions[col].width = 15

    destino = RAMO / "planillas" / "u04_depreciacion.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)

    v1, v2 = evaluar(ws), evaluar(w2)
    j = R["ej"]
    cerca = lambda a, b: abs(a - b) < 1e-6
    chk = {
        "planilla: los cuatro métodos, año por año": all(
            cerca(v1(f"C{16 + k}"), j["L"][k]) and cerca(v1(f"D{16 + k}"), j["U"][k]) and cerca(v1(f"E{16 + k}"), j["S"][k])
            and cerca(v1(f"F{16 + k}"), j["D"][k]) for k in range(EJ["vida"])),
        "planilla: los cuatro métodos suman la base": all(cerca(v1(f"{col}22"), j["base"]) for col in "CDEF"),
        "planilla: los cuatro terminan en el valor residual": all(cerca(v1(f"{col}21"), EJ["residual"]) for col in "GHIJ"),
        "planilla: 2026 y venta": cerca(v1("B25"), j["dep26"]) and cerca(v1("B28"), j["util"]),
        "planilla: caso, mes a mes": all(cerca(v2(f"C{15 + k}"), fl["dp"]) and cerca(v2(f"F{15 + k}"), fl["ds"])
                                         and cerca(v2(f"E{15 + k}"), fl["vlp"]) and cerca(v2(f"H{15 + k}"), fl["vls"])
                                         for k, fl in enumerate(c["filas"])),
        "planilla: caso, total de marzo": cerca(v2("I15"), c["total"]),
        "planilla: caso, terminan en su residual": cerca(v2(f"E{ultima}"), 0) and cerca(v2("H74"), c["s"]["residual"]),
    }
    return destino, chk


# ---------------------------------------------------------------- esperados

def esperados(R, T):
    p = fmt
    e = []
    j = R["ej"]
    e += [p(j[k]) for k in ("C", "base", "iva", "banco", "mensual", "dep26", "vl26", "da_venta", "vl_venta", "util",
                            "acum2_L", "acum2_S", "acum2_D")]
    e += [p(j["tu"]), fmt(100 * j["g"], 2) + " %", str(int(j["sd"]))]
    j = R["e1"]; e += [p(j[k]) for k in ("precio", "desc", "C", "iva", "total")]
    j = R["e2"]; e += [p(j[k]) for k in ("anual", "mensual", "da", "vl", "vl_fin")]
    j = R["e3"]; e += [p(j[k]) for k in ("anual", "mensual", "dep", "vl")]
    j = R["e4"]; e += [p(j[k]) for k in ("tot", "terreno", "galpon", "dep", "mal", "exceso")] + [fmt(100 * j["pct_terreno"], 0) + " %"]
    j = R["e5"]; e += [p(j["tasa"]), p(j["total_km"]), p(j["lineal"])]
    j = R["e6"]; e += [p(j["lineal"]), p(j["dif1"])]
    j = R["e7"]; e += [p(j[k]) for k in ("lineal", "dif1", "acum2", "acum2_lin", "vl2", "vl2_lin")]
    j = R["e8"]; e += [p(j[k]) for k in ("mensual", "da25", "d26", "da", "vl", "perdida")]
    j = R["e9"]; e += [p(j[k]) for k in ("mensual", "anual", "d26", "da", "vl")] + [str(j["meses"]), str(j["quedan"])]
    j = R["e10"]; e += [p(j[k]) for k in ("a1", "da3", "vl3", "a2", "da5", "vl5", "perdida", "mantener", "ventaja")]
    c = R["caso"]
    e += [p(c["dp"]), p(c["ds"], 2), p(c["total"], 2), p(c["da0"]), p(c["vl0_p"]), str(c["usados"]), str(c["restan_p"]),
          p(c["f1"]["dap"]), p(c["f1"]["vlp"]), p(c["f1"]["vls"], 2), p(c["maquinaria"]), p(c["da_31"], 2),
          p(c["vl_31"], 2), p(c["cif_fijos"], 2), p(12 * c["ds"]), p(12 * c["dp"])]
    for filas in T.values():
        e += filas
    return e


# ---------------------------------------------------------------- principal

if __name__ == "__main__":
    R = resultados()
    T = tablas(R)
    for k, v in R.items():
        if k == "caso":
            continue
        print(f"== {k}")
        for kk, vv in v.items():
            if kk != "asientos":
                print(f"   {kk}: {vv}")
    c = R["caso"]
    print("== caso", {k: v for k, v in c.items() if k not in ("C", "filas", "anual")})
    for k, filas in T.items():
        print(f"\n--- tabla {k}")
        print("\n".join(filas))
    destino, chk = planilla(R)
    chk["caso: la depreciación de las dos máquinas calza con el script"] = chequear_caso(R)
    print(f"\nPlanilla: {destino}")
    print("CHEQUEOS:")
    for k, v in chk.items():
        print(f"   {'OK ' if v else 'MAL'} {k}")
    if not all(chk.values()):
        sys.exit(1)
    print("CADENA: OK, la pieza del caso calza con caso_estampa_sur_contable.py")
    if "--revisar" in sys.argv:
        md = (RAMO / "unidades" / "04-depreciacion.md").read_text(encoding="utf-8")
        md = md.replace("{,}", ",").replace("{.}", ".").replace("\\$", "$").replace("\\ \\%", " %")
        faltan = [x for x in esperados(R, T) if x not in md]
        print("\nREVISIÓN:", "OK, todas las cifras están en el texto" if not faltan else f"faltan {len(faltan)}: {faltan}")
        sys.exit(1 if faltan else 0)
